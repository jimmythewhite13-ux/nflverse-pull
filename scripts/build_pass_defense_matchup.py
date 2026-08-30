"""
Builds "Pass Defense Matchup" -- Part B of claude_code_spec_defensive_matchup_engine.md
(Version 1 of the post-backtest roadmap). Team-level, 5-section decay-weighted pattern,
identical shape to Front Seven Index / Secondary Index (no rookie concept, no Section 2B,
no individual-player Section 6 -- a defense's real pass-defense rate stats are already
team-level).

ALL FIVE scored metrics are real, phase-SPLIT pass-defense-allowed rates (efficiency.
compute_team_season_matchup_metrics -- confirmed with the user before building that
efficiency.py's OLD epa_def/success_def are COMBINED across pass+run plays, which would put
the identical non-split number in both this tab and Run Defense Matchup, defeating the
entire point of this spec):
  - EPA/Dropback Allowed
  - Pass Success Rate Allowed
  - Completion % Allowed
  - NY/A Allowed -- NOT recomputed: efficiency.compute_team_season_efficiency's existing
    `NY/A Allowed (Def)` is ALREADY pass-only (built from real pass_attempt/sack plays, no
    rush yardage), so this tab merges that column in directly rather than duplicating it.
  - Explosive Pass Rate Allowed

All five are "lower raw value = better defense" -- Section 5's Z-scores are uniformly
sign-flipped (league average minus raw), same convention efficiency.py's own epa_def/
success_def already use.

PLUS three REFERENCED (not recomputed) columns, live-linked from Front Seven Index's own
Section 1 by a SUMPRODUCT-array MATCH on (Team, Season) -- avoids a two-criteria array/CSE
formula without needing to add a helper key column to Front Seven Index's already-shipped
Section 1. CONTEXT ONLY, not part of this tab's own weighted composite (same reasoning
Front Seven Index itself already uses for its own Blitz Rate/Avg Box Count):
  - Sack Rate (real, from Front Seven Index)
  - "Pressure proxy" -- QB Hit Rate (real, from Front Seven Index; the spec's own wording
    calls for a pressure proxy, and Front Seven Index doesn't have a separate pressure
    number beyond Sack Rate/QB Hit Rate -- QB Hit Rate captures pressure that didn't finish
    as a sack, a defensible real proxy, documented as such rather than presented as
    something more precise than it is)
  - Blitz Rate (real nflverse participation data, from Front Seven Index)

NOT VALIDATED YET: this tab is real, verified-correct plumbing -- there is no backtesting
harness yet to prove phase-specific matchups actually improve predictions over the existing
PPG-based blend. Stated explicitly on the tab itself, per the spec's own instruction. Feeds
Part C's matchup differential on Week 1 Matchups, which adds this ALONGSIDE the existing
PPG-based prediction rather than replacing it, specifically so a future backtest can compare
the two directly.

Usage:
    uv run python scripts/build_pass_defense_matchup.py "C:\\path\\to\\NFL_Prediction_Model.xlsx"
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from nflverse_pull.efficiency import (  # noqa: E402
    compute_team_season_efficiency,
    compute_team_season_matchup_metrics,
    fetch_pbp,
)

HISTORICAL_YEARS = [2023, 2024, 2025]
SHEET_NAME = "Pass Defense Matchup"
FRONT_SEVEN_SHEET = "Front Seven Index"
SPECIAL_TEAMS_SHEET = "Special Teams Index"
TEAM_RATINGS_SHEET = "Team Ratings"

# Front Seven Index's own real Section 1 range (Team | Season | Sack Rate | TFL Rate |
# QB Hit Rate | Blitz Rate | Avg Box Count) -- verified against that tab's own build script:
# sec1_first_row=5, 32 teams x 3 years = 96 rows.
FRONT_SEVEN_SEC1_FIRST_ROW = 5
FRONT_SEVEN_SEC1_LAST_ROW = 100

METRICS = [
    {"key": "epa_db", "col": "EPA/Dropback Allowed (Def)", "label": "EPA/Dropback\nAllowed",
     "fmt": "0.000"},
    {"key": "psr", "col": "Pass Success Rate Allowed (Def)", "label": "Pass Success\nRate Allowed",
     "fmt": "0.00%"},
    {"key": "comp", "col": "Completion % Allowed (Def)", "label": "Completion %\nAllowed",
     "fmt": "0.00%"},
    {"key": "nya", "col": "NY/A Allowed (Def)", "label": "NY/A\nAllowed", "fmt": "0.00"},
    {"key": "exp", "col": "Explosive Pass Rate Allowed (Def)",
     "label": "Explosive Pass\nRate Allowed", "fmt": "0.00%"},
]

TEAM_ORDER = [
    "Buffalo Bills", "Miami Dolphins", "New England Patriots", "New York Jets",
    "Baltimore Ravens", "Cincinnati Bengals", "Cleveland Browns", "Pittsburgh Steelers",
    "Houston Texans", "Indianapolis Colts", "Jacksonville Jaguars", "Tennessee Titans",
    "Denver Broncos", "Kansas City Chiefs", "Las Vegas Raiders", "Los Angeles Chargers",
    "Dallas Cowboys", "New York Giants", "Philadelphia Eagles", "Washington Commanders",
    "Chicago Bears", "Detroit Lions", "Green Bay Packers", "Minnesota Vikings",
    "Atlanta Falcons", "Carolina Panthers", "New Orleans Saints", "Tampa Bay Buccaneers",
    "Arizona Cardinals", "Los Angeles Rams", "San Francisco 49ers", "Seattle Seahawks",
]

TITLE_FONT = Font(name="Arial", size=10, bold=True)
TITLE_FILL = PatternFill("solid", fgColor="FFD9E1F2")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_ALIGN = Alignment(wrap_text=True, horizontal="center", vertical="center")
INPUT_FONT = Font(name="Arial", size=10, color="FF0000FF")
FORMULA_FONT = Font(name="Arial", size=10, color="FF000000")
LINK_FONT = Font(name="Arial", size=10, color="FF008000")
NOTE_FONT = Font(name="Arial", size=9, color="FF808080")
ASSUMPTION_FILL = PatternFill("solid", fgColor="FFFFFF00")


def _section_title(ws, row: int, last_col: int, text: str) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL


def _header_row(ws, row: int, headers: list[str], height: float = 28) -> None:
    ws.row_dimensions[row].height = height
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN


def add_model_assumptions_weights(wb: openpyxl.Workbook) -> None:
    ws = wb["Model Assumptions"]
    title_row = 86
    ws.merge_cells(f"A{title_row}:D{title_row}")
    title = ws.cell(row=title_row, column=1, value=(
        "Pass Defense Matchup Weighting & Conversion (pts per std. dev.; reuses QB Index's "
        "points-scale constants C37/C38 -- see 'Pass Defense Matchup' tab)"
    ))
    title.font = HEADER_FONT
    title.fill = HEADER_FILL

    rows = [
        (87, "EPA/Dropback Allowed Weight (Pass Defense Matchup, pts per SD)", 0.35,
         "The most complete single pass-defense value stat -- weighted highest, same "
         "logic as every other EPA-based weighting in this model."),
        (88, "Pass Success Rate Allowed Weight (Pass Defense Matchup, pts per SD)", 0.25,
         "Correlates with EPA but isolates consistency (positive-EPA-allowed rate) "
         "rather than magnitude."),
        (89, "Completion % Allowed Weight (Pass Defense Matchup, pts per SD)", 0.15,
         "Traditional counting stat -- doesn't account for depth of target or "
         "yards-after-catch the way EPA does, weighted lower as a sanity check."),
        (90, "NY/A Allowed Weight (Pass Defense Matchup, pts per SD)", 0.15,
         "Traditional counting stat, real and pass-only (efficiency.py's existing "
         "NY/A Allowed (Def) -- not recomputed here)."),
        (91, "Explosive Pass Rate Allowed Weight (Pass Defense Matchup, pts per SD)", 0.10,
         "A real tail-risk signal (big-play rate allowed) distinct from the average-"
         "based metrics above -- weighted lowest as a supplementary signal, not the "
         "primary one."),
        (92, "Pass Defense Matchup Points-to-Game-Points Conversion", 0.10,
         "A starting guess, like every other coefficient in this model -- a "
         "structurally different, dedicated constant, NOT reused from any other "
         "tab's conversion factor (this measures a phase-specific MATCHUP "
         "differential, not a team-level quality adjustment or a Replacement Value "
         "swap). NOT YET VALIDATED against real outcomes -- see the tab's own "
         "closing note and claude_code_spec_defensive_matchup_engine.md."),
    ]
    for row, label, value, note in rows:
        ws.cell(row=row, column=2, value=label)
        c = ws.cell(row=row, column=3, value=value)
        c.font = INPUT_FONT
        c.fill = ASSUMPTION_FILL
        c.number_format = "0.00"
        n = ws.cell(row=row, column=4, value=note)
        n.font = NOTE_FONT
        n.alignment = Alignment(wrap_text=True, vertical="top")


def _pull_data() -> pd.DataFrame:
    print(f"Pulling {HISTORICAL_YEARS} play-by-play data for pass-defense matchup metrics...")
    pbp = fetch_pbp(HISTORICAL_YEARS)
    matchup = compute_team_season_matchup_metrics(pbp)
    efficiency = compute_team_season_efficiency(pbp)

    # NY/A Allowed (Def) is real and ALREADY pass-only in efficiency.py -- merged in here,
    # not recomputed (see module docstring).
    nya = efficiency[["Team", "Season", "NY/A Allowed (Def)"]]
    out = matchup.merge(nya, on=["Team", "Season"], how="left")
    return out


def build(workbook_path: str) -> dict:
    season_stats = _pull_data()

    wb = openpyxl.load_workbook(workbook_path)
    add_model_assumptions_weights(wb)

    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    for candidate in (SPECIAL_TEAMS_SHEET, FRONT_SEVEN_SHEET):
        if candidate in wb.sheetnames:
            insert_after = candidate
            break
    else:
        insert_after = wb.sheetnames[0]
    ws = wb.create_sheet(SHEET_NAME, index=wb.sheetnames.index(insert_after) + 1)
    ws.column_dimensions["A"].width = 20.0
    ws.column_dimensions["C"].width = 20.0

    ws.merge_cells("A1:J1")
    t = ws.cell(row=1, column=1, value=(
        "Pass Defense Matchup -- Multi-Year Decay-Weighted TEAM-Level Pass Defense Rating "
        "(EPA/Dropback, Pass Success Rate, Completion %, NY/A, Explosive Pass Rate, all "
        "ALLOWED; real, phase-SPLIT pbp metrics). NOT YET VALIDATED against real outcomes "
        "-- see the closing note. Feeds Part C's matchup differential on Week 1 Matchups, "
        "ALONGSIDE (not replacing) the existing PPG-based prediction."
    ))
    t.font = Font(name="Arial", size=12, bold=True)

    # ==== Section 1: Raw 3-year TEAM-level data + referenced Front Seven context ===========
    sec1_first_row = 5
    sec1_last_row = sec1_first_row + len(season_stats) - 1
    _section_title(
        ws, 3, 10,
        "Section 1 — Raw 3-Year TEAM-Level Pass-Defense-Allowed Rates (real nflverse "
        "play-by-play, phase-split -- see this tab's opening note). Sack Rate / Pressure "
        "Proxy (QB Hit Rate) / Blitz Rate (H/I/J) are REFERENCED, not recomputed -- live "
        "links to Front Seven Index's own real Section 1 data, CONTEXT ONLY, not part of "
        "the weighted composite in Section 5.",
    )
    _header_row(ws, 4, [
        "Team", "Season", *[m["label"] for m in METRICS],
        "Sack Rate\n(ref, context only)", "Pressure Proxy\n(QB Hit Rate, ref)",
        "Blitz Rate\n(ref, context only)",
    ])
    for i, r in enumerate(season_stats.to_dict("records")):
        row = sec1_first_row + i
        values = [r["Team"], int(r["Season"])]
        for m in METRICS:
            v = r.get(m["col"])
            values.append(float(v) if pd.notna(v) else None)
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.font = INPUT_FONT
        for j, m in enumerate(METRICS):
            ws.cell(row=row, column=3 + j).number_format = m["fmt"]

        fs_range = (
            f"('{FRONT_SEVEN_SHEET}'!$A${FRONT_SEVEN_SEC1_FIRST_ROW}:"
            f"$A${FRONT_SEVEN_SEC1_LAST_ROW}=A{row})*"
            f"('{FRONT_SEVEN_SHEET}'!$B${FRONT_SEVEN_SEC1_FIRST_ROW}:"
            f"$B${FRONT_SEVEN_SEC1_LAST_ROW}=B{row})"
        )
        for col, fs_col in ((8, "C"), (9, "E"), (10, "F")):
            formula = (
                f"=IFERROR(INDEX('{FRONT_SEVEN_SHEET}'!${fs_col}${FRONT_SEVEN_SEC1_FIRST_ROW}:"
                f"${fs_col}${FRONT_SEVEN_SEC1_LAST_ROW},SUMPRODUCT(MATCH(1,{fs_range},0))),\"\")"
            )
            cell = ws.cell(row=row, column=col, value=formula)
            cell.font = LINK_FONT
            cell.number_format = "0.00%"

    team_range = f"$A${sec1_first_row}:$A${sec1_last_row}"
    season_range = f"$B${sec1_first_row}:$B${sec1_last_row}"
    sec1_col_of = {m["key"]: get_column_letter(3 + j) for j, m in enumerate(METRICS)}
    metric_ranges = {
        m["key"]: (
            f"${sec1_col_of[m['key']]}${sec1_first_row}:${sec1_col_of[m['key']]}${sec1_last_row}"
        )
        for m in METRICS
    }

    # ==== Section 2: League average per season ==============================================
    sec2_title_row = sec1_last_row + 2
    sec2_header_row = sec2_title_row + 1
    season_rows = {yr: sec2_header_row + 1 + i for i, yr in enumerate(HISTORICAL_YEARS)}

    _section_title(
        ws, sec2_title_row, 6,
        "Section 2 — League Average per Season (simple average across all 32 teams)",
    )
    _header_row(ws, sec2_header_row, ["Season", *[m["label"] for m in METRICS]], height=20)

    for yr in HISTORICAL_YEARS:
        row = season_rows[yr]
        c = ws.cell(row=row, column=1, value=yr)
        c.font = INPUT_FONT
        for j, m in enumerate(METRICS):
            formula = f"=AVERAGEIF({season_range},{yr},{metric_ranges[m['key']]})"
            cell = ws.cell(row=row, column=2 + j, value=formula)
            cell.font = FORMULA_FONT
            cell.number_format = m["fmt"]

    season_avg_col = {m["key"]: get_column_letter(2 + j) for j, m in enumerate(METRICS)}

    # ==== Section 3: Per-team decay-weighted, regressed baseline ===========================
    sec3_title_row = season_rows[HISTORICAL_YEARS[-1]] + 2
    sec3_header_row = sec3_title_row + 1
    sec3_first_row = sec3_header_row + 1
    n_teams = len(TEAM_ORDER)
    sec3_last_row = sec3_first_row + n_teams - 1
    sec3_last_col = 2 + len(METRICS) * 7

    _section_title(
        ws, sec3_title_row, sec3_last_col,
        "Section 3 — Per-Team 3-Yr Decay-Weighted, Regressed Baseline. A missing year "
        "substitutes THAT season's own Section 2 league average (no rookie concept for a "
        "team-level stat).",
    )
    headers = ["Team", "Years of\nReal History"]
    metric_block_start_col: dict[str, int] = {}
    col_cursor = 3
    for m in METRICS:
        metric_block_start_col[m["key"]] = col_cursor
        headers += [
            f"{m['label']} Y-1", f"{m['label']} Y-2", f"{m['label']} Y-3",
            f"Weighted\n{m['label']} Avg\n(3-Yr decay)", f"Team History\n{m['label']}",
            f"League Baseline\n{m['label']} (Y-1)", f"Projected 3-Yr\n{m['label']} Baseline",
        ]
        col_cursor += 7
    _header_row(ws, sec3_header_row, headers)

    for i, team in enumerate(TEAM_ORDER):
        row = sec3_first_row + i
        ws.cell(row=row, column=1, value=team).font = FORMULA_FONT

        years_hist_terms = "+".join(
            f"--(COUNTIFS({team_range},$A{row},{season_range},'Model Assumptions'!$C$18-{k})>0)"
            for k in (1, 2, 3)
        )
        yh = ws.cell(row=row, column=2, value=f"={years_hist_terms}")
        yh.font = FORMULA_FONT
        yh.number_format = "0"

        for m in METRICS:
            base = metric_block_start_col[m["key"]]
            y1, y2, y3, wavg, th, lb, pb = (get_column_letter(base + k) for k in range(7))
            mrange = metric_ranges[m["key"]]
            sec2_col = season_avg_col[m["key"]]

            def _ysub(offset: int, mrange=mrange, sec2_col=sec2_col) -> str:
                season_avg_lookup = (
                    f"INDEX(${sec2_col}${season_rows[HISTORICAL_YEARS[0]]}:"
                    f"${sec2_col}${season_rows[HISTORICAL_YEARS[-1]]},"
                    f"MATCH('Model Assumptions'!$C$18-{offset},"
                    f"$A${season_rows[HISTORICAL_YEARS[0]]}:"
                    f"$A${season_rows[HISTORICAL_YEARS[-1]]},0))"
                )
                return (
                    f"=IF(COUNTIFS({team_range},$A{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset})=0,{season_avg_lookup},"
                    f"SUMIFS({mrange},{team_range},$A{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset}))"
                )

            f_y1 = ws.cell(row=row, column=base, value=_ysub(1))
            f_y2 = ws.cell(row=row, column=base + 1, value=_ysub(2))
            f_y3 = ws.cell(row=row, column=base + 2, value=_ysub(3))
            f_wavg = ws.cell(row=row, column=base + 3, value=(
                f"=({y1}{row}*1+{y2}{row}*'Model Assumptions'!$C$20+"
                f"{y3}{row}*('Model Assumptions'!$C$20^2))/"
                f"(1+'Model Assumptions'!$C$20+'Model Assumptions'!$C$20^2)"
            ))
            f_th = ws.cell(row=row, column=base + 4, value=(
                f"={y1}{row}*'Model Assumptions'!$C$22+{wavg}{row}*(1-'Model Assumptions'!$C$22)"
            ))
            f_lb = ws.cell(row=row, column=base + 5, value=(
                f"=INDEX(${sec2_col}${season_rows[HISTORICAL_YEARS[0]]}:"
                f"${sec2_col}${season_rows[HISTORICAL_YEARS[-1]]},"
                f"MATCH('Model Assumptions'!$C$18-1,$A${season_rows[HISTORICAL_YEARS[0]]}:"
                f"$A${season_rows[HISTORICAL_YEARS[-1]]},0))"
            ))
            f_pb = ws.cell(row=row, column=base + 6, value=(
                f"={th}{row}*'Model Assumptions'!$C$21+{lb}{row}*(1-'Model Assumptions'!$C$21)"
            ))
            for cell in (f_y1, f_y2, f_y3, f_wavg, f_th, f_lb, f_pb):
                cell.font = FORMULA_FONT
                cell.number_format = m["fmt"]

    # ==== Section 4: League average/std-dev of Section 3's Projected Baselines =============
    sec4_title_row = sec3_last_row + 2
    sec4_header_row = sec4_title_row + 1
    avg_row, std_row = sec4_header_row + 1, sec4_header_row + 2

    _section_title(
        ws, sec4_title_row, 6, "Section 4 — League Average & Std. Dev. of the 3-Yr Baselines"
    )
    _header_row(ws, sec4_header_row, ["Stat", *[m["label"] for m in METRICS]], height=18)

    ws.cell(row=avg_row, column=1, value="League Average").font = FORMULA_FONT
    ws.cell(row=std_row, column=1, value="League Std. Dev.").font = FORMULA_FONT

    proj_baseline_col = {
        m["key"]: get_column_letter(metric_block_start_col[m["key"]] + 6) for m in METRICS
    }
    for j, m in enumerate(METRICS):
        pcol = proj_baseline_col[m["key"]]
        pb_range = f"{pcol}${sec3_first_row}:{pcol}${sec3_last_row}"
        a = ws.cell(row=avg_row, column=2 + j, value=f"=AVERAGE({pb_range})")
        s = ws.cell(row=std_row, column=2 + j, value=f"=STDEVP({pb_range})")
        for cell in (a, s):
            cell.font = FORMULA_FONT
            cell.number_format = m["fmt"]

    # ==== Section 5: Z-scores, weighted composite, points-scale Pass Defense Score =========
    sec5_title_row = std_row + 2
    sec5_header_row = sec5_title_row + 1
    sec5_first_row = sec5_header_row + 1
    sec5_last_row = sec5_first_row + n_teams - 1

    _section_title(
        ws, sec5_title_row, 8,
        "Section 5 — Z-Scores and Pass Defense Score. ALL FIVE metrics are ALLOWED rates "
        "(lower raw value = better defense) -- Z-scores are uniformly sign-flipped (league "
        "average minus raw), same convention efficiency.py's own epa_def/success_def "
        "already use. Baseline/points-per-SD reuse QB Index's own Model Assumptions cells "
        "C37/C38.",
    )
    _header_row(
        ws, sec5_header_row,
        ["Team", *[f"{m['label']}\nZ" for m in METRICS], "Weighted\nZ-Score Sum",
         "Pass Defense\nScore (Points)", "Years of\nReal History"],
    )

    avg_cell_ref = {
        m["key"]: f"${get_column_letter(2 + j)}${avg_row}" for j, m in enumerate(METRICS)
    }
    std_cell_ref = {
        m["key"]: f"${get_column_letter(2 + j)}${std_row}" for j, m in enumerate(METRICS)
    }
    weight_cells = {"epa_db": "$C$87", "psr": "$C$88", "comp": "$C$89", "nya": "$C$90",
                     "exp": "$C$91"}

    for i in range(n_teams):
        sec3_row = sec3_first_row + i
        row = sec5_first_row + i
        f = ws.cell(row=row, column=1, value=f"=A{sec3_row}")
        f.font = FORMULA_FONT

        z_cols = []
        for j, m in enumerate(METRICS):
            pcol = proj_baseline_col[m["key"]]
            # ALL five metrics are "allowed" rates -- sign-flipped uniformly (league avg
            # minus raw), so a higher Z always means a better pass defense.
            formula = f"=({avg_cell_ref[m['key']]}-{pcol}{sec3_row})/{std_cell_ref[m['key']]}"
            col = 2 + j
            cell = ws.cell(row=row, column=col, value=formula)
            cell.font = FORMULA_FONT
            cell.number_format = "0.00"
            z_cols.append(get_column_letter(col))

        weighted_terms = " + ".join(
            f"{z_cols[j]}{row}*'Model Assumptions'!{weight_cells[m['key']]}"
            for j, m in enumerate(METRICS)
        )
        wz_col = 2 + len(METRICS)
        wz = ws.cell(row=row, column=wz_col, value=f"={weighted_terms}")
        wz.font = FORMULA_FONT
        wz.number_format = "0.00"

        score = ws.cell(row=row, column=wz_col + 1, value=(
            f"='Model Assumptions'!$C$37+{get_column_letter(wz_col)}{row}*'Model Assumptions'!$C$38"
        ))
        score.font = FORMULA_FONT
        score.number_format = "0.0;(0.0)"

        yh_link = ws.cell(row=row, column=wz_col + 2, value=f"=B{sec3_row}")
        yh_link.font = FORMULA_FONT
        yh_link.number_format = "0"

    # ---- Closing note --------------------------------------------------------------------
    note_row = sec5_last_row + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
    note = ws.cell(row=note_row, column=1, value=(
        "claude_code_spec_defensive_matchup_engine.md, Version 1 of the post-backtest "
        "roadmap. NOT YET VALIDATED: there is no backtesting harness in this project yet "
        "to prove phase-specific pass-defense matchups actually improve predictions over "
        "the existing PPG-based Off/Def blend -- this tab is real, verified-correct "
        "plumbing (every metric traces to real nflverse pbp, the same decay-weighted/"
        "regressed/Z-scored chain every prior tab uses), NOT a proven improvement. Do not "
        "treat its output as more reliable than the existing prediction until a real "
        "backtest exists to check that claim. All five scored metrics are real, phase-"
        "SPLIT pass-defense-ALLOWED rates (efficiency.compute_team_season_matchup_metrics/"
        "compute_team_season_efficiency) -- deliberately NOT the combined pass+run epa_def/"
        "success_def efficiency.py already had, which would have put the identical "
        "non-split number in this tab and Run Defense Matchup, defeating the point of "
        "splitting by phase (confirmed with the user before building this). Sack Rate / "
        "Pressure Proxy (QB Hit Rate) / Blitz Rate (Section 1, H/I/J) are REFERENCED live "
        "from Front Seven Index's own real Section 1 data by a SUMPRODUCT-array MATCH on "
        "(Team, Season) -- not recomputed, not part of this tab's own weighted composite. "
        "'Pressure Proxy' is QB Hit Rate specifically, not a distinct pressure metric -- "
        "Front Seven Index doesn't have one beyond Sack Rate/QB Hit Rate, and this is "
        "labeled as a proxy rather than presented as something more precise. NO man/zone "
        "coverage tendency column exists anywhere on this tab -- confirmed proprietary "
        "(paid FTN/Fantasy Points charting only), not approximated with a fabricated "
        "proxy. This tab has NO Team Ratings wiring -- Part C references its Section 5 "
        "Score directly from Week 1 Matchups instead, alongside (not replacing) the "
        "existing PPG-based prediction, per the spec's own instruction."
    ))
    note.font = NOTE_FONT
    note.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(workbook_path)
    print(
        f"Built '{SHEET_NAME}': Section 1 {len(season_stats)} rows, Section 3/5 {n_teams} "
        f"teams. NOT wired into Team Ratings -- Part C references Section 5 directly from "
        f"Week 1 Matchups."
    )
    print(f"Saved to {workbook_path}")
    return {
        "sec1_rows": len(season_stats), "n_teams": n_teams,
        "sec3_range": (sec3_first_row, sec3_last_row),
        "sec5_range": (sec5_first_row, sec5_last_row),
    }


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print('Usage: uv run python scripts/build_pass_defense_matchup.py "path/to/workbook.xlsx"')
        sys.exit(1)
    build(sys.argv[1])
