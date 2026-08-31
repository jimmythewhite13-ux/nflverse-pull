"""
Builds "Coaching Index" -- claude_code_spec_coaching_index.md.

CRITICAL DESIGN CHOICE (per the spec's own explicit instruction, see coaching_stats.py's own
module docstring for the full real-data verification): blends by COACH TENURE, not team
history. Section 1 is keyed by (Coach, Season, Team) -- each real coach-season's metrics
come from ONLY the real games that coach actually coached (verified live: the 3 real 2023
in-season interim changes, CAR/LAC/LV, split correctly at real per-game precision). Section
3's Y-1/Y-2/Y-3 blend is a SUMIFS on the CURRENT (2026) coach's own name + season offset --
the exact same identity-plus-season-offset pattern every other tab's Section 3 already uses,
just swapping Team for Coach. A coach with ZERO real seasons anywhere in the pulled window
(a true rookie HC) gets every Y-k year substituted with that season's real league average --
the EXACT SAME mechanism QB Index's own Rookie Baseline already uses (no new fallback
mechanism invented here), which is what the spec calls "League-Average New HC Baseline."

MANDATORY CORRELATION CHECK (per the spec): Section 5 includes a LIVE `=CORREL(...)` formula
comparing this tab's own raw composite score against Team Ratings' existing Net Power Rating
(N) -- not a one-off Python snapshot, so it stays honest as the underlying data changes.
CAVEAT stated on the tab itself: this correlation necessarily includes Coaching Index's own
contribution to N (Coaching Adjustment becomes one of N's summed terms once this tab is
wired in), the same self-inclusion effect every other Team-Ratings-wired tab's own composite
already has relative to N -- a fully independent check would need N computed WITHOUT this
tab's own term, which isn't isolated here. Treat the reported correlation as a conservative
(if anything, slightly inflated) estimate of shared signal, not a deflated one.

METRICS, with the spec's own honest confidence levels carried onto the tab:
  - 4th Down Aggressiveness (highest confidence) -- see coaching_stats.py's own HONESTY NOTE:
    a real Go-For-It-Rate-vs-league-average proxy in the 4th-and-2-or-less bucket, NOT a true
    win-probability-optimal decision-model comparison (no such model exists in nfl_data_py).
  - 1Q Net EPA/Play (real, buildable).
  - Penalty Discipline -- needs the SAME regression-to-mean shape Turnover & Red-Zone
    Regression already established (raw rate regressed toward that season's real league
    average, weighted by real sample size via a DEDICATED blend-weight -- not reused BY
    REFERENCE from that other tab, since the sample-size scale differs entirely: real total
    plays [~400-2400] vs. real red-zone drives [~33-76] -- but built to the identical shape).
  - 2H EPA Delta -- LOWEST confidence, explicit uncertainty note, visibly lower default
    weight than 4th Down Aggressiveness (per the spec's own explicit instruction).
  - Red Zone Conversion is deliberately NOT rebuilt here -- see Turnover & Red-Zone
    Regression instead.

NO Section 7 Manual Override on this tab -- unlike a QB/RB depth-chart competition, a team's
real head coach is unambiguous, single-sourced directly from schedules (0 nulls, verified
live), so there's no real uncertainty an override would resolve.

Usage:
    uv run python scripts/build_coaching_index.py "C:\\path\\to\\NFL_Prediction_Model.xlsx"
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from nflverse_pull.availability import fetch_schedules_with_dates  # noqa: E402
from nflverse_pull.coaching_stats import (  # noqa: E402
    compute_coach_season_stats,
    compute_current_coach_by_team,
    compute_game_coach_map,
)
from nflverse_pull.efficiency import fetch_pbp  # noqa: E402

HISTORICAL_YEARS = [2023, 2024, 2025]
CURRENT_SEASON = 2026
SHEET_NAME = "Coaching Index"
INSERT_AFTER_SHEET = "Team-Specific HFA"

TEAM_ORDER = [
    "Buffalo Bills", "Miami Dolphins", "New England Patriots", "New York Jets",
    "Baltimore Ravens", "Cincinnati Bengals", "Cleveland Browns", "Pittsburgh Steelers",
    "Houston Texans", "Indianapolis Colts", "Jacksonville Jaguars", "Tennessee Titans",
    "Denver Broncos", "Kansas City Chiefs", "Las Vegas Raiders", "Los Angeles Chargers",
    "Dallas Cowboys", "New York Giants", "Philadelphia Eagles", "Washington Commanders",
    "Chicago Bears", "Detroit Lions", "Green Bay Packers", "Minnesota Vikings",
    "Atlanta Falcons", "Carolina Panthers", "New Orleans Saints", "Tampa Bay Buccaneers",
    "Arizona Cardinals", "Los Angeles Rams", "San Francisco 49ers", "Seattle Seahawks",
]

# METRICS: key (Section 1 raw column) | label | invert (True = lower raw value is better,
# Z-flipped) | weight cell (Model Assumptions).
METRICS = [
    {"key": "go_rate", "label": "4th Down\nGo Rate (Short)", "invert": False,
     "weight_cell": "C162"},
    {"key": "q1_epa", "label": "1Q Net\nEPA/Play", "invert": False, "weight_cell": "C163"},
    {"key": "penalty", "label": "Regressed\nPenalty Rate", "invert": True,
     "weight_cell": "C164"},
    {"key": "h2_delta", "label": "2H EPA\nDelta", "invert": False, "weight_cell": "C165"},
]

TITLE_FONT = Font(name="Arial", size=10, bold=True)
TITLE_FILL = PatternFill("solid", fgColor="FFD9E1F2")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_ALIGN = Alignment(wrap_text=True, horizontal="center", vertical="center")
INPUT_FONT = Font(name="Arial", size=10, color="FF0000FF")
FORMULA_FONT = Font(name="Arial", size=10, color="FF000000")
LINK_FONT = Font(name="Arial", size=10, color="FF008000")
NOTE_FONT = Font(name="Arial", size=9, color="FF808080")
ASSUMPTION_FILL = PatternFill("solid", fgColor="FFFFFF00")
WARN_FILL = PatternFill("solid", fgColor="FFFFC7CE")
WARN_FONT = Font(name="Arial", size=10, bold=True, color="FF9C0006")


def _section_title(ws, row: int, last_col: int, text: str) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL


def _header_row(ws, row: int, headers: list[str], height: float = 28) -> None:
    ws.row_dimensions[row].height = height
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN


def add_model_assumptions_weights(wb: openpyxl.Workbook) -> None:
    ws = wb["Model Assumptions"]
    title_row = 161
    ws.merge_cells(f"A{title_row}:D{title_row}")
    title = ws.cell(row=title_row, column=1, value=(
        "Coaching Index (pts per std. dev., + Penalty Regression Blend Weight; see "
        "'Coaching Index' tab)"
    ))
    title.font = HEADER_FONT
    title.fill = HEADER_FILL

    rows = [
        (162, "4th Down Aggressiveness Weight (pts per SD)", 0.4,
         "Highest confidence per the spec -- a real Go-For-It-Rate-vs-league-average proxy "
         "in the 4th-and-2-or-less bucket (see 'Coaching Index' tab's own honesty note: "
         "NOT a true win-probability-optimal decision-model comparison -- no such model "
         "exists in the real pbp data this project pulls from)."),
        (163, "1Q Scripting Net EPA Weight (pts per SD)", 0.4,
         "Real, buildable -- own offense's real Q1 EPA/play minus own defense's real Q1 "
         "EPA/play allowed."),
        (164, "Penalty Discipline Weight (pts per SD)", 0.3,
         "Uses the REGRESSED Penalty Rate (already shrunk toward league average by real "
         "sample size, C166-C168 below), not the raw rate."),
        (165, "Second-Half EPA Delta Weight (pts per SD)", 0.15,
         "LOWEST CONFIDENCE, deliberately weighted well below 4th Down Aggressiveness "
         "(C162) -- 'halftime adjustment skill' is a popular football-media narrative with "
         "real, published skepticism in analytics circles about whether it's a persistent, "
         "measurable coaching trait versus noise that gets a story fitted to it after the "
         "fact. Set to 0 to exclude this dimension entirely if you don't trust it at all."),
        (166, "Penalty Regression Blend Weight Base", 0.10,
         "Real minimum trust (lowest real sample size in the pulled data, an interim "
         "coach's partial season) before any sample-size credit is added."),
        (167, "Penalty Regression Blend Weight Increment (per real play)", 0.0004,
         "Calibrated live against the real pulled sample: Total Plays per real coach-season "
         "ranges ~430 (a real partial-season interim coach) to ~2435 (a full real season) "
         "-- a ~2,000-play spread, a completely different scale than Turnover & Red-Zone "
         "Regression's own red-zone-drive sample (33-76). This increment carries a full-"
         "season coach to the cap (C168) while still meaningfully differentiating a "
         "partial-season interim coach's smaller sample -- verified live before finalizing "
         "this value (0.03-per-drive from the Turnover engine would have saturated "
         "instantly here at this much larger scale, the same mistake caught and fixed "
         "there)."),
        (168, "Penalty Regression Blend Weight Cap", 0.90,
         "Same cap convention as Turnover & Red-Zone Regression's own blend weight -- never "
         "fully trust a single season's raw rate, even at a full sample."),
        (169, "Coaching Composite-to-Points Conversion", 0.5,
         "A starting guess, like every other coefficient in this model. Converts the "
         "Weighted Z-Score Sum (Section 5) into real game points -- the spec suggests a "
         "roughly -1.5 to +1.5 range as a sanity check, not a hardcoded bound; this stays "
         "a tunable assumption."),
    ]
    for row, label, value, note in rows:
        ws.cell(row=row, column=2, value=label)
        c = ws.cell(row=row, column=3, value=value)
        c.font = INPUT_FONT
        c.fill = ASSUMPTION_FILL
        c.number_format = "0.0000" if row == 167 else "0.00"
        n = ws.cell(row=row, column=4, value=note)
        n.font = NOTE_FONT
        n.alignment = Alignment(wrap_text=True, vertical="top")


def _pull_data():
    print(f"Pulling {HISTORICAL_YEARS + [CURRENT_SEASON]} schedule + pbp data for coaching...")
    sched = fetch_schedules_with_dates(HISTORICAL_YEARS + [CURRENT_SEASON])
    game_coach_map = compute_game_coach_map(sched)
    pbp = fetch_pbp(HISTORICAL_YEARS)
    coach_season = compute_coach_season_stats(pbp, game_coach_map)
    current_coach = compute_current_coach_by_team(sched, CURRENT_SEASON)
    return coach_season, current_coach


def build(workbook_path: str) -> dict:
    coach_season, current_coach = _pull_data()

    wb = openpyxl.load_workbook(workbook_path)
    add_model_assumptions_weights(wb)

    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    insert_after = INSERT_AFTER_SHEET if INSERT_AFTER_SHEET in wb.sheetnames else wb.sheetnames[0]
    ws = wb.create_sheet(SHEET_NAME, index=wb.sheetnames.index(insert_after) + 1)
    ws.column_dimensions["A"].width = 20.0

    ws.merge_cells("A1:K1")
    t = ws.cell(row=1, column=1, value=(
        "Coaching Index -- Real Coach-Tenure-Aware 3-Yr Decay-Weighted Composite (4th Down "
        "Aggressiveness, 1Q Scripting EPA, Regressed Penalty Discipline, 2H EPA Delta). "
        "Blended by COACH, not team history -- see this tab's own Section 3 and "
        "coaching_stats.py's module docstring for the real 2023 in-season interim-coach "
        "verification (CAR/LAC/LV)."
    ))
    t.font = Font(name="Arial", size=12, bold=True)

    # ==== Section 1: Raw per-Coach-Season data =============================================
    sec1_first = 5
    sec1_last = sec1_first + len(coach_season) - 1
    _section_title(
        ws, 3, 11,
        "Section 1 — Raw per-Coach-Season Data (each row = one real coach's real games at "
        "one real team in one real season -- a mid-season coaching change produces TWO "
        "rows, each built only from the games that coach actually coached, verified live "
        "against the real 2023 CAR/LAC/LV interim changes). Penalty Blend Weight/Regressed "
        "Penalty Rate apply the SAME regression-to-mean shape Turnover & Red-Zone "
        "Regression already uses, with its OWN dedicated constants (C166-C168) -- the "
        "real sample-size scale here (total plays) is nothing like that tab's real "
        "red-zone-drive sample.",
    )
    _header_row(ws, 4, [
        "Coach", "Season", "Team", "4th Down Go\nRate (Short)", "4th-and-Short\nAttempts",
        "1Q Net\nEPA/Play", "Penalty Rate\n(Raw)", "Total Plays\n(real)",
        "Penalty Blend\nWeight", "Regressed\nPenalty Rate", "2H EPA\nDelta",
    ])
    for i, r in enumerate(coach_season.to_dict("records")):
        row = sec1_first + i
        values = [
            r["Coach"], int(r["Season"]), r["Team"], float(r["4th Down Go Rate (Short)"]),
            int(r["4th-and-Short Attempts"]), float(r["1Q Net EPA/Play"]),
            float(r["Penalty Rate"]), int(r["Total Plays"]),
        ]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.font = INPUT_FONT
            if col in (4, 6, 7):
                cell.number_format = "0.000"
        # Penalty Blend Weight -- SUMIFS-free simple lookup: this row's own real sample size
        # against the dedicated constants above, capped.
        bw = ws.cell(row=row, column=9, value=(
            f"=MIN('Model Assumptions'!$C$168,'Model Assumptions'!$C$166+"
            f"'Model Assumptions'!$C$167*H{row})"
        ))
        bw.font = FORMULA_FONT
        bw.number_format = "0.000"
        # Regressed Penalty Rate = Raw * BlendWeight + (that season's real league-average
        # raw rate, Section 2) * (1 - BlendWeight). Section 2 is below in row order but
        # Excel forward-references cells with no issue.
        reg = ws.cell(row=row, column=10, value=(
            f"=G{row}*I{row}+IFERROR(INDEX($D${{s2first}}:$D${{s2last}},"
            f"MATCH(B{row},$A${{s2first}}:$A${{s2last}},0)),G{row})*(1-I{row})"
        ))
        reg.font = FORMULA_FONT
        reg.number_format = "0.000"
        hd = ws.cell(row=row, column=11, value=float(r["2H EPA Delta"]))
        hd.font = INPUT_FONT
        hd.number_format = "0.000"

    id_range = f"$A${sec1_first}:$A${sec1_last}"
    season_range = f"$B${sec1_first}:$B${sec1_last}"
    go_range = f"$D${sec1_first}:$D${sec1_last}"
    q1_range = f"$F${sec1_first}:$F${sec1_last}"
    penalty_raw_range = f"$G${sec1_first}:$G${sec1_last}"
    penalty_reg_range = f"$J${sec1_first}:$J${sec1_last}"
    h2_range = f"$K${sec1_first}:$K${sec1_last}"

    # ==== Section 2: League Average per Season =============================================
    sec2_title_row = sec1_last + 2
    sec2_header_row = sec2_title_row + 1
    season_rows = {yr: sec2_header_row + 1 + i for i, yr in enumerate(HISTORICAL_YEARS)}
    sec2_first, sec2_last = season_rows[HISTORICAL_YEARS[0]], season_rows[HISTORICAL_YEARS[-1]]

    _section_title(
        ws, sec2_title_row, 5,
        "Section 2 — League Average per Season (simple average across every real "
        "coach-season that year; column D, Avg Penalty Rate, uses the RAW rate -- what "
        "Section 1's own Regressed Penalty Rate blends toward -- not a regressed figure)",
    )
    _header_row(ws, sec2_header_row, [
        "Season", "Avg 4th Down\nGo Rate (Short)", "Avg 1Q Net\nEPA/Play",
        "Avg Penalty\nRate (Raw)", "Avg 2H EPA\nDelta",
    ], height=20)
    for yr in HISTORICAL_YEARS:
        row = season_rows[yr]
        ws.cell(row=row, column=1, value=yr).font = INPUT_FONT
        for col, rng in ((2, go_range), (3, q1_range), (4, penalty_raw_range), (5, h2_range)):
            f = ws.cell(row=row, column=col, value=f"=AVERAGEIF({season_range},{yr},{rng})")
            f.font = FORMULA_FONT
            f.number_format = "0.000"

    # Backfill Section 1's Regressed Penalty Rate formula with the real Section 2 row range.
    for i in range(len(coach_season)):
        row = sec1_first + i
        ws.cell(row=row, column=10).value = ws.cell(row=row, column=10).value.replace(
            "{s2first}", str(sec2_first)
        ).replace("{s2last}", str(sec2_last))

    league_avg_col = {"go_rate": "B", "q1_epa": "C", "penalty": "D", "h2_delta": "E"}

    # ==== Section 3: Per-current-team Coach-Tenure-Aware 3-Yr Decay-Weighted Baseline ======
    sec3_title_row = sec2_last + 2
    sec3_header_row = sec3_title_row + 1
    sec3_first_row = sec3_header_row + 1
    n_teams = len(TEAM_ORDER)
    sec3_last_row = sec3_first_row + n_teams - 1
    sec3_last_col = 5 + len(METRICS) * 7

    _section_title(
        ws, sec3_title_row, sec3_last_col,
        "Section 3 — Per-Team, Current-Coach 3-Yr Decay-Weighted, Regressed Baseline. "
        "Blended by COACH (Y-1/Y-2/Y-3 = SUMIFS on the CURRENT coach's own name + season "
        "offset, across ANY team he coached) -- NOT by team history. A coach missing a "
        "given year (never coached anywhere that season, including a true rookie HC with "
        "0 years anywhere) substitutes that season's real Section 2 league average -- the "
        "EXACT SAME mechanism QB Index's own Rookie Baseline uses, which is what the spec "
        "calls the League-Average New HC Baseline; no separate flat constant needed.",
    )
    headers3 = ["Team", "Current Coach\n(2026)", "Years as HC\n(real, any team)", "", ""]
    metric_block_start_col: dict[str, int] = {}
    col_cursor = 6
    for m in METRICS:
        metric_block_start_col[m["key"]] = col_cursor
        headers3 += [
            f"{m['label']} Y-1", f"{m['label']} Y-2", f"{m['label']} Y-3",
            f"Weighted\n{m['label']} Avg\n(3-Yr decay)", f"Coach History\n{m['label']}",
            f"League Baseline\n{m['label']} (Y-1)", f"Projected 3-Yr\n{m['label']} Baseline",
        ]
        col_cursor += 7
    _header_row(ws, sec3_header_row, headers3[:sec3_last_col])

    metric_source_range = {
        "go_rate": go_range, "q1_epa": q1_range, "penalty": penalty_reg_range,
        "h2_delta": h2_range,
    }

    for i, team in enumerate(TEAM_ORDER):
        row = sec3_first_row + i
        coach_row = current_coach[current_coach["Team"] == team]
        coach_name = coach_row.iloc[0]["Coach"] if len(coach_row) else ""
        ws.cell(row=row, column=1, value=team).font = FORMULA_FONT
        cc = ws.cell(row=row, column=2, value=coach_name)
        cc.font = INPUT_FONT

        years_terms = "+".join(
            f"--(COUNTIFS({id_range},$B{row},{season_range},'Model Assumptions'!$C$18-{k})>0)"
            for k in (1, 2, 3)
        )
        yh = ws.cell(row=row, column=3, value=f"={years_terms}")
        yh.font = FORMULA_FONT
        yh.number_format = "0"

        for m in METRICS:
            base = metric_block_start_col[m["key"]]
            y1, y2, y3, wavg, th, lb, pb = (get_column_letter(base + k) for k in range(7))
            mrange = metric_source_range[m["key"]]
            avg_col = league_avg_col[m["key"]]

            def _ysub(offset: int, row=row, mrange=mrange, avg_col=avg_col) -> str:
                season_avg_lookup = (
                    f"INDEX(${avg_col}${sec2_first}:${avg_col}${sec2_last},"
                    f"MATCH('Model Assumptions'!$C$18-{offset},$A${sec2_first}:"
                    f"$A${sec2_last},0))"
                )
                return (
                    f"=IF(COUNTIFS({id_range},$B{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset})=0,{season_avg_lookup},"
                    f"SUMIFS({mrange},{id_range},$B{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset}))"
                )

            f_y1 = ws.cell(row=row, column=base, value=_ysub(1))
            f_y2 = ws.cell(row=row, column=base + 1, value=_ysub(2))
            f_y3 = ws.cell(row=row, column=base + 2, value=_ysub(3))
            f_wavg = ws.cell(row=row, column=base + 3, value=(
                f"=({y1}{row}*1+{y2}{row}*'Model Assumptions'!$C$20+"
                f"{y3}{row}*('Model Assumptions'!$C$20^2))/"
                f"(1+'Model Assumptions'!$C$20+'Model Assumptions'!$C$20^2)"
            ))
            f_th = ws.cell(row=row, column=base + 4, value=(
                f"={y1}{row}*'Model Assumptions'!$C$22+{wavg}{row}*(1-'Model Assumptions'!$C$22)"
            ))
            f_lb = ws.cell(row=row, column=base + 5, value=(
                f"=INDEX(${avg_col}${sec2_first}:${avg_col}${sec2_last},"
                f"MATCH('Model Assumptions'!$C$18-1,$A${sec2_first}:$A${sec2_last},0))"
            ))
            f_pb = ws.cell(row=row, column=base + 6, value=(
                f"={th}{row}*'Model Assumptions'!$C$21+{lb}{row}*(1-'Model Assumptions'!$C$21)"
            ))
            for cell in (f_y1, f_y2, f_y3, f_wavg, f_th, f_lb, f_pb):
                cell.font = FORMULA_FONT
                cell.number_format = "0.000"

    # ==== Section 4: League Average & Std Dev of the 3-Yr Baseline =========================
    sec4_title_row = sec3_last_row + 2
    sec4_header_row = sec4_title_row + 1
    avg_row, std_row = sec4_header_row + 1, sec4_header_row + 2
    _section_title(
        ws, sec4_title_row, sec3_last_col,
        "Section 4 — League Average & Std. Dev. of the 3-Yr Baseline (per metric, for "
        "Z-scoring)",
    )
    ws.cell(row=sec4_header_row, column=1, value="Stat").font = HEADER_FONT
    ws.cell(row=avg_row, column=1, value="League Average").font = FORMULA_FONT
    ws.cell(row=std_row, column=1, value="League Std. Dev.").font = FORMULA_FONT
    pb_col_letter = {}
    for m in METRICS:
        base = metric_block_start_col[m["key"]]
        pb_col = get_column_letter(base + 6)
        pb_col_letter[m["key"]] = pb_col
        rng = f"{pb_col}${sec3_first_row}:{pb_col}${sec3_last_row}"
        a = ws.cell(row=avg_row, column=base + 6, value=f"=AVERAGE({rng})")
        s = ws.cell(row=std_row, column=base + 6, value=f"=STDEVP({rng})")
        a.font = FORMULA_FONT
        s.font = FORMULA_FONT
        a.number_format = "0.0000"
        s.number_format = "0.0000"

    # ==== Section 5: Z-Scores, Coaching Index Score, and the CORRELATION CHECK =============
    sec5_title_row = std_row + 2
    sec5_header_row = sec5_title_row + 1
    sec5_first_row = sec5_header_row + 1
    sec5_last_row = sec5_first_row + n_teams - 1
    _section_title(
        ws, sec5_title_row, 8,
        "Section 5 — Z-Scores, Coaching Index Score, and the MANDATORY correlation check "
        "against Team Ratings' existing Net Power Rating (see this tab's own module "
        "docstring for the self-inclusion caveat on that correlation).",
    )
    _header_row(ws, sec5_header_row, [
        "Team", "4th Down Z", "1Q EPA Z", "Penalty Z\n(flipped)", "2H EPA Z",
        "Weighted\nZ-Score Sum", "Coaching\nAdjustment (pts)", "Net Power\nRating (ref)",
    ])
    z_cols = {}
    for i, team in enumerate(TEAM_ORDER):
        row = sec5_first_row + i
        s3row = sec3_first_row + i
        ws.cell(row=row, column=1, value=f"=A{s3row}").font = FORMULA_FONT
        for j, m in enumerate(METRICS):
            base = metric_block_start_col[m["key"]]
            pb_col = pb_col_letter[m["key"]]
            zcol = 2 + j
            z_cols[m["key"]] = zcol
            sign = "-1*" if m["invert"] else ""
            f = ws.cell(row=row, column=zcol, value=(
                f"={sign}({pb_col}{s3row}-${pb_col}${avg_row})/${pb_col}${std_row}"
            ))
            f.font = FORMULA_FONT
            f.number_format = "0.00"
        wz_formula = "=" + "+".join(
            f"{get_column_letter(2 + j)}{row}*'Model Assumptions'!${m['weight_cell'][0]}"
            f"${m['weight_cell'][1:]}"
            for j, m in enumerate(METRICS)
        )
        wz = ws.cell(row=row, column=6, value=wz_formula)
        wz.font = FORMULA_FONT
        wz.number_format = "0.00"
        adj = ws.cell(row=row, column=7, value=f"=F{row}*'Model Assumptions'!$C$169")
        adj.font = FORMULA_FONT
        adj.number_format = "0.00;(0.00)"
        npr = ws.cell(row=row, column=8, value=(
            f"=IFERROR(INDEX('Team Ratings'!$N$3:$N$34,"
            f"MATCH(A{row},'Team Ratings'!$A$3:$A$34,0)),\"\")"
        ))
        npr.font = LINK_FONT
        npr.number_format = "0.00;(0.00)"

    wz_range = f"$F${sec5_first_row}:$F${sec5_last_row}"
    npr_range = f"$H${sec5_first_row}:$H${sec5_last_row}"
    corr_row = sec5_last_row + 2
    corr_label = ws.cell(
        row=corr_row, column=1,
        value="Correlation (Weighted Z-Score Sum vs. Net Power Rating)",
    )
    corr_label.font = FORMULA_FONT
    corr_cell = ws.cell(row=corr_row, column=2, value=f"=CORREL({wz_range},{npr_range})")
    corr_cell.font = FORMULA_FONT
    corr_cell.number_format = "0.000"
    flag_cell = ws.cell(row=corr_row, column=3, value=(
        f'=IF(ABS(B{corr_row})>0.6,"INDEPENDENCE CONCERN -- correlation exceeds 0.6, this '
        f'index may not be adding signal beyond existing Net Power Rating inputs",'
        f'"No independence concern at the 0.6 threshold")'
    ))
    flag_cell.font = WARN_FONT
    ws.merge_cells(start_row=corr_row, start_column=3, end_row=corr_row, end_column=8)

    # ==== Section 6: Coaching Adjustment wired into Team Ratings ===========================
    tr = wb["Team Ratings"]
    tr.cell(row=2, column=31, value="Coaching\nAdjustment (pts)").font = HEADER_FONT
    tr.cell(row=2, column=31).fill = HEADER_FILL
    tr.cell(row=2, column=31).alignment = HEADER_ALIGN
    for i, team in enumerate(TEAM_ORDER):
        tr_row = 3 + i
        c = tr.cell(row=tr_row, column=31, value=(
            f"=IFERROR(INDEX('{SHEET_NAME}'!$G${sec5_first_row}:$G${sec5_last_row},"
            f"MATCH(A{tr_row},'{SHEET_NAME}'!$A${sec5_first_row}:$A${sec5_last_row},0)),0)"
        ))
        c.font = LINK_FONT
        c.number_format = "0.00;(0.00)"
        net = tr.cell(row=tr_row, column=14, value=(
            f"=J{tr_row}-K{tr_row}+L{tr_row}+M{tr_row}+P{tr_row}+R{tr_row}+S{tr_row}+"
            f"T{tr_row}+U{tr_row}+V{tr_row}+W{tr_row}+X{tr_row}+Y{tr_row}+Z{tr_row}+"
            f"AA{tr_row}+AB{tr_row}+AC{tr_row}+AD{tr_row}+AE{tr_row}"
        ))
        net.font = FORMULA_FONT
        net.number_format = "0.00;(0.00)"

    # ---- Closing note ----------------------------------------------------------------------
    note_row = corr_row + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)
    note = ws.cell(row=note_row, column=1, value=(
        "claude_code_spec_coaching_index.md. Blended by COACH TENURE, not team history -- "
        "see Section 3's own header note and coaching_stats.py's module docstring for the "
        "real 2023 CAR/LAC/LV in-season interim-coach verification. 4th Down Aggressiveness "
        "(highest confidence per the spec) is a real Go-For-It-Rate-vs-league-average proxy "
        "on 4th-and-2-or-less -- NOT a true win-probability-optimal decision-model "
        "comparison (no such model exists in the real pbp data this project pulls from; "
        "building one from scratch would fabricate rigor this project doesn't have). "
        "Second-Half EPA Delta carries the LOWEST default weight (Model Assumptions C165, "
        "well below 4th Down's C162) and this explicit uncertainty note: 'halftime "
        "adjustment skill' has real, published skepticism in analytics circles about "
        "whether it's a persistent coaching trait versus noise. Penalty Discipline uses a "
        "REGRESSED rate (Section 1 col J), shrunk toward that season's real league average "
        "by real sample size via a DEDICATED blend-weight (C166-C168, calibrated live "
        "against this tab's own real ~430-2435 total-play sample range -- NOT reused by "
        "reference from Turnover & Red-Zone Regression's own red-zone-drive-scaled "
        "constants, though built to the identical shape the spec calls for). Red Zone "
        "Conversion is deliberately NOT rebuilt here -- see Turnover & Red-Zone Regression. "
        "The correlation check (row above) is a LIVE formula, re-evaluating automatically "
        "as inputs change -- see this tab's own module docstring for why it necessarily "
        "includes this tab's own contribution to Net Power Rating. NO Manual Override table "
        "on this tab -- a team's real head coach is unambiguous, single-sourced directly "
        "from real schedule data (0 nulls, verified live), unlike a depth-chart "
        "competition. NOT YET VALIDATED against real outcomes -- same disclaimer as every "
        "tab in the Defensive Matchup Engine family."
    ))
    note.font = NOTE_FONT
    note.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(workbook_path)
    print(
        f"Built '{SHEET_NAME}': Section 1 {len(coach_season)} coach-season rows, Section "
        f"3/5 {n_teams} teams. Wired into 'Team Ratings' (col AE) and took over ownership "
        "of the Net Power Rating (N) formula."
    )
    print(f"Saved to {workbook_path}")
    return {"sec1_rows": len(coach_season), "sec5_range": (sec5_first_row, sec5_last_row)}


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print('Usage: uv run python scripts/build_coaching_index.py "path/to/workbook.xlsx"')
        sys.exit(1)
    build(sys.argv[1])
