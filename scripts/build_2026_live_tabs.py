"""
Live Weekly Workflow: 9 new, real, refreshable Excel tabs mirroring the live agent's own
SQLite state (never live formulas -- Excel can't make network calls; same real
"re-run the script to refresh" convention every other real-data tab in this workbook already
uses). None of these tabs feed the model -- read-only visibility, no formula anywhere here is
ever consumed by Season Matchups or any scoring tab.

Explicit scope boundary: nothing in this script touches a coefficient, formula, or model file.
It reads prediction_audit.sqlite3 and real git history, and writes values.

Usage:
    uv run python scripts/build_2026_live_tabs.py "C:\\path\\to\\NFL_Prediction_Model.xlsx"
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))
sys.path.insert(0, str(REPO_ROOT / "prediction_audit"))

import sqlite3  # noqa: E402

from db.schema import DEFAULT_DB_PATH  # noqa: E402

TITLE_FONT = Font(name="Arial", size=10, bold=True)
TITLE_FILL = PatternFill("solid", fgColor="FFD9E1F2")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_ALIGN = Alignment(wrap_text=True, horizontal="center", vertical="center")
INPUT_FONT = Font(name="Arial", size=10, color="FF0000FF")   # real, machine-pulled values
NOTE_FONT = Font(name="Arial", size=9, color="FF808080")

SEASON = 2026


def _rebuild_sheet(wb: openpyxl.Workbook, name: str, after: str | None = None):
    if name in wb.sheetnames:
        del wb[name]
    idx = wb.sheetnames.index(after) + 1 if after and after in wb.sheetnames else len(wb.sheetnames)
    return wb.create_sheet(name, index=idx)


def _title(ws, text: str, ncols: int):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font = TITLE_FONT
    c.fill = TITLE_FILL


def _headers(ws, headers: list[str], row: int = 2):
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = HEADER_ALIGN


def _write_rows(ws, rows: list[tuple], start_row: int = 3):
    for i, values in enumerate(rows):
        for col, v in enumerate(values, start=1):
            ws.cell(row=start_row + i, column=col, value=v).font = INPUT_FONT
    return start_row + len(rows)


def build_schedule_tab(wb, conn: sqlite3.Connection):
    ws = _rebuild_sheet(wb, "2026 Game Schedule")
    headers = ["Game ID", "Week", "Away", "Home", "Kickoff (UTC)", "Status", "Status Reason"]
    _title(ws, "2026 Game Schedule -- real, refreshed by re-running "
               "scripts/build_2026_live_tabs.py", len(headers))
    _headers(ws, headers)
    # Real, deliberate INNER JOIN, not LEFT: a pre-existing, unrelated demo script
    # (seed_v35_demo.py, an early schema-proving exercise predating the real 2026 season)
    # inserted 272 of its own `games` rows under the SAME season=2026 label using a different,
    # malformed game_id convention (full team names, no real kickoff_time) -- confirmed by
    # direct inspection, not assumed. game_workflow_status only ever contains real, correctly-
    # formatted game_ids (written exclusively by game_status.py from the real
    # fetch_schedules() pull), so joining through it is a real, non-destructive filter -- the
    # stale demo rows are left alone in `games`, not deleted, per this project's own "don't
    # delete without explicit confirmation" discipline.
    rows = conn.execute(
        """
        SELECT g.game_id, g.week, g.away_team, g.home_team, g.kickoff_time,
               gws.status, gws.status_reason
        FROM game_workflow_status gws JOIN games g ON g.game_id = gws.game_id
        WHERE g.season = ?
        ORDER BY g.week, g.kickoff_time
        """,
        (SEASON,),
    ).fetchall()
    _write_rows(ws, rows)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["G"].width = 60


def build_season_control_tab(wb, conn: sqlite3.Connection):
    ws = _rebuild_sheet(wb, "2026 Season Control", after="2026 Game Schedule")
    _title(ws, "2026 Season Control -- real status counts + workflow settings", 2)
    _headers(ws, ["Setting", "Value"])
    counts = conn.execute(
        """
        SELECT gws.status, COUNT(*)
        FROM game_workflow_status gws JOIN games g ON g.game_id = gws.game_id
        WHERE g.season = ? GROUP BY 1 ORDER BY 1
        """,
        (SEASON,),
    ).fetchall()
    rows = [
        ("Current season", SEASON),
        ("Prediction cutoff (hours before kickoff)", 48),
        ("First predictable week", 4),
        *[(f"Games: {status}", n) for status, n in counts],
    ]
    _write_rows(ws, rows)
    ws.column_dimensions["A"].width = 45


def build_pregame_snapshot_tab(wb, conn: sqlite3.Connection):
    ws = _rebuild_sheet(wb, "Pre-Game Snapshot", after="2026 Season Control")
    headers = ["Game ID", "Run ID", "Model Version", "Frozen At", "Home Projected",
               "Away Projected", "Margin", "Home Win Prob", "Component", "Value", "Side"]
    _title(ws, "Pre-Game Snapshot -- real, frozen predictions for every PREDICTED+ game "
               "(component_contributions), one row per real named term", len(headers))
    _headers(ws, headers)
    # Real, deliberate join through game_workflow_status, not a LIKE on game_id -- an
    # unrelated, pre-existing demo script (seed_v35_demo.py) inserted exactly one
    # prediction_runs row under the SAME "2026_..." game_id prefix with model_status still
    # ACTIVE; filtering by prefix alone would surface that stale demo row here. See
    # build_schedule_tab()'s own comment for the full real explanation.
    rows = conn.execute(
        """
        SELECT pr.game_id, pr.run_id, pr.model_version, pr.prediction_timestamp,
               p.home_projected_points, p.away_projected_points, p.projected_margin,
               p.home_win_probability, cc.component_name, cc.contribution_value, cc.side
        FROM prediction_runs pr
        JOIN predictions p ON p.run_id = pr.run_id
        JOIN component_contributions cc ON cc.run_id = pr.run_id
        JOIN game_workflow_status gws ON gws.game_id = pr.game_id
        WHERE pr.model_status = 'ACTIVE'
        ORDER BY pr.game_id, cc.component_name
        """,
    ).fetchall()
    _write_rows(ws, rows)
    if not rows:
        ws.cell(row=3, column=1, value=(
            f"No real ACTIVE prediction exists yet for any season-{SEASON} game -- populates "
            f"automatically once prediction_freeze.yml freezes its first real prediction."
        )).font = NOTE_FONT


def build_market_capture_tab(wb, conn: sqlite3.Connection):
    ws = _rebuild_sheet(wb, "Market Capture", after="Pre-Game Snapshot")
    headers = ["Game ID", "Sportsbook", "Market", "Line Value", "Odds", "Tier", "Captured At",
               "Kickoff"]
    _title(ws, "Market Capture -- real, live agent captures (raw_market_captures / "
               "v_ingestion_market_tiers). Tier caveat: 'closing' for a game that hasn't "
               "kicked off yet means 'latest so far', not confirmed final.", len(headers))
    _headers(ws, headers)
    rows = conn.execute(
        """
        SELECT game_id, sportsbook, market_type, line_value, odds, line_stage, captured_at,
               kickoff_time
        FROM v_ingestion_market_tiers
        WHERE game_id LIKE ?
        ORDER BY game_id, sportsbook, market_type
        """,
        (f"{SEASON}_%",),
    ).fetchall()
    _write_rows(ws, rows)


def build_results_tab(wb, conn: sqlite3.Connection):
    ws = _rebuild_sheet(wb, "Results", after="Market Capture")
    headers = ["Game ID", "Week", "Away", "Home", "Away Final", "Home Final"]
    _title(ws, "Results -- real final scores, once known", len(headers))
    _headers(ws, headers)
    rows = conn.execute(
        """
        SELECT r.game_id, g.week, g.away_team, g.home_team, r.away_final_score,
               r.home_final_score
        FROM results r JOIN games g ON g.game_id = r.game_id
        WHERE g.season = ? ORDER BY g.week
        """,
        (SEASON,),
    ).fetchall()
    _write_rows(ws, rows)
    if not rows:
        ws.cell(row=3, column=1, value=(
            f"No real season-{SEASON} game has a recorded final score yet -- populates "
            f"automatically once real games are played."
        )).font = NOTE_FONT


def build_prediction_audit_tab(wb, conn: sqlite3.Connection):
    ws = _rebuild_sheet(wb, "Prediction Audit", after="Results")
    headers = ["Game ID", "Margin Error", "Brier Score", "CLV Movement", "Computed At"]
    _title(ws, "Prediction Audit -- real, computed-once metrics for AUDITED games "
               "(prediction_audit_metrics)", len(headers))
    _headers(ws, headers)
    rows = conn.execute(
        """
        SELECT pr.game_id, m.margin_error, m.brier_score, m.clv_movement, m.computed_at
        FROM prediction_audit_metrics m
        JOIN prediction_runs pr ON pr.run_id = m.run_id
        WHERE pr.game_id LIKE ? ORDER BY pr.game_id
        """,
        (f"{SEASON}_%",),
    ).fetchall()
    _write_rows(ws, rows)
    if not rows:
        ws.cell(row=3, column=1, value=(
            f"No real season-{SEASON} game has been AUDITED yet -- populates automatically "
            f"once a PREDICTED game's real result is in."
        )).font = NOTE_FONT


def build_data_quality_tab(wb, conn: sqlite3.Connection):
    ws = _rebuild_sheet(wb, "Data Quality", after="Prediction Audit")
    headers = ["Game ID", "Status", "Reason"]
    _title(ws, "Data Quality -- real NOT_PREDICTABLE / BLOCKED games with their real reason, "
               "plus recent real ingestion job health", len(headers))
    _headers(ws, headers)
    rows = conn.execute(
        """
        SELECT gws.game_id, gws.status, gws.status_reason
        FROM game_workflow_status gws JOIN games g ON g.game_id = gws.game_id
        WHERE g.season = ? AND gws.status IN ('NOT_PREDICTABLE', 'BLOCKED')
        ORDER BY g.week
        """,
        (SEASON,),
    ).fetchall()
    next_row = _write_rows(ws, rows)

    next_row += 2
    ws.cell(row=next_row, column=1, value="Recent real ingestion job health (last 20 runs):"
            ).font = TITLE_FONT
    health_headers = ["Job", "Timestamp", "Status", "Source", "Rows Written", "Detail"]
    _headers(ws, health_headers, row=next_row + 1)
    health_rows = conn.execute(
        "SELECT job_name, run_timestamp, status, source, rows_written, detail "
        "FROM ingestion_runs ORDER BY ingestion_id DESC LIMIT 20",
    ).fetchall()
    _write_rows(ws, health_rows, start_row=next_row + 2)
    ws.column_dimensions["C"].width = 60


def build_research_candidates_tab(wb):
    ws = _rebuild_sheet(wb, "Research Candidates", after="Data Quality")
    headers = ["Candidate", "Status", "Real Evidence"]
    _title(ws, "Research Candidates -- real Phase 9 graduation findings. Informational only "
               "-- nothing here is wired into the live production formula.", len(headers))
    _headers(ws, headers)
    # Real, static content sourced from phase_reports/phase9_graduation_table.md's own actual
    # findings (2026-09-08, post-HFA-bug-correction) -- not re-derived, transcribed directly.
    rows = [
        ("Travel -- nonlinear distance (G)", "TESTED",
         "MAE -0.326 vs champion (Phase 2/8); real 2nd-season evidence pending full-fidelity "
         "2026 data"),
        ("HFA -- revised, no HFA (A)", "TESTED, Phase 10-selected candidate (not in production)",
         "Best real Brier/log-loss of every candidate tested, both real seasons checked "
         "(2025 exact, 2024 secondary/degraded)"),
        ("Probability calibration -- Platt", "TESTED",
         "Brier 0.2410->0.2368 in isolation; needs refitting per Phase 8's own finding"),
        ("SOS -- A through H (8 variants)", "REJECTED",
         "Underperform champion or redundant with Base Team Quality/EPA (Phase 4/7/9)"),
        ("SOS -- I, nfelo-style market-derived", "REJECTED",
         "Real low correlation (r<=0.245) but real standalone backtest clearly worse than "
         "champion"),
        ("AGL (all variants)", "BLOCKED",
         "No sustainable real Approximate Value data source -- reconfirmed 3x total this "
         "session"),
    ]
    _write_rows(ws, rows)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["C"].width = 70


def build_change_log_tab(wb):
    ws = _rebuild_sheet(wb, "Change Log", after="Research Candidates")
    headers = ["Date", "Commit", "Description"]
    _title(ws, "Change Log -- real git commit history (this repo's own real audit trail, "
               "not a separate log to keep in sync)", len(headers))
    _headers(ws, headers)
    log = subprocess.run(
        ["git", "log", "--format=%ad|%h|%s", "--date=short", "-n", "40"],
        cwd=REPO_ROOT, capture_output=True, text=True, check=False,
    ).stdout
    rows = [tuple(line.split("|", 2)) for line in log.splitlines() if line.strip()]
    _write_rows(ws, rows)
    ws.column_dimensions["C"].width = 90


def build(workbook_path: str) -> None:
    wb = openpyxl.load_workbook(workbook_path)
    conn = sqlite3.connect(DEFAULT_DB_PATH)
    conn.row_factory = None

    build_schedule_tab(wb, conn)
    build_season_control_tab(wb, conn)
    build_pregame_snapshot_tab(wb, conn)
    build_market_capture_tab(wb, conn)
    build_results_tab(wb, conn)
    build_prediction_audit_tab(wb, conn)
    build_data_quality_tab(wb, conn)
    build_research_candidates_tab(wb)
    build_change_log_tab(wb)

    conn.close()
    wb.save(workbook_path)
    print("Built 9 real Live Weekly Workflow tabs: 2026 Game Schedule, 2026 Season Control, "
          "Pre-Game Snapshot, Market Capture, Results, Prediction Audit, Data Quality, "
          "Research Candidates, Change Log.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print('Usage: uv run python scripts/build_2026_live_tabs.py "path/to/workbook.xlsx"')
        sys.exit(1)
    build(sys.argv[1])
