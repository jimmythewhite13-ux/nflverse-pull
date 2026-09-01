"""
Builds "Player Prop Projections" per claude_code_spec_player_prop_projections.md --
deterministic Volume x Per-Unit-Efficiency point-estimate projections for every current
starter at QB/RB/WR1/WR2/WR3/TE1 (192 real player-roles), one row per that player's own real
game this season (17 games each -> 3,264 rows -- confirmed with the user: full season, all
starters, not a current-week-only ~192-row table).

WHAT THIS IS NOT (per the spec's own explicit "what this is NOT" section): no scraping of
any sportsbook's own prop lines (Part C's Sportsbook Line is a manual, weekly-updated blue
input, orange-filled like every other "needs weekly update" cell in this workbook, e.g.
Season Win Totals' own Vegas Win Total); no Monte Carlo or simulated-distribution logic
anywhere -- every projected number here is a single deterministic point estimate built from
real, already-computed inputs via plain arithmetic.

Part 0 -- Team Pass/Rush Attempt Pace mini-engine (Sections 1-3 below, this tab's own,
mirroring Secondary Index's own Section 1->2->3 TEAM-level chain): real 3-Yr decay-weighted,
regressed Pass Attempts/Game and Rush Attempts/Game per team
(player_props.compute_team_season_pass_rush_volume), current-season-blended via the shared
current_season_blend.add_current_season_blend() (claude_code_spec_current_season_blending.md
-- the exact same mechanism, same Model Assumptions C12/C13/C14 cells, every other tab's own
Section 3 already uses). Reuses the SAME shared decay/regression constants (Model Assumptions
C18/C20/C21/C22) every other tab's own Section 3 already reuses -- no new constants needed
for this chain.

Main table columns, one row per (player, real game):
  Player Name | Player ID | Position | Team | Role | Week | Opponent | Home/Away | Game Key
  (helper, real "Week|Away|Home" string identical to Season Matchups' own Game Key helper
  column -- player_props.compute_player_game_schedule() precomputes this, the real Opponent,
  and the real Home/Away flag in PYTHON rather than an Excel multi-criteria array match,
  keeping every downstream lookup on this ~3,264-row tab a single-criterion MATCH, the same
  design choice that avoided the "unwrapped multi-criteria array MATCH" bug class already
  caught once this project, Explosive Play Matchup Explanation Engine).

  Part A -- Expected Volume:
    Team Pace (Att/Gm, blended) -- this tab's own Section 3 blended value, Pass Attempts/Game
      for QB/WR/TE rows, Rush Attempts/Game for RB rows.
    DK Spread (Home, ref) -- Season Matchups' own real DraftKings Spread (Home) input (col
      AD, negative = home favored), cross-referenced by this row's real Game Key.
    Own Spread (+ = Underdog) -- DK Spread reoriented to THIS row's own team (Home rows keep
      the sign, Away rows flip it): positive = this team is the underdog.
    Game-Script Adj. Volume (Att/Gm) -- Team Pace + Own Spread * a NEW tunable Game-Script
      Volume Sensitivity constant (Model Assumptions C185), sign baked in per position at
      BUILD time (Python knows each row's own Position): +1 for QB/WR/TE (a bigger underdog
      -> more real pass volume, trailing-script), -1 for RB (a bigger favorite -> more real
      rush volume, clock-control script). ONE tunable constant, not two -- the direction is
      structural (which side of "pass or rush" this row's position falls on), not a second
      free parameter.
    Player Share (%) -- 100% for QB (the starter gets the whole team's real pass-attempt
      pool); real Carry Share (Y-1) for RB, cross-referenced from RB Value Index's own
      Section 6 (Team-keyed, that team's real starter/backup carry split); real, decay-
      weighted, current-season-blended Target Share for WR/TE, cross-referenced from WR-TE
      Value Index's own Section 3 (Player-ID-keyed, claude_code_spec_player_prop_
      projections.md's own new metric -- see build_wr_te_index.py's own METRICS entry).
    Projected Volume (Att/Tgt) -- Game-Script Adj. Volume * Player Share.

  Part B -- Expected Per-Unit Efficiency:
    Base Efficiency (Blended) -- QB: the new PURE Y/A blended baseline (QB Index Section 3,
      claude_code_spec_player_prop_projections.md's own metric, weight pinned to 0 in QB
      Index Score -- deliberately distinct from ANY/A, see build_qb_index.py's own METRICS
      entry). RB: the EXISTING YPC blended baseline (RB Value Index Section 3, already
      current-season-blended). WR/TE: the EXISTING YPT blended baseline (WR-TE Value Index
      Section 3, already current-season-blended).
    Matchup Differential (ref) -- the Defensive Matchup Engine's own real per-game
      differential, already wired onto Season Matchups (claude_code_spec_defensive_
      matchup_engine.md): Home/Away Pass Matchup Differential (cols AW/AZ) for QB/WR/TE rows,
      Home/Away Run Matchup Differential (cols BC/BF) for RB rows -- selected by this row's
      own real Home/Away flag.
    Matchup-Adj. Efficiency -- Base Efficiency + Matchup Differential * a NEW tunable
      Matchup-Adjusted Efficiency Sensitivity constant (Model Assumptions C186).
    Projected Yards -- Projected Volume * Matchup-Adj. Efficiency.
    Projected Receptions (WR/TE only, blank for QB/RB -- a real "not applicable", not a
      fabricated 0) -- Projected Volume (Targets) * WR-TE Value Index's own real, blended
      Catch Rate (Section 3, claude_code_spec_player_prop_projections.md's own metric). Not
      matchup-adjusted -- the spec's own Part B calls out Y/A/YPC/YPT as the matchup-adjusted
      efficiency metrics; Catch Rate is a separate volume-side reception-count stat.

  Part C -- Manual Sportsbook Line + Edge (blue input, orange "needs weekly update" fill,
  same convention as Season Win Totals' own Vegas Win Total column; read back across a
  rebuild by real "Player ID|Week" key, same "preserve manual inputs" pattern as Market
  Comparison & Confidence's own Kalshi/Polymarket prices):
    Sportsbook Yards Line | Yards Edge (Proj-Line) | Yards Recommended Play (Over/Under,
      threshold-based -- Model Assumptions C187, same IF(ABS(edge)>=threshold,...) pattern
      Season Matchups' own DK/MyBookie Recommended Play columns already use).
    Sportsbook Receptions Line | Receptions Edge (Proj-Line) | Receptions Recommended Play
      (WR/TE rows only, Model Assumptions C188).

No Team Ratings / Net Power Rating (N) involvement -- this is a pure downstream consolidation
tab, reading from Season Matchups, QB Index, RB Value Index, and WR-TE Value Index (all
already built earlier in the pipeline) via real INDEX/MATCH references, never writing back
into any of them. Placed last in main.py's pipeline for that reason.

Usage:
    uv run python scripts/build_player_prop_projections.py "C:\\path\\to\\NFL_Prediction_Model.xlsx"
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from add_manual_override_table import read_existing_overrides  # noqa: E402
from current_season_blend import add_current_season_blend  # noqa: E402

from nflverse_pull.availability import fetch_schedules_with_dates  # noqa: E402
from nflverse_pull.current_roster import (  # noqa: E402
    compute_current_starters,
    fetch_depth_charts,
    resolve_scored_population,
)
from nflverse_pull.efficiency import fetch_pbp  # noqa: E402
from nflverse_pull.player_props import (  # noqa: E402
    compute_player_game_schedule,
    compute_team_season_pass_rush_volume,
)
from nflverse_pull.season_schedule import (  # noqa: E402
    compute_roof_fallback,
    compute_season_schedule,
)

HISTORICAL_YEARS = [2023, 2024, 2025]
CURRENT_ROSTER_YEAR = 2026
CURRENT_SEASON = 2026
SHEET_NAME = "Player Prop Projections"
QB_SHEET = "QB Index"
RB_SHEET = "RB Value Index"
WR_TE_SHEET = "WR-TE Value Index"
SEASON_SHEET = "Season Matchups"

TEAM_ORDER = [
    "Buffalo Bills", "Miami Dolphins", "New England Patriots", "New York Jets",
    "Baltimore Ravens", "Cincinnati Bengals", "Cleveland Browns", "Pittsburgh Steelers",
    "Houston Texans", "Indianapolis Colts", "Jacksonville Jaguars", "Tennessee Titans",
    "Denver Broncos", "Kansas City Chiefs", "Las Vegas Raiders", "Los Angeles Chargers",
    "Dallas Cowboys", "New York Giants", "Philadelphia Eagles", "Washington Commanders",
    "Chicago Bears", "Detroit Lions", "Green Bay Packers", "Minnesota Vikings",
    "Atlanta Falcons", "Carolina Panthers", "New Orleans Saints", "Tampa Bay Buccaneers",
    "Arizona Cardinals", "Los Angeles Rams", "San Francisco 49ers", "Seattle Seahawks",
]

PACE_METRICS = [
    {"key": "pass_att", "col": "Pass Attempts/Game", "label": "Pass\nAtt/Gm", "fmt": "0.00"},
    {"key": "rush_att", "col": "Rush Attempts/Game", "label": "Rush\nAtt/Gm", "fmt": "0.00"},
]

TITLE_FONT = Font(name="Arial", size=10, bold=True)
TITLE_FILL = PatternFill("solid", fgColor="FFD9E1F2")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_ALIGN = Alignment(wrap_text=True, horizontal="center", vertical="center")
INPUT_FONT = Font(name="Arial", size=10, color="FF0000FF")
FORMULA_FONT = Font(name="Arial", size=10, color="FF000000")
LINK_FONT = Font(name="Arial", size=10, color="FF008000")
NOTE_FONT = Font(name="Arial", size=9, color="FF808080")
ASSUMPTION_FILL = PatternFill("solid", fgColor="FFFFFF00")
NEEDS_UPDATE_FILL = PatternFill("solid", fgColor="FFFFA500")

COLUMNS = [
    "Player Name", "Player ID", "Position", "Team", "Role", "Week", "Opponent", "Home/Away",
    "Game Key\n(helper)",
    "Team Pace\n(Att/Gm, blended)", "DK Spread\n(Home, ref)", "Own Spread\n(+ = Underdog)",
    "Game-Script Adj.\nVolume (Att/Gm)", "Player Share\n(%)", "Projected Volume\n(Att/Tgt)",
    "Base Efficiency\n(Blended)", "Matchup Differential\n(ref)", "Matchup-Adj.\nEfficiency",
    "Projected Yards", "Projected\nReceptions",
    "Sportsbook Yards\nLine", "Yards Edge\n(Proj-Line)", "Yards Recommended\nPlay",
    "Sportsbook Receptions\nLine", "Receptions Edge\n(Proj-Line)",
    "Receptions Recommended\nPlay",
]
COL = {name: i + 1 for i, name in enumerate(COLUMNS)}
LET = {name: get_column_letter(i + 1) for i, name in enumerate(COLUMNS)}


def _section_title(ws, row: int, last_col: int, text: str) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL


def _header_row(ws, row: int, headers: list[str], height: float = 28) -> None:
    ws.row_dimensions[row].height = height
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN


def _find_title_row(ws, needle: str) -> int:
    for row in range(1, ws.max_row + 1):
        val = ws.cell(row=row, column=1).value
        if val and str(val).strip().startswith(needle):
            return row
    raise ValueError(f"Could not find a row starting with {needle!r} in '{ws.title}'.")


def _find_header_col(ws, header_row: int, text: str) -> int:
    for col in range(1, ws.max_column + 1):
        if ws.cell(row=header_row, column=col).value == text:
            return col
    raise ValueError(f"Could not find header {text!r} in row {header_row} of '{ws.title}'.")


def _last_row_by_col(ws, first_row: int, col: int) -> int:
    last_row = first_row
    while ws.cell(row=last_row + 1, column=col).value is not None:
        last_row += 1
    return last_row


def add_model_assumptions_weights(wb: openpyxl.Workbook) -> None:
    ws = wb["Model Assumptions"]
    title_row = 184
    ws.merge_cells(f"A{title_row}:D{title_row}")
    title = ws.cell(row=title_row, column=1, value=(
        "Player Prop Projections -- Volume x Per-Unit-Efficiency (see 'Player Prop "
        "Projections' tab). Deterministic point estimates only -- no Monte Carlo/simulated "
        "distribution anywhere in this project."
    ))
    title.font = HEADER_FONT
    title.fill = HEADER_FILL

    rows = [
        (185, "Game-Script Volume Sensitivity (Att/Gm per pt of own spread)", 0.05,
         "claude_code_spec_player_prop_projections.md Part A: how much a team's real real "
         "expected pass/rush volume shifts per point of its own real spread (positive = "
         "underdog). A starting guess, like every other coefficient in this model -- applied "
         "with sign +1 (pass volume rises for a bigger underdog) for QB/WR/TE rows, -1 (rush "
         "volume rises for a bigger favorite, clock-control script) for RB rows, baked in "
         "per-row at build time rather than a second free parameter."),
        (186, "Matchup-Adjusted Efficiency Sensitivity (per pt of Index-Score differential)",
         0.01,
         "claude_code_spec_player_prop_projections.md Part B: how much a player's real "
         "blended per-unit efficiency (Y/A for QB, YPC for RB, YPT for WR/TE) shifts per "
         "point of the Defensive Matchup Engine's own real Pass/Run Matchup Differential "
         "(already wired onto Season Matchups). A starting guess, like every other "
         "coefficient in this model."),
        (187, "Player Prop Yards Edge Threshold (min |edge| to recommend Over/Under)", 5.0,
         "Same threshold-based Recommended Play pattern Season Matchups' own DK/MyBookie "
         "Spread/Total Recommended Play columns already use -- a separate constant since "
         "this measures yards, not game points."),
        (188, "Player Prop Receptions Edge Threshold (min |edge| to recommend Over/Under)",
         0.5,
         "Same threshold-based Recommended Play pattern as C187, in receptions rather than "
         "yards -- WR/TE rows only."),
    ]
    for row, label, value, note in rows:
        ws.cell(row=row, column=2, value=label)
        c = ws.cell(row=row, column=3, value=value)
        c.font = INPUT_FONT
        c.fill = ASSUMPTION_FILL
        c.number_format = "0.00"
        n = ws.cell(row=row, column=4, value=note)
        n.font = NOTE_FONT
        n.alignment = Alignment(wrap_text=True, vertical="top")


def _read_back(wb: openpyxl.Workbook) -> dict[str, dict[str, float]]:
    """
    Reads Part C's manual Sportsbook Line inputs back from an existing 'Player Prop
    Projections' sheet (if present) BEFORE this script deletes and rebuilds the whole sheet
    -- keyed by real "Player ID|Week" (a player plays exactly one real game per week, so this
    is a real, stable, unique key across a rebuild even though the sheet's own row order can
    shift week to week with the real depth chart). Same "preserve manual inputs across a
    rebuild" pattern as Market Comparison & Confidence's own Kalshi/Polymarket price read-back
    and Season Matchups' own DK/MyBookie odds read-back.
    """
    if SHEET_NAME not in wb.sheetnames:
        return {}
    ws = wb[SHEET_NAME]
    header_row = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "Player Name":
            header_row = r
            break
    if header_row is None:
        return {}
    headers = {ws.cell(row=header_row, column=c).value: c for c in range(1, ws.max_column + 1)}
    pid_col = headers.get("Player ID")
    week_col = headers.get("Week")
    yards_col = headers.get("Sportsbook Yards\nLine")
    rec_col = headers.get("Sportsbook Receptions\nLine")
    if not (pid_col and week_col):
        return {}

    out: dict[str, dict[str, float]] = {}
    row = header_row + 1
    while ws.cell(row=row, column=1).value is not None:
        pid = ws.cell(row=row, column=pid_col).value
        wk = ws.cell(row=row, column=week_col).value
        if pid is not None and wk is not None:
            key = f"{pid}|{int(wk)}"
            vals = {}
            if yards_col:
                v = ws.cell(row=row, column=yards_col).value
                if v not in (None, "") and not (isinstance(v, str) and v.startswith("=")):
                    vals["yards"] = v
            if rec_col:
                v = ws.cell(row=row, column=rec_col).value
                if v not in (None, "") and not (isinstance(v, str) and v.startswith("=")):
                    vals["rec"] = v
            if vals:
                out[key] = vals
        row += 1
    return out


def build(workbook_path: str) -> dict:
    print(f"Pulling {HISTORICAL_YEARS} play-by-play data for Team Pace / real Y-A / Target "
          "Share / Catch Rate inputs...")
    pbp = fetch_pbp(HISTORICAL_YEARS)
    team_pace = compute_team_season_pass_rush_volume(pbp)

    print(f"Pulling {CURRENT_SEASON} real schedule for real Week/Opponent/Home-Away...")
    sched_multi = fetch_schedules_with_dates(HISTORICAL_YEARS)
    sched_current = fetch_schedules_with_dates([CURRENT_SEASON])
    roof_fallback = compute_roof_fallback(sched_multi)
    schedule = compute_season_schedule(sched_current, CURRENT_SEASON, roof_fallback)

    print(f"Pulling {CURRENT_ROSTER_YEAR} depth charts for the QB/RB/WR/TE population...")
    depth_charts = fetch_depth_charts([CURRENT_ROSTER_YEAR])
    current_starters = compute_current_starters(depth_charts)

    wb = openpyxl.load_workbook(workbook_path)
    existing_lines = _read_back(wb)
    print(f"Read back {len(existing_lines)} existing manual Sportsbook Line entr(y/ies) "
          "before rebuilding the sheet.")

    # Population = the CURRENT real starter at each scored role, reading each source tab's own
    # ALREADY-REBUILT Section 7 Manual Roster Override (this script runs last in main.py's
    # pipeline, after QB Index/RB Value Index/WR-TE Value Index and their own Section 7
    # rebuilds) -- so an override entered on any of those tabs is reflected here too, not a
    # second, independently-pulled population that could silently disagree with them.
    qb_overrides = read_existing_overrides(wb, QB_SHEET, ["Starter", "Backup"])
    rb_overrides = read_existing_overrides(wb, RB_SHEET, ["Starter", "Backup"])
    wrte_overrides = read_existing_overrides(wb, WR_TE_SHEET, ["WR1", "WR2", "WR3", "TE1"])

    qb_pop = resolve_scored_population(current_starters, qb_overrides, "QB")
    qb_pop = qb_pop[qb_pop["Role"] == "Starter"].copy()
    qb_pop["Position"] = "QB"
    rb_pop = resolve_scored_population(current_starters, rb_overrides, "RB")
    rb_pop = rb_pop[rb_pop["Role"] == "Starter"].copy()
    rb_pop["Position"] = "RB"
    wr_pop = resolve_scored_population(current_starters, wrte_overrides, "WR")
    wr_pop = wr_pop.copy()
    wr_pop["Position"] = "WR"
    te_pop = resolve_scored_population(current_starters, wrte_overrides, "TE")
    te_pop = te_pop.copy()
    te_pop["Position"] = "TE"
    population = pd.concat([qb_pop, rb_pop, wr_pop, te_pop], ignore_index=True)
    print(f"{len(population)} current-roster player-roles identified "
          f"({(population['Position'] == 'QB').sum()} QB, "
          f"{(population['Position'] == 'RB').sum()} RB, "
          f"{(population['Position'] == 'WR').sum()} WR, "
          f"{(population['Position'] == 'TE').sum()} TE).")

    schedule_rows = compute_player_game_schedule(population, schedule)
    games_per_role = len(schedule_rows) // max(len(population), 1)
    print(f"{len(schedule_rows)} real player-game rows ({len(population)} player-roles x "
          f"{games_per_role} real games per player-role).")

    add_model_assumptions_weights(wb)

    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    if "Season Win Totals" in wb.sheetnames:
        insert_after = "Season Win Totals"
    else:
        insert_after = wb.sheetnames[-1]
    ws = wb.create_sheet(SHEET_NAME, index=wb.sheetnames.index(insert_after) + 1)
    ws.column_dimensions["A"].width = 20.0

    ws.merge_cells(f"A1:{get_column_letter(len(COLUMNS))}1")
    t = ws.cell(row=1, column=1, value=(
        "Player Prop Projections -- deterministic Volume x Per-Unit-Efficiency point "
        "estimates for every current QB/RB/WR1-3/TE1 starter's own real game this season. "
        "NO scraped sportsbook data, NO Monte Carlo/simulated distribution anywhere -- "
        "every number here is plain arithmetic over real, already-computed inputs. See the "
        "closing note."
    ))
    t.font = Font(name="Arial", size=12, bold=True)

    # ==== Part 0: Team Pass/Rush Attempt Pace mini-engine (this tab's own Section 1-3, =====
    # mirroring Secondary Index's own TEAM-level Section 1->2->3 chain) ======================
    pace_sec1_title_row = 3
    pace_sec1_header_row = 4
    pace_sec1_first_row = 5
    pace_sec1_last_row = pace_sec1_first_row + len(team_pace) - 1
    _section_title(
        ws, pace_sec1_title_row, 4,
        "Section 1 \u2014 Raw 3-Year TEAM-Level Pass/Rush Attempts per Game (real nflverse "
        "play-by-play, sacks excluded from Pass Attempts -- see player_props.py's own "
        "docstring)",
    )
    _header_row(ws, pace_sec1_header_row, ["Team", "Season", *[m["label"] for m in PACE_METRICS]])
    for i, r in enumerate(team_pace.to_dict("records")):
        row = pace_sec1_first_row + i
        values = [r["Team"], int(r["Season"]), r["Pass Attempts/Game"], r["Rush Attempts/Game"]]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.font = INPUT_FONT
        for col in (3, 4):
            ws.cell(row=row, column=col).number_format = "0.00"

    pace_team_range_sec1 = f"$A${pace_sec1_first_row}:$A${pace_sec1_last_row}"
    pace_season_range = f"$B${pace_sec1_first_row}:$B${pace_sec1_last_row}"
    pace_sec1_col_of = {"pass_att": "C", "rush_att": "D"}
    pace_metric_ranges = {
        m["key"]: (
            f"${pace_sec1_col_of[m['key']]}${pace_sec1_first_row}:"
            f"${pace_sec1_col_of[m['key']]}${pace_sec1_last_row}"
        )
        for m in PACE_METRICS
    }

    pace_sec2_title_row = pace_sec1_last_row + 2
    pace_sec2_header_row = pace_sec2_title_row + 1
    pace_season_rows = {
        yr: pace_sec2_header_row + 1 + i for i, yr in enumerate(HISTORICAL_YEARS)
    }
    _section_title(
        ws, pace_sec2_title_row, 3,
        "Section 2 \u2014 League Average per Season (simple average across all 32 teams)",
    )
    _header_row(
        ws, pace_sec2_header_row, ["Season", *[m["label"] for m in PACE_METRICS]], height=20
    )
    for yr in HISTORICAL_YEARS:
        row = pace_season_rows[yr]
        c = ws.cell(row=row, column=1, value=yr)
        c.font = INPUT_FONT
        for j, m in enumerate(PACE_METRICS):
            formula = f"=AVERAGEIF({pace_season_range},{yr},{pace_metric_ranges[m['key']]})"
            cell = ws.cell(row=row, column=2 + j, value=formula)
            cell.font = FORMULA_FONT
            cell.number_format = m["fmt"]
    pace_season_avg_col = {m["key"]: get_column_letter(2 + j) for j, m in enumerate(PACE_METRICS)}

    pace_sec3_title_row = pace_season_rows[HISTORICAL_YEARS[-1]] + 2
    pace_sec3_header_row = pace_sec3_title_row + 1
    pace_sec3_first_row = pace_sec3_header_row + 1
    pace_sec3_last_row = pace_sec3_first_row + len(TEAM_ORDER) - 1
    pace_sec3_last_col = 2 + len(PACE_METRICS) * 7
    _section_title(
        ws, pace_sec3_title_row, pace_sec3_last_col,
        "Section 3 \u2014 Per-Team 3-Yr Decay-Weighted, Regressed Baseline (reuses the SAME "
        "shared decay/regression constants, Model Assumptions C18/C20/C21/C22, every other "
        "tab's own Section 3 already uses -- no new constants for this chain). A missing "
        "year substitutes THAT season's own Section 2 league average (no rookie concept for "
        "a team-level pace stat).",
    )
    headers = ["Team", "Years of\nReal History"]
    pace_metric_block_start_col: dict[str, int] = {}
    col_cursor = 3
    for m in PACE_METRICS:
        pace_metric_block_start_col[m["key"]] = col_cursor
        headers += [
            f"{m['label']} Y-1", f"{m['label']} Y-2", f"{m['label']} Y-3",
            f"Weighted\n{m['label']} Avg\n(3-Yr decay)", f"Team History\n{m['label']}",
            f"League Baseline\n{m['label']} (Y-1)", f"Projected 3-Yr\n{m['label']} Baseline",
        ]
        col_cursor += 7
    _header_row(ws, pace_sec3_header_row, headers)

    for i, team in enumerate(TEAM_ORDER):
        row = pace_sec3_first_row + i
        ws.cell(row=row, column=1, value=team).font = FORMULA_FONT

        years_hist_terms = "+".join(
            f"--(COUNTIFS({pace_team_range_sec1},$A{row},{pace_season_range},"
            f"'Model Assumptions'!$C$18-{k})>0)"
            for k in (1, 2, 3)
        )
        yh = ws.cell(row=row, column=2, value=f"={years_hist_terms}")
        yh.font = FORMULA_FONT
        yh.number_format = "0"

        for m in PACE_METRICS:
            base = pace_metric_block_start_col[m["key"]]
            y1, y2, y3, wavg, th, lb, pb = (get_column_letter(base + k) for k in range(7))
            mrange = pace_metric_ranges[m["key"]]
            sec2_col = pace_season_avg_col[m["key"]]

            def _ysub(offset: int, mrange=mrange, sec2_col=sec2_col) -> str:
                season_avg_lookup = (
                    f"INDEX(${sec2_col}${pace_season_rows[HISTORICAL_YEARS[0]]}:"
                    f"${sec2_col}${pace_season_rows[HISTORICAL_YEARS[-1]]},"
                    f"MATCH('Model Assumptions'!$C$18-{offset},"
                    f"$A${pace_season_rows[HISTORICAL_YEARS[0]]}:"
                    f"$A${pace_season_rows[HISTORICAL_YEARS[-1]]},0))"
                )
                return (
                    f"=IF(COUNTIFS({pace_team_range_sec1},$A{row},{pace_season_range},"
                    f"'Model Assumptions'!$C$18-{offset})=0,{season_avg_lookup},"
                    f"SUMIFS({mrange},{pace_team_range_sec1},$A{row},{pace_season_range},"
                    f"'Model Assumptions'!$C$18-{offset}))"
                )

            f_y1 = ws.cell(row=row, column=base, value=_ysub(1))
            f_y2 = ws.cell(row=row, column=base + 1, value=_ysub(2))
            f_y3 = ws.cell(row=row, column=base + 2, value=_ysub(3))
            f_wavg = ws.cell(row=row, column=base + 3, value=(
                f"=({y1}{row}*1+{y2}{row}*'Model Assumptions'!$C$20+"
                f"{y3}{row}*('Model Assumptions'!$C$20^2))/"
                f"(1+'Model Assumptions'!$C$20+'Model Assumptions'!$C$20^2)"
            ))
            f_th = ws.cell(row=row, column=base + 4, value=(
                f"={y1}{row}*'Model Assumptions'!$C$22+{wavg}{row}*(1-'Model Assumptions'!$C$22)"
            ))
            f_lb = ws.cell(row=row, column=base + 5, value=(
                f"=INDEX(${sec2_col}${pace_season_rows[HISTORICAL_YEARS[0]]}:"
                f"${sec2_col}${pace_season_rows[HISTORICAL_YEARS[-1]]},"
                f"MATCH('Model Assumptions'!$C$18-1,$A${pace_season_rows[HISTORICAL_YEARS[0]]}:"
                f"$A${pace_season_rows[HISTORICAL_YEARS[-1]]},0))"
            ))
            f_pb = ws.cell(row=row, column=base + 6, value=(
                f"={th}{row}*'Model Assumptions'!$C$21+{lb}{row}*(1-'Model Assumptions'!$C$21)"
            ))
            for cell in (f_y1, f_y2, f_y3, f_wavg, f_th, f_lb, f_pb):
                cell.font = FORMULA_FONT
                cell.number_format = m["fmt"]

    pace_proj_baseline_col = {
        m["key"]: get_column_letter(pace_metric_block_start_col[m["key"]] + 6)
        for m in PACE_METRICS
    }
    pace_blended_cols = add_current_season_blend(
        ws, team_col="A", sec3_first_row=pace_sec3_first_row, sec3_last_row=pace_sec3_last_row,
        start_col=pace_sec3_last_col + 1,
        metrics=[
            {"key": m["key"], "label": m["label"], "pb_col": pace_proj_baseline_col[m["key"]]}
            for m in PACE_METRICS
        ],
    )
    pace_team_range = f"$A${pace_sec3_first_row}:$A${pace_sec3_last_row}"
    pace_pass_range = (
        f"${pace_blended_cols['pass_att']}${pace_sec3_first_row}:"
        f"${pace_blended_cols['pass_att']}${pace_sec3_last_row}"
    )
    pace_rush_range = (
        f"${pace_blended_cols['rush_att']}${pace_sec3_first_row}:"
        f"${pace_blended_cols['rush_att']}${pace_sec3_last_row}"
    )

    # ==== Real ranges on the other tabs, discovered dynamically (survives any row-count =====
    # shift -- same convention build_market_comparison_confidence.py's own sm_col() uses) ====
    qb_ws = wb[QB_SHEET]
    qb_sec3_title = _find_title_row(qb_ws, "Section 3")
    qb_sec3_header = qb_sec3_title + 1
    qb_sec3_first = qb_sec3_header + 1
    qb_sec3_last = _last_row_by_col(qb_ws, qb_sec3_first, 2)
    qb_ya_col = get_column_letter(
        _find_header_col(qb_ws, qb_sec3_header, "Blended\nY/A\n(volume-proj.\nonly)")
    )
    qb_id_range = f"'{QB_SHEET}'!$B${qb_sec3_first}:$B${qb_sec3_last}"
    qb_ya_range = f"'{QB_SHEET}'!${qb_ya_col}${qb_sec3_first}:${qb_ya_col}${qb_sec3_last}"

    rb_ws = wb[RB_SHEET]
    rb_sec3_title = _find_title_row(rb_ws, "Section 3")
    rb_sec3_header = rb_sec3_title + 1
    rb_sec3_first = rb_sec3_header + 1
    rb_sec3_last = _last_row_by_col(rb_ws, rb_sec3_first, 2)
    rb_ypc_col = get_column_letter(_find_header_col(rb_ws, rb_sec3_header, "Blended\nYPC"))
    rb_id_range = f"'{RB_SHEET}'!$B${rb_sec3_first}:$B${rb_sec3_last}"
    rb_ypc_range = f"'{RB_SHEET}'!${rb_ypc_col}${rb_sec3_first}:${rb_ypc_col}${rb_sec3_last}"

    rb_sec6_title = _find_title_row(rb_ws, "Section 6")
    rb_sec6_header = rb_sec6_title + 1
    rb_sec6_first = rb_sec6_header + 1
    rb_sec6_last = _last_row_by_col(rb_ws, rb_sec6_first, 1)
    rb_team_range = f"'{RB_SHEET}'!$A${rb_sec6_first}:$A${rb_sec6_last}"
    rb_carryshare_range = f"'{RB_SHEET}'!$H${rb_sec6_first}:$H${rb_sec6_last}"

    wrte_ws = wb[WR_TE_SHEET]
    wrte_sec3_title = _find_title_row(wrte_ws, "Section 3")
    wrte_sec3_header = wrte_sec3_title + 1
    wrte_sec3_first = wrte_sec3_header + 1
    wrte_sec3_last = _last_row_by_col(wrte_ws, wrte_sec3_first, 2)
    wrte_ypt_col = get_column_letter(_find_header_col(wrte_ws, wrte_sec3_header, "Blended\nYPT"))
    wrte_ts_col = get_column_letter(_find_header_col(
        wrte_ws, wrte_sec3_header, "Blended\nTarget Share\n(volume-proj.\nonly)"
    ))
    wrte_cr_col = get_column_letter(_find_header_col(
        wrte_ws, wrte_sec3_header, "Blended\nCatch Rate\n(volume-proj.\nonly)"
    ))
    wrte_id_range = f"'{WR_TE_SHEET}'!$B${wrte_sec3_first}:$B${wrte_sec3_last}"
    wrte_ypt_range = (
        f"'{WR_TE_SHEET}'!${wrte_ypt_col}${wrte_sec3_first}:${wrte_ypt_col}${wrte_sec3_last}"
    )
    wrte_ts_range = (
        f"'{WR_TE_SHEET}'!${wrte_ts_col}${wrte_sec3_first}:${wrte_ts_col}${wrte_sec3_last}"
    )
    wrte_cr_range = (
        f"'{WR_TE_SHEET}'!${wrte_cr_col}${wrte_sec3_first}:${wrte_cr_col}${wrte_sec3_last}"
    )

    sm_ws = wb[SEASON_SHEET]
    sm_first, sm_last = 3, 3
    while sm_ws.cell(row=sm_last + 1, column=1).value is not None:
        sm_last += 1

    def sm_col(letter: str) -> str:
        return f"'{SEASON_SHEET}'!${letter}${sm_first}:${letter}${sm_last}"

    sm_key_range = sm_col(get_column_letter(110))
    dk_spread_range = sm_col(get_column_letter(30))
    home_pass_diff_range = sm_col(get_column_letter(49))
    away_pass_diff_range = sm_col(get_column_letter(52))
    home_run_diff_range = sm_col(get_column_letter(55))
    away_run_diff_range = sm_col(get_column_letter(58))

    # ==== Main table: one row per (player, real game) ========================================
    main_title_row = pace_sec3_last_row + 3
    main_header_row = main_title_row + 1
    main_first_row = main_header_row + 1
    _section_title(
        ws, main_title_row, len(COLUMNS),
        "Player Prop Projections -- one row per current starter's own real game this season "
        "(Part A: Expected Volume, Part B: Expected Per-Unit Efficiency, Part C: manual "
        "Sportsbook Line + Edge). See this tab's own top-of-sheet title and closing note.",
    )
    ws.row_dimensions[main_header_row].height = 34
    for name in COLUMNS:
        c = ws.cell(row=main_header_row, column=COL[name], value=name)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = HEADER_ALIGN

    def L(name: str, row: int) -> str:
        return f"{LET[name]}{row}"

    for i, rec in enumerate(schedule_rows.to_dict("records")):
        row = main_first_row + i
        position = rec["Position"]
        home_away = rec["Home/Away"]
        week = int(rec["Week"])

        ref = {
            "team": L("Team", row), "game_key": L("Game Key\n(helper)", row),
            "pid": L("Player ID", row),
            "pace": L("Team Pace\n(Att/Gm, blended)", row),
            "dk": L("DK Spread\n(Home, ref)", row),
            "own_spread": L("Own Spread\n(+ = Underdog)", row),
            "gs_vol": L("Game-Script Adj.\nVolume (Att/Gm)", row),
            "share": L("Player Share\n(%)", row),
            "vol": L("Projected Volume\n(Att/Tgt)", row),
            "base_eff": L("Base Efficiency\n(Blended)", row),
            "diff": L("Matchup Differential\n(ref)", row),
            "madj_eff": L("Matchup-Adj.\nEfficiency", row),
            "proj_yards": L("Projected Yards", row),
            "proj_rec": L("Projected\nReceptions", row),
            "yards_line": L("Sportsbook Yards\nLine", row),
            "yards_edge": L("Yards Edge\n(Proj-Line)", row),
            "rec_line": L("Sportsbook Receptions\nLine", row),
            "rec_edge": L("Receptions Edge\n(Proj-Line)", row),
        }

        for name, val in (
            ("Player Name", rec["Player Name"]), ("Player ID", rec["Player ID"]),
            ("Position", position), ("Team", rec["Team"]), ("Role", rec["Role"]),
            ("Week", week), ("Opponent", rec["Opponent"]), ("Home/Away", home_away),
            ("Game Key\n(helper)", rec["Game Key"]),
        ):
            cell = ws.cell(row=row, column=COL[name], value=val)
            cell.font = INPUT_FONT
        ws.cell(row=row, column=COL["Week"]).number_format = "0"

        # ---- Part A: Expected Volume -------------------------------------------------------
        is_pass_position = position in ("QB", "WR", "TE")
        pace_range = pace_pass_range if is_pass_position else pace_rush_range
        pace_c = ws.cell(row=row, column=COL["Team Pace\n(Att/Gm, blended)"], value=(
            f'=IFERROR(INDEX({pace_range},MATCH({ref["team"]},{pace_team_range},0)),"")'
        ))
        pace_c.font = FORMULA_FONT
        pace_c.number_format = "0.00"

        dk_c = ws.cell(row=row, column=COL["DK Spread\n(Home, ref)"], value=(
            f'=IFERROR(INDEX({dk_spread_range},MATCH({ref["game_key"]},{sm_key_range},0)),"")'
        ))
        dk_c.font = LINK_FONT
        dk_c.number_format = "0.0;(0.0)"

        own_sign = "" if home_away == "Home" else "-"
        own_c = ws.cell(row=row, column=COL["Own Spread\n(+ = Underdog)"], value=(
            f'=IF({ref["dk"]}="","",{own_sign}{ref["dk"]})'
        ))
        own_c.font = FORMULA_FONT
        own_c.number_format = "0.0;(0.0)"

        dir_sign = "+" if is_pass_position else "-"
        gs_c = ws.cell(row=row, column=COL["Game-Script Adj.\nVolume (Att/Gm)"], value=(
            f'=IF(OR({ref["pace"]}="",{ref["own_spread"]}=""),"",'
            f'{ref["pace"]}{dir_sign}{ref["own_spread"]}*\'Model Assumptions\'!$C$185)'
        ))
        gs_c.font = FORMULA_FONT
        gs_c.number_format = "0.00"

        share_cell = ws.cell(row=row, column=COL["Player Share\n(%)"])
        if position == "QB":
            share_cell.value = "=1"
            share_cell.font = FORMULA_FONT
        elif position == "RB":
            share_cell.value = (
                f'=IFERROR(INDEX({rb_carryshare_range},MATCH({ref["team"]},'
                f'{rb_team_range},0)),"")'
            )
            share_cell.font = LINK_FONT
        else:
            share_cell.value = (
                f'=IFERROR(INDEX({wrte_ts_range},MATCH({ref["pid"]},{wrte_id_range},0)),"")'
            )
            share_cell.font = LINK_FONT
        share_cell.number_format = "0.0%"

        vol_c = ws.cell(row=row, column=COL["Projected Volume\n(Att/Tgt)"], value=(
            f'=IF(OR({ref["gs_vol"]}="",{ref["share"]}=""),"",{ref["gs_vol"]}*{ref["share"]})'
        ))
        vol_c.font = FORMULA_FONT
        vol_c.number_format = "0.00"

        # ---- Part B: Expected Per-Unit Efficiency -------------------------------------------
        if position == "QB":
            eff_range, eff_id_range = qb_ya_range, qb_id_range
        elif position == "RB":
            eff_range, eff_id_range = rb_ypc_range, rb_id_range
        else:
            eff_range, eff_id_range = wrte_ypt_range, wrte_id_range
        base_c = ws.cell(row=row, column=COL["Base Efficiency\n(Blended)"], value=(
            f'=IFERROR(INDEX({eff_range},MATCH({ref["pid"]},{eff_id_range},0)),"")'
        ))
        base_c.font = LINK_FONT
        base_c.number_format = "0.00"

        if is_pass_position:
            diff_range = home_pass_diff_range if home_away == "Home" else away_pass_diff_range
        else:
            diff_range = home_run_diff_range if home_away == "Home" else away_run_diff_range
        diff_c = ws.cell(row=row, column=COL["Matchup Differential\n(ref)"], value=(
            f'=IFERROR(INDEX({diff_range},MATCH({ref["game_key"]},{sm_key_range},0)),"")'
        ))
        diff_c.font = LINK_FONT
        diff_c.number_format = "0.0;(0.0)"

        madj_c = ws.cell(row=row, column=COL["Matchup-Adj.\nEfficiency"], value=(
            f'=IF(OR({ref["base_eff"]}="",{ref["diff"]}=""),"",'
            f'{ref["base_eff"]}+{ref["diff"]}*\'Model Assumptions\'!$C$186)'
        ))
        madj_c.font = FORMULA_FONT
        madj_c.number_format = "0.00"

        yards_c = ws.cell(row=row, column=COL["Projected Yards"], value=(
            f'=IF(OR({ref["vol"]}="",{ref["madj_eff"]}=""),"",{ref["vol"]}*{ref["madj_eff"]})'
        ))
        yards_c.font = FORMULA_FONT
        yards_c.number_format = "0.0;(0.0)"

        is_receiver = position in ("WR", "TE")
        if is_receiver:
            rec_c = ws.cell(row=row, column=COL["Projected\nReceptions"], value=(
                f'=IF({ref["vol"]}="","",IFERROR({ref["vol"]}*INDEX({wrte_cr_range},'
                f'MATCH({ref["pid"]},{wrte_id_range},0)),""))'
            ))
            rec_c.font = FORMULA_FONT
            rec_c.number_format = "0.00"

        # ---- Part C: manual Sportsbook Line + Edge (read back, orange "needs weekly ---------
        # update" fill, same convention as Season Win Totals' own Vegas Win Total). -----------
        key = f"{rec['Player ID']}|{week}"
        existing = existing_lines.get(key, {})

        yl_c = ws.cell(
            row=row, column=COL["Sportsbook Yards\nLine"], value=existing.get("yards")
        )
        yl_c.font = INPUT_FONT
        yl_c.fill = NEEDS_UPDATE_FILL
        yl_c.number_format = "0.0"

        ye_c = ws.cell(row=row, column=COL["Yards Edge\n(Proj-Line)"], value=(
            f'=IF(OR({ref["yards_line"]}="",{ref["proj_yards"]}=""),"",'
            f'{ref["proj_yards"]}-{ref["yards_line"]})'
        ))
        ye_c.font = FORMULA_FONT
        ye_c.number_format = "0.0;(0.0)"

        yp_c = ws.cell(row=row, column=COL["Yards Recommended\nPlay"], value=(
            f'=IF({ref["yards_edge"]}="","No Line",IF(ABS({ref["yards_edge"]})>='
            f'\'Model Assumptions\'!$C$187,IF({ref["yards_edge"]}>0,"Over","Under"),'
            f'"No Edge"))'
        ))
        yp_c.font = FORMULA_FONT

        if is_receiver:
            rl_c = ws.cell(
                row=row, column=COL["Sportsbook Receptions\nLine"], value=existing.get("rec")
            )
            rl_c.font = INPUT_FONT
            rl_c.fill = NEEDS_UPDATE_FILL
            rl_c.number_format = "0.00"

            re_c = ws.cell(row=row, column=COL["Receptions Edge\n(Proj-Line)"], value=(
                f'=IF(OR({ref["rec_line"]}="",{ref["proj_rec"]}=""),"",'
                f'{ref["proj_rec"]}-{ref["rec_line"]})'
            ))
            re_c.font = FORMULA_FONT
            re_c.number_format = "0.00;(0.00)"

            rp_c = ws.cell(row=row, column=COL["Receptions Recommended\nPlay"], value=(
                f'=IF({ref["rec_edge"]}="","No Line",IF(ABS({ref["rec_edge"]})>='
                f'\'Model Assumptions\'!$C$188,IF({ref["rec_edge"]}>0,"Over","Under"),'
                f'"No Edge"))'
            ))
            rp_c.font = FORMULA_FONT

    main_last_row = main_first_row + len(schedule_rows) - 1

    note_row = main_last_row + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=14)
    note = ws.cell(row=note_row, column=1, value=(
        "NO scraping of any sportsbook's own prop lines anywhere in this project -- Part C's "
        "Sportsbook Line is a manual, weekly-updated input (orange fill), same convention as "
        "every other 'needs weekly update' cell in this workbook. NO Monte Carlo or "
        "simulated-distribution logic anywhere -- every projected number is a single "
        "deterministic point estimate from real, already-computed inputs via plain "
        "arithmetic. QB's Pure Y/A (Base Efficiency) is a SEPARATE metric from QB Index's "
        "own ANY/A, weight pinned to 0 in QB Index Score -- ANY/A itself is untouched. "
        "WR/TE's Target Share is genuinely NEW and distinct from Pass-Play Snap "
        "Participation % (routes run vs. targets actually received) -- see build_wr_te_"
        "index.py's own METRICS entries for both."
    ))
    note.font = NOTE_FONT
    note.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(workbook_path)
    print(
        f"Built '{SHEET_NAME}': Part 0 Section 1 {len(team_pace)} rows, Section 3 "
        f"{len(TEAM_ORDER)} teams, main table {len(schedule_rows)} player-game rows "
        f"({len(population)} player-roles x {games_per_role} real games each)."
    )
    print(f"Saved to {workbook_path}")
    return {
        "pace_sec3_range": (pace_sec3_first_row, pace_sec3_last_row),
        "main_range": (main_first_row, main_last_row),
        "n_rows": len(schedule_rows),
    }


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(
            'Usage: uv run python scripts/build_player_prop_projections.py '
            '"path/to/workbook.xlsx"'
        )
        sys.exit(1)
    build(sys.argv[1])
