"""
Builds "Kicking Index" -- Phase 3 of the multi-phase roadmap sketched in
claude_code_spec_rb_index.md's own header ("3. Kicking -- clean data, different metric
shape, lower game impact."). No separate written spec exists for this phase -- designed
directly (per explicit user instruction), reusing the Section 1-5 pattern from build_rb_
index.py / build_wr_te_index.py where it fits, and departing from it where kicking's shape
genuinely differs.

Where this departs from the other tabs, and why:

- Scores ONE role per team (K1) -- a team practically always has exactly one roster kicker,
  so there's no Starter/Backup binary and no multi-slot group either (unlike QB/RB or
  WR/TE). current_roster.py's new "PK": {1: "K1"} support (see that module).
- NO Section 2B (Individual Rookie Assumptions). Checked live before writing this: the
  api.fantasycalc.com dynasty-value dataset used for every other position's rookie
  crosswalk carries ZERO "K"/"PK" entries at all -- kickers simply aren't rostered in
  dynasty fantasy formats, so there is no real market signal to build a Section 2B from for
  this position. A zero-history kicker uses the flat Section 2 Rookie Baseline directly,
  same as every other tab's fallback path, just without an individual-assumption layer on
  top of it.
- Section 6 is a DIRECT quality adjustment, not a Replacement Value swap. QB/RB's
  Replacement Value answers "what happens if I bench the starter for the backup" -- there is
  no second roster kicker to swap in, so that question doesn't apply here. Instead, Section 6
  answers a different, equally well-defined question: "how many points better or worse than
  a league-average kicker is this team's kicker," computed directly from the Kicking Index
  Score's own distance from the 50-point league-average baseline (Model Assumptions C37) --
  no manual "K2 In" status toggle exists because there's nothing to toggle to.

Section 1: raw 3-year data per kicker-season (from kicking_stats.py)
Section 2: league average per season (ALL qualifying rows) + a flat Rookie Baseline
Section 3: per-kicker 3-Yr decay-weighted, regressed baseline -- population is the CURRENT-
           roster-identified K1 (current_roster.resolve_scored_population); missing years
           substitute the flat Section 2 Rookie Baseline directly (no Section 2B, see above)
Section 4: league average/std-dev of Section 3's projected baselines
Section 5: Z-scores, weighted composite, points-scale score (reuses QB Index's own
           Model Assumptions C37/C38 scale constants, same design choice as RB/WR-TE Index)
Section 6: direct Kicking Adjustment (pts) -- wired into Team Ratings as a new, unconditional
           additive column (no status toggle, see above)

Usage:
    uv run python scripts/build_kicking_index.py "C:\\path\\to\\NFL_Prediction_Model.xlsx"
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from nflverse_pull.current_roster import (  # noqa: E402
    compute_current_starters,
    fetch_depth_charts,
    resolve_scored_population,
)
from nflverse_pull.efficiency import fetch_pbp  # noqa: E402
from nflverse_pull.kicking_stats import compute_team_season_kicking_stats  # noqa: E402

HISTORICAL_YEARS = [2023, 2024, 2025]
CURRENT_ROSTER_YEAR = 2026
SHEET_NAME = "Kicking Index"
WR_TE_SHEET = "WR-TE Value Index"
RB_INDEX_SHEET = "RB Value Index"
QB_INDEX_SHEET = "QB Index"
TEAM_RATINGS_SHEET = "Team Ratings"

METRICS = [
    {"key": "foe", "col": "FG% Over Expected", "label": "FG%\nOver Expected", "fmt": "0.00%"},
    {"key": "fgpct", "col": "FG%", "label": "FG%", "fmt": "0.00%"},
    {"key": "xppct", "col": "XP%", "label": "XP%", "fmt": "0.00%"},
]

TEAM_ORDER = [
    "Buffalo Bills", "Miami Dolphins", "New England Patriots", "New York Jets",
    "Baltimore Ravens", "Cincinnati Bengals", "Cleveland Browns", "Pittsburgh Steelers",
    "Houston Texans", "Indianapolis Colts", "Jacksonville Jaguars", "Tennessee Titans",
    "Denver Broncos", "Kansas City Chiefs", "Las Vegas Raiders", "Los Angeles Chargers",
    "Dallas Cowboys", "New York Giants", "Philadelphia Eagles", "Washington Commanders",
    "Chicago Bears", "Detroit Lions", "Green Bay Packers", "Minnesota Vikings",
    "Atlanta Falcons", "Carolina Panthers", "New Orleans Saints", "Tampa Bay Buccaneers",
    "Arizona Cardinals", "Los Angeles Rams", "San Francisco 49ers", "Seattle Seahawks",
]

TITLE_FONT = Font(name="Arial", size=10, bold=True)
TITLE_FILL = PatternFill("solid", fgColor="FFD9E1F2")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_ALIGN = Alignment(wrap_text=True, horizontal="center", vertical="center")
INPUT_FONT = Font(name="Arial", size=10, color="FF0000FF")
FORMULA_FONT = Font(name="Arial", size=10, color="FF000000")
LINK_FONT = Font(name="Arial", size=10, color="FF008000")
NOTE_FONT = Font(name="Arial", size=9, color="FF808080")
ASSUMPTION_FILL = PatternFill("solid", fgColor="FFFFFF00")


def _section_title(ws, row: int, last_col: int, text: str) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL


def _header_row(ws, row: int, headers: list[str], height: float = 28) -> None:
    ws.row_dimensions[row].height = height
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN


def add_model_assumptions_weights(wb: openpyxl.Workbook) -> None:
    ws = wb["Model Assumptions"]
    title_row = 52
    ws.merge_cells(f"A{title_row}:D{title_row}")
    title = ws.cell(row=title_row, column=1, value=(
        "Kicking Index Weighting & Conversion (pts per std. dev.; reuses QB Index's "
        "points-scale constants C37/C38 -- see 'Kicking Index' tab)"
    ))
    title.font = HEADER_FONT
    title.fill = HEADER_FILL

    rows = [
        (53, "FG% Over Expected Weight (Kicking Index, pts per SD)", 0.5,
         "Distance-adjusted make rate is the most complete single kicking stat -- "
         "weighted highest, same logic as every other efficiency weighting in this "
         "model."),
        (54, "FG% Weight (Kicking Index, pts per SD)", 0.3,
         "Raw make rate correlates with FG% Over Expected but isn't distance-adjusted "
         "-- a team with an unusually short or long average attempt distance would "
         "skew this on its own, so it's weighted lower as a sanity check."),
        (55, "XP% Weight (Kicking Index, pts per SD)", 0.2,
         "Extra points are made at a very high, tightly clustered rate league-wide "
         "(essentially all NFL kickers convert 90%+) -- weighted lowest since it's "
         "the least discriminating of the three metrics."),
        (56, "Kicking Index Points-to-Game-Points Conversion", 0.05,
         "A starting guess, like every other coefficient in this model -- smaller "
         "than every other position's conversion constant (RB's is 0.08, QB's is "
         "0.15), reflecting kicking's genuinely lower game impact per the roadmap's "
         "own framing. This is a DIRECT quality adjustment, not a Replacement Value "
         "swap -- see this tab's own closing note."),
    ]
    for row, label, value, note in rows:
        ws.cell(row=row, column=2, value=label)
        c = ws.cell(row=row, column=3, value=value)
        c.font = INPUT_FONT
        c.fill = ASSUMPTION_FILL
        c.number_format = "0.00"
        n = ws.cell(row=row, column=4, value=note)
        n.font = NOTE_FONT
        n.alignment = Alignment(wrap_text=True, vertical="top")


def _pull_data() -> dict:
    print(f"Pulling {HISTORICAL_YEARS} play-by-play data for kicking stats...")
    pbp = fetch_pbp(HISTORICAL_YEARS)
    season_stats = compute_team_season_kicking_stats(pbp)

    print(f"Pulling {CURRENT_ROSTER_YEAR} depth charts for current-roster K1 population...")
    depth_charts = fetch_depth_charts([CURRENT_ROSTER_YEAR])
    current_starters = compute_current_starters(depth_charts)
    empty_overrides = pd.DataFrame(
        columns=["Team", "Manual Starter Override", "Manual Backup Override"]
    )
    population = resolve_scored_population(current_starters, empty_overrides, "PK")
    print(f"{len(population)} current-roster K1s identified.")

    return {"season_stats": season_stats, "population": population}


def build(workbook_path: str) -> dict:
    data = _pull_data()
    season_stats = data["season_stats"]
    population = data["population"]
    n_kickers = len(population)

    zero_history = population[~population["Player ID"].isin(set(season_stats["Player ID"]))]
    if len(zero_history):
        print(f"{len(zero_history)} current K1(s) with zero qualifying history: "
              f"{list(zero_history['Player Name'])} -- no dynasty fantasy market exists for "
              "kickers (verified live: 0 K/PK entries in the fantasycalc.com dataset), so "
              "these use the flat Section 2 Rookie Baseline directly, no Section 2B.")
    else:
        print("No current K1 has zero qualifying pbp history.")

    wb = openpyxl.load_workbook(workbook_path)
    add_model_assumptions_weights(wb)

    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    for candidate in (WR_TE_SHEET, RB_INDEX_SHEET, QB_INDEX_SHEET):
        if candidate in wb.sheetnames:
            insert_after = candidate
            break
    else:
        insert_after = wb.sheetnames[0]
    ws = wb.create_sheet(SHEET_NAME, index=wb.sheetnames.index(insert_after) + 1)
    ws.column_dimensions["A"].width = 18.0
    ws.column_dimensions["C"].width = 20.0

    ws.merge_cells("A1:H1")
    t = ws.cell(row=1, column=1, value=(
        "Kicking Index -- Multi-Year Decay-Weighted Kicking Rating (FG% Over Expected, "
        "distance-adjusted; FG%; XP%). Scores ONE role per team (K1) -- see the closing "
        "note below for why Section 6 is a direct quality adjustment, not a Replacement "
        "Value swap like QB/RB Index."
    ))
    t.font = Font(name="Arial", size=12, bold=True)

    # ==== Section 1: Raw 3-year data per kicker-season =====================================
    sec1_first_row = 5
    sec1_last_row = sec1_first_row + len(season_stats) - 1
    _section_title(
        ws, 3, 8, "Section 1 \u2014 Raw 3-Year Data per Kicker-Season (from nflverse pbp)"
    )
    _header_row(
        ws, 4,
        ["Player Name", "Player ID", "Team", "Season", "FG Attempts", "FG% Over Expected",
         "FG%", "XP%", "Is Rookie Season"],
    )
    for i, r in enumerate(season_stats.to_dict("records")):
        row = sec1_first_row + i
        values = [
            r["Player Name"], r["Player ID"], r["Team"], int(r["Season"]),
            int(r["FG Attempts"]), float(r["FG% Over Expected"]), float(r["FG%"]),
            float(r["XP%"]), bool(r["Is Rookie Season"]),
        ]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.font = INPUT_FONT
            if col in (6, 7, 8):
                cell.number_format = "0.00%"

    id_range = f"$B${sec1_first_row}:$B${sec1_last_row}"
    season_range = f"$D${sec1_first_row}:$D${sec1_last_row}"
    rookie_range = f"$I${sec1_first_row}:$I${sec1_last_row}"
    sec1_col_of = {"foe": "F", "fgpct": "G", "xppct": "H"}
    metric_ranges = {
        m["key"]: (
            f"${sec1_col_of[m['key']]}${sec1_first_row}:${sec1_col_of[m['key']]}${sec1_last_row}"
        )
        for m in METRICS
    }

    # ==== Section 2: League average per season (ALL qualifying rows) + flat Rookie Baseline
    sec2_title_row = sec1_last_row + 2
    sec2_header_row = sec2_title_row + 1
    season_rows = {yr: sec2_header_row + 1 + i for i, yr in enumerate(HISTORICAL_YEARS)}
    rookie_baseline_row = sec2_header_row + 1 + len(HISTORICAL_YEARS)
    rookie_count_row = rookie_baseline_row + 1

    _section_title(
        ws, sec2_title_row, 4,
        "Section 2 \u2014 League Average per Season (ALL qualifying rows -- only one "
        "role/team exists here, so there's no Starter/Backup-style filter to apply) and "
        "flat Rookie Baseline",
    )
    _header_row(ws, sec2_header_row, ["Season / Stat", *[m["label"] for m in METRICS]], height=20)

    for yr in HISTORICAL_YEARS:
        row = season_rows[yr]
        c = ws.cell(row=row, column=1, value=yr)
        c.font = INPUT_FONT
        for j, m in enumerate(METRICS):
            formula = f"=AVERAGEIF({season_range},{yr},{metric_ranges[m['key']]})"
            cell = ws.cell(row=row, column=2 + j, value=formula)
            cell.font = FORMULA_FONT
            cell.number_format = m["fmt"]

    lbl = ws.cell(row=rookie_baseline_row, column=1, value="Rookie Baseline (all years, flat)")
    lbl.font = FORMULA_FONT
    for j, m in enumerate(METRICS):
        formula = f"=AVERAGEIF({rookie_range},TRUE,{metric_ranges[m['key']]})"
        cell = ws.cell(row=rookie_baseline_row, column=2 + j, value=formula)
        cell.font = FORMULA_FONT
        cell.number_format = m["fmt"]

    lbl2 = ws.cell(
        row=rookie_count_row, column=1,
        value="Rookie Baseline Sample Size (qualifying rookie seasons) -- small-sample "
              "estimate; recomputes automatically as more seasons are added",
    )
    lbl2.font = NOTE_FONT
    lbl2.alignment = Alignment(wrap_text=True)
    for j in range(len(METRICS)):
        cell = ws.cell(row=rookie_count_row, column=2 + j, value=f"=COUNTIF({rookie_range},TRUE)")
        cell.font = FORMULA_FONT
        cell.number_format = "0"

    flat_rookie_cell = {
        m["key"]: f"${get_column_letter(2 + j)}${rookie_baseline_row}"
        for j, m in enumerate(METRICS)
    }
    season_avg_col = {m["key"]: get_column_letter(2 + j) for j, m in enumerate(METRICS)}

    # ==== Section 3: Per-kicker decay-weighted, regressed baseline =========================
    sec3_title_row = rookie_count_row + 2
    sec3_header_row = sec3_title_row + 1
    sec3_first_row = sec3_header_row + 1
    sec3_last_row = sec3_first_row + n_kickers - 1
    sec3_last_col = 5 + len(METRICS) * 7

    _section_title(
        ws, sec3_title_row, sec3_last_col,
        "Section 3 \u2014 Per-Kicker 3-Yr Decay-Weighted, Regressed Baseline. Population "
        "is the CURRENT-roster-identified K1 (current_roster.py's live 2026 depth-chart "
        "pull). Missing years substitute the flat Section 2 Rookie Baseline directly -- no "
        "Section 2B exists for this tab (no dynasty fantasy market for kickers, see the "
        "closing note).",
    )
    headers = ["Player Name", "Player ID", "Team", "Role", "Years of\nReal History"]
    metric_block_start_col: dict[str, int] = {}
    col_cursor = 6
    for m in METRICS:
        metric_block_start_col[m["key"]] = col_cursor
        headers += [
            f"{m['label']} Y-1", f"{m['label']} Y-2", f"{m['label']} Y-3",
            f"Weighted\n{m['label']} Avg\n(3-Yr decay)", f"Team History\n{m['label']}",
            f"League Baseline\n{m['label']} (Y-1)", f"Projected 3-Yr\n{m['label']} Baseline",
        ]
        col_cursor += 7
    _header_row(ws, sec3_header_row, headers)

    for i, p in enumerate(population.to_dict("records")):
        row = sec3_first_row + i
        ws.cell(row=row, column=1, value=p["Player Name"]).font = FORMULA_FONT
        ws.cell(row=row, column=2, value=p["Player ID"]).font = FORMULA_FONT
        ws.cell(row=row, column=3, value=p["Team"]).font = FORMULA_FONT
        ws.cell(row=row, column=4, value=p["Role"]).font = FORMULA_FONT

        years_hist_terms = "+".join(
            f"--(COUNTIFS({id_range},$B{row},{season_range},'Model Assumptions'!$C$18-{k})>0)"
            for k in (1, 2, 3)
        )
        yh = ws.cell(row=row, column=5, value=f"={years_hist_terms}")
        yh.font = FORMULA_FONT
        yh.number_format = "0"

        for m in METRICS:
            base = metric_block_start_col[m["key"]]
            y1, y2, y3, wavg, th, lb, pb = (get_column_letter(base + k) for k in range(7))
            mrange = metric_ranges[m["key"]]
            rbcell = flat_rookie_cell[m["key"]]

            def _ysub(offset: int, mrange=mrange, rbcell=rbcell) -> str:
                return (
                    f"=IF(COUNTIFS({id_range},$B{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset})=0,{rbcell},"
                    f"SUMIFS({mrange},{id_range},$B{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset}))"
                )

            f_y1 = ws.cell(row=row, column=base, value=_ysub(1))
            f_y2 = ws.cell(row=row, column=base + 1, value=_ysub(2))
            f_y3 = ws.cell(row=row, column=base + 2, value=_ysub(3))
            f_wavg = ws.cell(row=row, column=base + 3, value=(
                f"=({y1}{row}*1+{y2}{row}*'Model Assumptions'!$C$20+"
                f"{y3}{row}*('Model Assumptions'!$C$20^2))/"
                f"(1+'Model Assumptions'!$C$20+'Model Assumptions'!$C$20^2)"
            ))
            f_th = ws.cell(row=row, column=base + 4, value=(
                f"={y1}{row}*'Model Assumptions'!$C$22+{wavg}{row}*(1-'Model Assumptions'!$C$22)"
            ))
            sec2_col = season_avg_col[m["key"]]
            f_lb = ws.cell(row=row, column=base + 5, value=(
                f"=INDEX(${sec2_col}${season_rows[HISTORICAL_YEARS[0]]}:"
                f"${sec2_col}${season_rows[HISTORICAL_YEARS[-1]]},"
                f"MATCH('Model Assumptions'!$C$18-1,$A${season_rows[HISTORICAL_YEARS[0]]}:"
                f"$A${season_rows[HISTORICAL_YEARS[-1]]},0))"
            ))
            f_pb = ws.cell(row=row, column=base + 6, value=(
                f"={th}{row}*'Model Assumptions'!$C$21+{lb}{row}*(1-'Model Assumptions'!$C$21)"
            ))
            for cell in (f_y1, f_y2, f_y3, f_wavg, f_th, f_lb, f_pb):
                cell.font = FORMULA_FONT
                cell.number_format = m["fmt"]

    # ==== Section 4: League average/std-dev of Section 3's Projected Baselines =============
    sec4_title_row = sec3_last_row + 2
    sec4_header_row = sec4_title_row + 1
    avg_row, std_row = sec4_header_row + 1, sec4_header_row + 2

    _section_title(
        ws, sec4_title_row, 4, "Section 4 \u2014 League Average & Std. Dev. of the 3-Yr Baselines"
    )
    _header_row(ws, sec4_header_row, ["Stat", *[m["label"] for m in METRICS]], height=18)

    ws.cell(row=avg_row, column=1, value="League Average").font = FORMULA_FONT
    ws.cell(row=std_row, column=1, value="League Std. Dev.").font = FORMULA_FONT

    proj_baseline_col = {
        m["key"]: get_column_letter(metric_block_start_col[m["key"]] + 6) for m in METRICS
    }
    for j, m in enumerate(METRICS):
        pcol = proj_baseline_col[m["key"]]
        pb_range = f"{pcol}${sec3_first_row}:{pcol}${sec3_last_row}"
        a = ws.cell(row=avg_row, column=2 + j, value=f"=AVERAGE({pb_range})")
        s = ws.cell(row=std_row, column=2 + j, value=f"=STDEVP({pb_range})")
        for cell in (a, s):
            cell.font = FORMULA_FONT
            cell.number_format = m["fmt"]

    # ==== Section 5: Z-scores, weighted composite, points-scale score ======================
    sec5_title_row = std_row + 2
    sec5_header_row = sec5_title_row + 1
    sec5_first_row = sec5_header_row + 1
    sec5_last_row = sec5_first_row + n_kickers - 1

    _section_title(
        ws, sec5_title_row, 10,
        "Section 5 \u2014 Z-Scores and Kicking Index Score (all three metrics are \"higher "
        "is better\" -- no sign-flip needed. Baseline/points-per-SD reuse QB Index's own "
        "Model Assumptions cells C37/C38.)",
    )
    _header_row(
        ws, sec5_header_row,
        ["Player Name", "Player ID", "Team", "Role", *[f"{m['label']}\nZ" for m in METRICS],
         "Weighted\nZ-Score Sum", "Kicking Index\nScore (Points)", "Years of\nReal History"],
    )

    avg_cell_ref = {
        m["key"]: f"${get_column_letter(2 + j)}${avg_row}" for j, m in enumerate(METRICS)
    }
    std_cell_ref = {
        m["key"]: f"${get_column_letter(2 + j)}${std_row}" for j, m in enumerate(METRICS)
    }
    weight_cells = {"foe": "$C$53", "fgpct": "$C$54", "xppct": "$C$55"}

    for i in range(n_kickers):
        sec3_row = sec3_first_row + i
        row = sec5_first_row + i
        for col, src_col in ((1, "A"), (2, "B"), (3, "C"), (4, "D")):
            f = ws.cell(row=row, column=col, value=f"={src_col}{sec3_row}")
            f.font = FORMULA_FONT

        z_cols = []
        for j, m in enumerate(METRICS):
            pcol = proj_baseline_col[m["key"]]
            formula = f"=({pcol}{sec3_row}-{avg_cell_ref[m['key']]})/{std_cell_ref[m['key']]}"
            col = 5 + j
            cell = ws.cell(row=row, column=col, value=formula)
            cell.font = FORMULA_FONT
            cell.number_format = "0.00"
            z_cols.append(get_column_letter(col))

        weighted_terms = " + ".join(
            f"{z_cols[j]}{row}*'Model Assumptions'!{weight_cells[m['key']]}"
            for j, m in enumerate(METRICS)
        )
        wz_col = 5 + len(METRICS)
        wz = ws.cell(row=row, column=wz_col, value=f"={weighted_terms}")
        wz.font = FORMULA_FONT
        wz.number_format = "0.00"

        score = ws.cell(row=row, column=wz_col + 1, value=(
            f"='Model Assumptions'!$C$37+{get_column_letter(wz_col)}{row}*'Model Assumptions'!$C$38"
        ))
        score.font = FORMULA_FONT
        score.number_format = "0.0;(0.0)"

        yh_link = ws.cell(row=row, column=wz_col + 2, value=f"=E{sec3_row}")
        yh_link.font = FORMULA_FONT
        yh_link.number_format = "0"

    # ==== Section 6: Direct Kicking Adjustment (NOT a Replacement Value swap) ==============
    sec6_title_row = sec5_last_row + 2
    sec6_header_row = sec6_title_row + 1
    sec6_first_row = sec6_header_row + 1
    sec6_last_row = sec6_first_row + len(TEAM_ORDER) - 1

    _section_title(
        ws, sec6_title_row, 5,
        "Section 6 \u2014 Kicking Adjustment (pts) -- a DIRECT quality measure (Kicking "
        "Index Score minus the league-average baseline of 50, converted to game points), "
        "NOT a Replacement Value swap like QB/RB Index. There is no second roster kicker "
        "to swap in, so 'Backup In' doesn't apply here -- this is unconditional, applied "
        "to every team every time.",
    )
    _header_row(ws, sec6_header_row, [
        "Team", "K1 Name", "K1 Kicking\nIndex Score", "Kicking Adjustment\n(Index Points)",
        "Kicking Adjustment\n(Game Points)",
    ])

    # Section 5 columns: A=Name, B=ID, C=Team, D=Role, E-G=the 3 metric Z-scores,
    # H=Weighted Z-Score Sum, I=Kicking Index Score, J=Years of Real History -- the Score
    # column is I, NOT H (H is the intermediate weighted-Z-sum, not the points-scale score).
    score_range = f"$I${sec5_first_row}:$I${sec5_last_row}"
    name_range = f"$A${sec5_first_row}:$A${sec5_last_row}"
    team_range = f"$C${sec5_first_row}:$C${sec5_last_row}"

    for i, team in enumerate(TEAM_ORDER):
        row = sec6_first_row + i
        t = ws.cell(row=row, column=1, value=team)
        t.font = INPUT_FONT

        name = ws.cell(row=row, column=2, value=(
            f'=IFERROR(INDEX({name_range},MATCH("{team}",{team_range},0)),"")'
        ))
        score = ws.cell(row=row, column=3, value=(
            f'=IFERROR(INDEX({score_range},MATCH("{team}",{team_range},0)),"")'
        ))
        name.font = FORMULA_FONT
        score.font = FORMULA_FONT
        score.number_format = "0.0;(0.0)"

        adj_index = ws.cell(row=row, column=4, value=(
            f'=IF(C{row}="","",C{row}-\'Model Assumptions\'!$C$37)'
        ))
        adj_game = ws.cell(row=row, column=5, value=(
            f'=IF(D{row}="","",D{row}*\'Model Assumptions\'!$C$56)'
        ))
        adj_index.font = FORMULA_FONT
        adj_game.font = FORMULA_FONT
        adj_index.number_format = "0.0;(0.0)"
        adj_game.number_format = "0.00;(0.00)"

    note_row = sec6_last_row + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
    note = ws.cell(row=note_row, column=1, value=(
        "Kicking Adjustment (Index Points) = K1 Kicking Index Score - the league-average "
        "baseline (Model Assumptions C37, shared with QB Index); Kicking Adjustment (Game "
        "Points) applies the Kicking Index Points-to-Game-Points Conversion (Model "
        "Assumptions C56 -- its own constant, distinct from every other position's). Both "
        "recalculate live from Section 5. A team missing a current K1 entirely shows "
        "blank rather than a misleading 0. THIS IS NOT A REPLACEMENT VALUE SWAP: unlike "
        "QB/RB Index, there is no 'K2 In' status toggle and no second roster kicker being "
        "compared against -- a team practically always has exactly one kicker, so this "
        "section answers 'how many points better/worse than league-average is this team's "
        "kicker,' applied unconditionally to Net Power Rating for every team, not "
        "conditionally on a manual toggle. There is also NO SECTION 2B (Individual Rookie "
        "Assumptions) on this tab -- checked live before building this: the "
        "api.fantasycalc.com dynasty-value dataset used for every other position's rookie "
        "crosswalk carries zero K/PK entries at all (kickers aren't rostered in dynasty "
        "fantasy formats), so a zero-history kicker uses the flat Section 2 Rookie Baseline "
        "directly. The 2026 depth-chart snapshot this tab's K1 population is built from was "
        "pulled BEFORE final 53-man roster cuts -- same caveat as every other current-"
        "roster-driven tab in this workbook."
    ))
    note.font = NOTE_FONT
    note.alignment = Alignment(wrap_text=True, vertical="top")

    # ==== Wire into Team Ratings (new, unconditional additive column) ======================
    tr = wb[TEAM_RATINGS_SHEET]
    tr.cell(row=2, column=19, value="Kicking\nAdjustment (pts)").font = HEADER_FONT
    tr.cell(row=2, column=19).fill = HEADER_FILL
    tr.cell(row=2, column=19).alignment = HEADER_ALIGN

    adj_game_range = f"'{SHEET_NAME}'!$E${sec6_first_row}:$E${sec6_last_row}"
    adj_team_range = f"'{SHEET_NAME}'!$A${sec6_first_row}:$A${sec6_last_row}"

    for row in range(3, 3 + len(TEAM_ORDER)):
        adj = tr.cell(row=row, column=19, value=(
            f"=IFERROR(INDEX({adj_game_range},MATCH(A{row},{adj_team_range},0)),0)"
        ))
        adj.font = LINK_FONT
        adj.number_format = "0.00;(0.00)"

        # Net Power Rating (N) now also includes the Kicking Adjustment, additive alongside
        # the QB (P) and RB (R) adjustments -- none of the three disturb each other.
        net = tr.cell(row=row, column=14, value=(
            f"=J{row}-K{row}+L{row}+M{row}+P{row}+R{row}+S{row}"
        ))
        net.font = FORMULA_FONT
        net.number_format = "0.0;(0.0)"

    note_row_tr = 3 + len(TEAM_ORDER) + 2
    tr.merge_cells(start_row=note_row_tr, start_column=1, end_row=note_row_tr, end_column=19)
    tr_note = tr.cell(row=note_row_tr, column=1, value=(
        "Kicking Adjustment (S) is unconditional -- unlike QB Replacement Value Adj (P) and "
        "RB Replacement Value Adj (R), there's no manual status toggle to gate it, since "
        "there's no second roster kicker to switch to. It pulls directly from 'Kicking "
        "Index' Section 6 (that team's K1 Kicking Index Score minus the league-average "
        "baseline, converted to game points) and applies to Net Power Rating (N) for every "
        "team, every time, alongside whatever QB/RB adjustments are also active."
    ))
    tr_note.font = NOTE_FONT
    tr_note.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(workbook_path)
    print(
        f"Built '{SHEET_NAME}': Section 1 {len(season_stats)} rows, Section 3/5/6 "
        f"{n_kickers} kickers / {len(TEAM_ORDER)} teams. Wired into "
        f"'{TEAM_RATINGS_SHEET}' (col S)."
    )
    print(f"Saved to {workbook_path}")
    return {
        "sec1_rows": len(season_stats), "n_kickers": n_kickers,
        "sec3_range": (sec3_first_row, sec3_last_row),
        "sec5_range": (sec5_first_row, sec5_last_row),
        "sec6_range": (sec6_first_row, sec6_last_row),
    }


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print('Usage: uv run python scripts/build_kicking_index.py "path/to/workbook.xlsx"')
        sys.exit(1)
    build(sys.argv[1])
