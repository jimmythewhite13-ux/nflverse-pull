"""
Builds "Explosive Play Matchup" -- claude_code_spec_explosive_play_engine.md Parts A/B.

DESIGN TRAP AVOIDED (the spec's own explicit warning): 15+/20+/30+/40+ yard completions are
NOT independent categories -- a 40-yard completion is also a 30+/20+/15+ yard completion.
This tab uses ONE scored threshold per phase (15+ passing, 10+ rushing -- the SAME
EXPLOSIVE_PASS_YARDS/EXPLOSIVE_RUN_YARDS constants Pass/Run Defense Matchup already use, not
a new inconsistent threshold), with 20+/30+/40+ (passing) and 15+/20+ (rushing) carried as
CONTEXT-ONLY reference columns in Section 1 -- never Z-scored, never weighted, anywhere on
this tab (see efficiency.compute_team_season_explosive_tiers's own module docstring).

Part A finding, reported here rather than silently rebuilding: the Defensive Matchup
Engine's own compute_team_season_matchup_metrics() ALREADY computes real team-level
"Explosive Pass Rate (Off)" / "Explosive Run Rate (Off)" -- confirmed by reading that
function's own output columns -- but no tab built on top of it ever surfaced those two
columns; only the "(Def)"/Allowed side landed on Pass/Run Defense Matchup. This tab is where
the offensive side finally surfaces, referencing the SAME already-computed function rather
than recomputing it. (Separately, QB Index's own player-level Explosive Pass Rate now exists
too, via claude_code_spec_qb_environment_model.md's own Part A on the "QB Environment Model"
tab -- a different, complementary, player-level number; that one is NOT touched or
duplicated here, this tab is TEAM-level, matching Pass/Run Defense Matchup's own shape.)

Part B -- two genuinely distinct NEW defensive metrics, confirmed different by their real
formulas (not the same number twice under different names -- see efficiency.py's own
compute_team_season_deep_pass_yac_metrics docstring for the full reasoning):
  - Deep Pass Completion Rate Allowed: completions / attempts on throws with real
    air_yards >= 20 -- "getting beaten deep," independent of how many yards a short pass
    turned into after the catch.
  - YAC Allowed: real mean yards_after_catch across all real completions allowed -- a
    genuine, free proxy for "not making a play in space." MISSED TACKLES THEMSELVES ARE
    CONFIRMED NOT BUILDABLE ANYWHERE IN THIS PROJECT -- no free public play-by-play has
    charting/tracking data for that; NOT attempted here or anywhere else.
  - Explosive Pass/Run Rate Allowed: REFERENCED live from Pass/Run Defense Matchup's own
    already-computed Section 5 Z-scores (cols F/E there respectively) -- NOT recomputed.

Section 1: raw 3-year team data (4 scored metrics: Explosive Pass Rate Off, Explosive Run
           Rate Off, Deep Pass Completion Rate Allowed, YAC Allowed; 5 context-only
           reference tiers: Pass 20+/30+/40+, Run 15+/20+, all Off)
Section 2: league average per season (4 scored metrics)
Section 3: per-team 3-Yr decay-weighted, regressed baseline (4 scored metrics)
Section 4: league average/std-dev of Section 3's projected baselines
Section 5: Z-scores (4 own metrics, Deep Pass/YAC sign-flipped since lower=better defense) +
           2 REFERENCED Z's (Explosive Pass/Run Rate Allowed, from Pass/Run Defense Matchup)
           -> Pass Explosive-Prevention Composite Z (weighted: Explosive Pass Rate Allowed +
           Deep Pass Completion Rate Allowed + YAC Allowed) and Run Explosive-Prevention Z
           (= Explosive Run Rate Allowed directly -- no run-side equivalent to Deep Pass/YAC
           exists, a documented asymmetry, not padded to force parity with the pass side).

NO Team Ratings wiring -- same reasoning as Pass/Run Defense Matchup: a matchup-specific
composite, not a static team-quality number. Feeds Part C's differential (build_explosive_
play_matchup_wiring.py) directly from Section 5.

NOT YET VALIDATED against real outcomes -- same disclaimer as every tab in the Defensive
Matchup Engine family.

Usage:
    uv run python scripts/build_explosive_play_matchup.py "C:\\path\\to\\NFL_Prediction_Model.xlsx"
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from nflverse_pull.efficiency import (  # noqa: E402
    compute_team_season_deep_pass_yac_metrics,
    compute_team_season_explosive_tiers,
    compute_team_season_matchup_metrics,
    fetch_pbp,
)

HISTORICAL_YEARS = [2023, 2024, 2025]
SHEET_NAME = "Explosive Play Matchup"
PASS_DEFENSE_SHEET = "Pass Defense Matchup"
RUN_DEFENSE_SHEET = "Run Defense Matchup"

# Pass/Run Defense Matchup's own real Section 5 layout (verified against each tab's own
# build script): 32 teams, rows 150-181. Explosive Pass Rate Allowed Z lives at col F on
# Pass Defense Matchup; Explosive Run Rate Allowed Z lives at col E on Run Defense Matchup.
PASS_DEF_SEC5_RANGE = (150, 181)
PASS_DEF_EXP_Z_COL, PASS_DEF_TEAM_COL = "F", "A"
RUN_DEF_SEC5_RANGE = (150, 181)
RUN_DEF_EXP_Z_COL, RUN_DEF_TEAM_COL = "E", "A"

METRICS = [
    {"key": "exp_pass_off", "col": "Explosive Pass Rate (Off)",
     "label": "Explosive Pass\nRate (Off)", "fmt": "0.00%", "invert": False},
    {"key": "exp_run_off", "col": "Explosive Run Rate (Off)",
     "label": "Explosive Run\nRate (Off)", "fmt": "0.00%", "invert": False},
    {"key": "deep_allowed", "col": "Deep Pass Completion Rate Allowed (Def)",
     "label": "Deep Pass Comp.\nRate Allowed", "fmt": "0.00%", "invert": True},
    {"key": "yac_allowed", "col": "YAC Allowed (Def)", "label": "YAC\nAllowed",
     "fmt": "0.00", "invert": True},
]
TIER_COLS = [
    "Pass 20+ Rate (Off)", "Pass 30+ Rate (Off)", "Pass 40+ Rate (Off)",
    "Run 15+ Rate (Off)", "Run 20+ Rate (Off)",
]
TIER_LABELS = [
    "Pass 20+\n(ref only)", "Pass 30+\n(ref only)", "Pass 40+\n(ref only)",
    "Run 15+\n(ref only)", "Run 20+\n(ref only)",
]

TEAM_ORDER = [
    "Buffalo Bills", "Miami Dolphins", "New England Patriots", "New York Jets",
    "Baltimore Ravens", "Cincinnati Bengals", "Cleveland Browns", "Pittsburgh Steelers",
    "Houston Texans", "Indianapolis Colts", "Jacksonville Jaguars", "Tennessee Titans",
    "Denver Broncos", "Kansas City Chiefs", "Las Vegas Raiders", "Los Angeles Chargers",
    "Dallas Cowboys", "New York Giants", "Philadelphia Eagles", "Washington Commanders",
    "Chicago Bears", "Detroit Lions", "Green Bay Packers", "Minnesota Vikings",
    "Atlanta Falcons", "Carolina Panthers", "New Orleans Saints", "Tampa Bay Buccaneers",
    "Arizona Cardinals", "Los Angeles Rams", "San Francisco 49ers", "Seattle Seahawks",
]

TITLE_FONT = Font(name="Arial", size=10, bold=True)
TITLE_FILL = PatternFill("solid", fgColor="FFD9E1F2")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_ALIGN = Alignment(wrap_text=True, horizontal="center", vertical="center")
INPUT_FONT = Font(name="Arial", size=10, color="FF0000FF")
FORMULA_FONT = Font(name="Arial", size=10, color="FF000000")
LINK_FONT = Font(name="Arial", size=10, color="FF008000")
NOTE_FONT = Font(name="Arial", size=9, color="FF808080")
ASSUMPTION_FILL = PatternFill("solid", fgColor="FFFFFF00")


def _section_title(ws, row: int, last_col: int, text: str) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL


def _header_row(ws, row: int, headers: list[str], height: float = 28) -> None:
    ws.row_dimensions[row].height = height
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN


def add_model_assumptions_weights(wb: openpyxl.Workbook) -> None:
    ws = wb["Model Assumptions"]
    title_row = 135
    ws.merge_cells(f"A{title_row}:D{title_row}")
    title = ws.cell(row=title_row, column=1, value=(
        "Explosive Play Matchup -- Pass Explosive-Prevention Composite Weighting (Z-score "
        "composite, no points-scale conversion -- see 'Explosive Play Matchup' tab)"
    ))
    title.font = HEADER_FONT
    title.fill = HEADER_FILL

    rows = [
        (136, "Explosive Pass Rate Allowed Weight (Pass Prevention Composite, ref)", 0.4,
         "Real, already-established defensive metric (Pass Defense Matchup's own "
         "Section 5, referenced not recomputed) -- weighted highest as the most direct "
         "single real signal for 'defense giving up big plays through the air.'"),
        (137, "Deep Pass Completion Rate Allowed Weight (Pass Prevention Composite)", 0.35,
         "A genuinely distinct real signal ('getting beaten deep') from Explosive Pass "
         "Rate Allowed -- weighted close behind it, not treated as a lesser duplicate."),
        (138, "YAC Allowed Weight (Pass Prevention Composite)", 0.25,
         "A real proxy for 'not making plays in space,' weighted lowest of the three -- "
         "further from pure coverage/pass-rush quality than the other two (real missed "
         "tackles would be more direct, but that data doesn't exist anywhere free)."),
    ]
    for row, label, value, note in rows:
        ws.cell(row=row, column=2, value=label)
        c = ws.cell(row=row, column=3, value=value)
        c.font = INPUT_FONT
        c.fill = ASSUMPTION_FILL
        c.number_format = "0.00"
        n = ws.cell(row=row, column=4, value=note)
        n.font = NOTE_FONT
        n.alignment = Alignment(wrap_text=True, vertical="top")


def _pull_data() -> pd.DataFrame:
    print(f"Pulling {HISTORICAL_YEARS} play-by-play data for Explosive Play Matchup...")
    pbp = fetch_pbp(HISTORICAL_YEARS)
    matchup = compute_team_season_matchup_metrics(pbp)
    deep_yac = compute_team_season_deep_pass_yac_metrics(pbp)
    tiers = compute_team_season_explosive_tiers(pbp)

    out = matchup[["Team", "Season", "Explosive Pass Rate (Off)", "Explosive Run Rate (Off)"]]
    out = out.merge(deep_yac, on=["Team", "Season"], how="left")
    out = out.merge(tiers, on=["Team", "Season"], how="left")
    return out


def build(workbook_path: str) -> dict:
    season_stats = _pull_data()

    wb = openpyxl.load_workbook(workbook_path)
    for required in (PASS_DEFENSE_SHEET, RUN_DEFENSE_SHEET):
        if required not in wb.sheetnames:
            raise ValueError(
                f"'{required}' not found -- run its own build script first (this tab "
                "references its Section 5 Explosive Rate Allowed Z directly)."
            )

    add_model_assumptions_weights(wb)

    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    for candidate in (RUN_DEFENSE_SHEET, PASS_DEFENSE_SHEET):
        if candidate in wb.sheetnames:
            insert_after = candidate
            break
    else:
        insert_after = wb.sheetnames[0]
    ws = wb.create_sheet(SHEET_NAME, index=wb.sheetnames.index(insert_after) + 1)
    ws.column_dimensions["A"].width = 20.0
    ws.column_dimensions["C"].width = 20.0

    ws.merge_cells("A1:K1")
    t = ws.cell(row=1, column=1, value=(
        "Explosive Play Matchup -- Multi-Year Decay-Weighted Explosive-Play Offense/"
        "Prevention Rating (real nflverse pbp; ONE scored threshold per phase -- 15+ "
        "passing, 10+ rushing, same as Pass/Run Defense Matchup -- with 20+/30+/40+ "
        "passing and 15+/20+ rushing as CONTEXT-ONLY reference tiers, never weighted "
        "anywhere). NOT YET VALIDATED against real outcomes -- see the closing note."
    ))
    t.font = Font(name="Arial", size=12, bold=True)

    # ==== Section 1: Raw 3-year team data (4 scored + 5 context-only reference tiers) =====
    sec1_first_row = 5
    sec1_last_row = sec1_first_row + len(season_stats) - 1
    _section_title(
        ws, 3, 11,
        "Section 1 — Raw 3-Year TEAM-Level Data (real nflverse pbp -- see this tab's "
        "opening note). Pass 20+/30+/40+ and Run 15+/20+ (G-K) are CONTEXT ONLY -- a "
        "40-yard completion IS also a 30+/20+/15+ yard completion, so these are NEVER "
        "part of any Z-score or weighted composite on this tab.",
    )
    _header_row(
        ws, 4,
        ["Team", "Season", *[m["label"] for m in METRICS], *TIER_LABELS],
    )
    for i, r in enumerate(season_stats.to_dict("records")):
        row = sec1_first_row + i
        values = [r["Team"], int(r["Season"])]
        for m in METRICS:
            v = r.get(m["col"])
            values.append(float(v) if pd.notna(v) else None)
        for tc in TIER_COLS:
            v = r.get(tc)
            values.append(float(v) if pd.notna(v) else None)
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.font = INPUT_FONT
        for j, m in enumerate(METRICS):
            ws.cell(row=row, column=3 + j).number_format = m["fmt"]
        for j in range(len(TIER_COLS)):
            ws.cell(row=row, column=3 + len(METRICS) + j).number_format = "0.00%"

    team_range = f"$A${sec1_first_row}:$A${sec1_last_row}"
    season_range = f"$B${sec1_first_row}:$B${sec1_last_row}"
    sec1_col_of = {m["key"]: get_column_letter(3 + j) for j, m in enumerate(METRICS)}
    metric_ranges = {
        m["key"]: (
            f"${sec1_col_of[m['key']]}${sec1_first_row}:${sec1_col_of[m['key']]}${sec1_last_row}"
        )
        for m in METRICS
    }

    # ==== Section 2: League average per season ==============================================
    sec2_title_row = sec1_last_row + 2
    sec2_header_row = sec2_title_row + 1
    season_rows = {yr: sec2_header_row + 1 + i for i, yr in enumerate(HISTORICAL_YEARS)}

    _section_title(
        ws, sec2_title_row, 5,
        "Section 2 — League Average per Season (simple average across all 32 teams)",
    )
    _header_row(ws, sec2_header_row, ["Season", *[m["label"] for m in METRICS]], height=20)

    for yr in HISTORICAL_YEARS:
        row = season_rows[yr]
        c = ws.cell(row=row, column=1, value=yr)
        c.font = INPUT_FONT
        for j, m in enumerate(METRICS):
            formula = f"=AVERAGEIF({season_range},{yr},{metric_ranges[m['key']]})"
            cell = ws.cell(row=row, column=2 + j, value=formula)
            cell.font = FORMULA_FONT
            cell.number_format = m["fmt"]

    season_avg_col = {m["key"]: get_column_letter(2 + j) for j, m in enumerate(METRICS)}

    # ==== Section 3: Per-team decay-weighted, regressed baseline ===========================
    sec3_title_row = season_rows[HISTORICAL_YEARS[-1]] + 2
    sec3_header_row = sec3_title_row + 1
    sec3_first_row = sec3_header_row + 1
    n_teams = len(TEAM_ORDER)
    sec3_last_row = sec3_first_row + n_teams - 1
    sec3_last_col = 2 + len(METRICS) * 7

    _section_title(
        ws, sec3_title_row, sec3_last_col,
        "Section 3 — Per-Team 3-Yr Decay-Weighted, Regressed Baseline. A missing year "
        "substitutes THAT season's own Section 2 league average (no rookie concept for a "
        "team-level stat).",
    )
    headers = ["Team", "Years of\nReal History"]
    metric_block_start_col: dict[str, int] = {}
    col_cursor = 3
    for m in METRICS:
        metric_block_start_col[m["key"]] = col_cursor
        headers += [
            f"{m['label']} Y-1", f"{m['label']} Y-2", f"{m['label']} Y-3",
            f"Weighted\n{m['label']} Avg\n(3-Yr decay)", f"Team History\n{m['label']}",
            f"League Baseline\n{m['label']} (Y-1)", f"Projected 3-Yr\n{m['label']} Baseline",
        ]
        col_cursor += 7
    _header_row(ws, sec3_header_row, headers)

    for i, team in enumerate(TEAM_ORDER):
        row = sec3_first_row + i
        ws.cell(row=row, column=1, value=team).font = FORMULA_FONT

        years_hist_terms = "+".join(
            f"--(COUNTIFS({team_range},$A{row},{season_range},'Model Assumptions'!$C$18-{k})>0)"
            for k in (1, 2, 3)
        )
        yh = ws.cell(row=row, column=2, value=f"={years_hist_terms}")
        yh.font = FORMULA_FONT
        yh.number_format = "0"

        for m in METRICS:
            base = metric_block_start_col[m["key"]]
            y1, y2, y3, wavg, th, lb, pb = (get_column_letter(base + k) for k in range(7))
            mrange = metric_ranges[m["key"]]
            sec2_col = season_avg_col[m["key"]]

            def _ysub(offset: int, mrange=mrange, sec2_col=sec2_col) -> str:
                season_avg_lookup = (
                    f"INDEX(${sec2_col}${season_rows[HISTORICAL_YEARS[0]]}:"
                    f"${sec2_col}${season_rows[HISTORICAL_YEARS[-1]]},"
                    f"MATCH('Model Assumptions'!$C$18-{offset},"
                    f"$A${season_rows[HISTORICAL_YEARS[0]]}:"
                    f"$A${season_rows[HISTORICAL_YEARS[-1]]},0))"
                )
                return (
                    f"=IF(COUNTIFS({team_range},$A{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset})=0,{season_avg_lookup},"
                    f"SUMIFS({mrange},{team_range},$A{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset}))"
                )

            f_y1 = ws.cell(row=row, column=base, value=_ysub(1))
            f_y2 = ws.cell(row=row, column=base + 1, value=_ysub(2))
            f_y3 = ws.cell(row=row, column=base + 2, value=_ysub(3))
            f_wavg = ws.cell(row=row, column=base + 3, value=(
                f"=({y1}{row}*1+{y2}{row}*'Model Assumptions'!$C$20+"
                f"{y3}{row}*('Model Assumptions'!$C$20^2))/"
                f"(1+'Model Assumptions'!$C$20+'Model Assumptions'!$C$20^2)"
            ))
            f_th = ws.cell(row=row, column=base + 4, value=(
                f"={y1}{row}*'Model Assumptions'!$C$22+{wavg}{row}*(1-'Model Assumptions'!$C$22)"
            ))
            f_lb = ws.cell(row=row, column=base + 5, value=(
                f"=INDEX(${sec2_col}${season_rows[HISTORICAL_YEARS[0]]}:"
                f"${sec2_col}${season_rows[HISTORICAL_YEARS[-1]]},"
                f"MATCH('Model Assumptions'!$C$18-1,$A${season_rows[HISTORICAL_YEARS[0]]}:"
                f"$A${season_rows[HISTORICAL_YEARS[-1]]},0))"
            ))
            f_pb = ws.cell(row=row, column=base + 6, value=(
                f"={th}{row}*'Model Assumptions'!$C$21+{lb}{row}*(1-'Model Assumptions'!$C$21)"
            ))
            for cell in (f_y1, f_y2, f_y3, f_wavg, f_th, f_lb, f_pb):
                cell.font = FORMULA_FONT
                cell.number_format = m["fmt"]

    # ==== Section 4: League average/std-dev of Section 3's Projected Baselines =============
    sec4_title_row = sec3_last_row + 2
    sec4_header_row = sec4_title_row + 1
    avg_row, std_row = sec4_header_row + 1, sec4_header_row + 2

    _section_title(
        ws, sec4_title_row, 5, "Section 4 — League Average & Std. Dev. of the 3-Yr Baselines"
    )
    _header_row(ws, sec4_header_row, ["Stat", *[m["label"] for m in METRICS]], height=18)

    ws.cell(row=avg_row, column=1, value="League Average").font = FORMULA_FONT
    ws.cell(row=std_row, column=1, value="League Std. Dev.").font = FORMULA_FONT

    proj_baseline_col = {
        m["key"]: get_column_letter(metric_block_start_col[m["key"]] + 6) for m in METRICS
    }
    for j, m in enumerate(METRICS):
        pcol = proj_baseline_col[m["key"]]
        pb_range = f"{pcol}${sec3_first_row}:{pcol}${sec3_last_row}"
        a = ws.cell(row=avg_row, column=2 + j, value=f"=AVERAGE({pb_range})")
        s = ws.cell(row=std_row, column=2 + j, value=f"=STDEVP({pb_range})")
        for cell in (a, s):
            cell.font = FORMULA_FONT
            cell.number_format = m["fmt"]

    # ==== Section 5: Z-scores + Explosive-Prevention Composites =============================
    sec5_title_row = std_row + 2
    sec5_header_row = sec5_title_row + 1
    sec5_first_row = sec5_header_row + 1
    sec5_last_row = sec5_first_row + n_teams - 1

    _section_title(
        ws, sec5_title_row, 9,
        "Section 5 — Z-Scores and Explosive-Prevention Composites. Explosive Pass/Run "
        "Rate (Off) are \"higher is better\" for the offense -- no sign-flip. Deep Pass "
        "Completion Rate Allowed / YAC Allowed are \"lower is better\" for the defense -- "
        "sign-flipped (league avg minus raw). Explosive Pass/Run Rate Allowed Z (F/G) are "
        "LIVE-REFERENCED from Pass/Run Defense Matchup's own Section 5 -- not recomputed. "
        "Pass Prevention Composite (H) blends 3 real signals (weighted, C136-138); Run "
        "Prevention (I) is just Explosive Run Rate Allowed's own Z -- no run-side "
        "equivalent to Deep Pass/YAC exists, a documented asymmetry.",
    )
    _header_row(
        ws, sec5_header_row,
        ["Team", "Explosive Pass\nRate (Off) Z", "Explosive Run\nRate (Off) Z",
         "Deep Pass Comp.\nAllowed Z", "YAC Allowed\nZ", "Explosive Pass\nAllowed Z (ref)",
         "Explosive Run\nAllowed Z (ref)", "Pass Prevention\nComposite Z",
         "Run Prevention\nZ", "Years of\nReal History"],
    )

    avg_cell_ref = {
        m["key"]: f"${get_column_letter(2 + j)}${avg_row}" for j, m in enumerate(METRICS)
    }
    std_cell_ref = {
        m["key"]: f"${get_column_letter(2 + j)}${std_row}" for j, m in enumerate(METRICS)
    }
    pass_def_score_range = (
        f"'{PASS_DEFENSE_SHEET}'!${PASS_DEF_EXP_Z_COL}${PASS_DEF_SEC5_RANGE[0]}:"
        f"${PASS_DEF_EXP_Z_COL}${PASS_DEF_SEC5_RANGE[1]}"
    )
    pass_def_team_range = (
        f"'{PASS_DEFENSE_SHEET}'!${PASS_DEF_TEAM_COL}${PASS_DEF_SEC5_RANGE[0]}:"
        f"${PASS_DEF_TEAM_COL}${PASS_DEF_SEC5_RANGE[1]}"
    )
    run_def_score_range = (
        f"'{RUN_DEFENSE_SHEET}'!${RUN_DEF_EXP_Z_COL}${RUN_DEF_SEC5_RANGE[0]}:"
        f"${RUN_DEF_EXP_Z_COL}${RUN_DEF_SEC5_RANGE[1]}"
    )
    run_def_team_range = (
        f"'{RUN_DEFENSE_SHEET}'!${RUN_DEF_TEAM_COL}${RUN_DEF_SEC5_RANGE[0]}:"
        f"${RUN_DEF_TEAM_COL}${RUN_DEF_SEC5_RANGE[1]}"
    )

    for i in range(n_teams):
        sec3_row = sec3_first_row + i
        row = sec5_first_row + i
        f = ws.cell(row=row, column=1, value=f"=A{sec3_row}")
        f.font = FORMULA_FONT

        for j, m in enumerate(METRICS):
            pcol = proj_baseline_col[m["key"]]
            if m["invert"]:
                formula = (
                    f"=({avg_cell_ref[m['key']]}-{pcol}{sec3_row})/{std_cell_ref[m['key']]}"
                )
            else:
                formula = (
                    f"=({pcol}{sec3_row}-{avg_cell_ref[m['key']]})/{std_cell_ref[m['key']]}"
                )
            col = 2 + j
            cell = ws.cell(row=row, column=col, value=formula)
            cell.font = FORMULA_FONT
            cell.number_format = "0.00"

        f_ref = ws.cell(row=row, column=6, value=(
            f"=IFERROR(INDEX({pass_def_score_range},MATCH(A{row},{pass_def_team_range},0)),\"\")"
        ))
        r_ref = ws.cell(row=row, column=7, value=(
            f"=IFERROR(INDEX({run_def_score_range},MATCH(A{row},{run_def_team_range},0)),\"\")"
        ))
        f_ref.font = LINK_FONT
        r_ref.font = LINK_FONT
        f_ref.number_format = "0.00"
        r_ref.number_format = "0.00"

        pass_composite = ws.cell(row=row, column=8, value=(
            f'=IF(F{row}="","",F{row}*\'Model Assumptions\'!$C$136+'
            f"D{row}*'Model Assumptions'!$C$137+E{row}*'Model Assumptions'!$C$138)"
        ))
        run_composite = ws.cell(row=row, column=9, value=f"=G{row}")
        pass_composite.font = FORMULA_FONT
        run_composite.font = FORMULA_FONT
        pass_composite.number_format = "0.00"
        run_composite.number_format = "0.00"

        yh_link = ws.cell(row=row, column=10, value=f"=B{sec3_row}")
        yh_link.font = FORMULA_FONT
        yh_link.number_format = "0"

    # ---- Closing note --------------------------------------------------------------------
    note_row = sec5_last_row + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
    note = ws.cell(row=note_row, column=1, value=(
        "claude_code_spec_explosive_play_engine.md. NOT YET VALIDATED: no backtesting "
        "harness exists yet to prove explosive-play matchups improve predictions -- this is "
        "real, verified-correct plumbing, not a proven improvement, same disclaimer as "
        "every tab in the Defensive Matchup Engine family. ONE scored threshold per phase "
        "(15+ passing, 10+ rushing, cols C/D) -- Pass 20+/30+/40+ and Run 15+/20+ (Section "
        "1, cols G-K) are CONTEXT ONLY, confirmed never referenced by any Z-score/weighted "
        "formula on this tab (a 40-yard completion is also a 30+/20+/15+ yard completion; "
        "weighting nested categories would double-count the same real play). Deep Pass "
        "Completion Rate Allowed and YAC Allowed are GENUINELY DISTINCT real metrics from "
        "Explosive Pass Rate Allowed -- different formulas, different real-world meaning "
        "(getting beaten deep vs. giving up yards after the catch vs. any pass gaining "
        "15+ total yards) -- see efficiency.py's own module docstring. NO 'missed tackles' "
        "column exists anywhere in this project -- confirmed not buildable (requires "
        "charting/tracking data no free public source has); YAC Allowed is a real, "
        "documented, honestly-labeled proxy for a RELATED but different concept, not a "
        "substitute presented as the real thing. NO Team Ratings wiring -- this is a "
        "matchup-specific composite, same reasoning as Pass/Run Defense Matchup; feeds "
        "Part C's differential (build_explosive_play_matchup_wiring.py) directly."
    ))
    note.font = NOTE_FONT
    note.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(workbook_path)
    print(
        f"Built '{SHEET_NAME}': Section 1 {len(season_stats)} rows, Section 3/5 {n_teams} "
        f"teams. NOT wired into Team Ratings -- feeds the Explosive Play Matchup wiring "
        f"directly."
    )
    print(f"Saved to {workbook_path}")
    return {
        "sec1_rows": len(season_stats), "n_teams": n_teams,
        "sec3_range": (sec3_first_row, sec3_last_row),
        "sec5_range": (sec5_first_row, sec5_last_row),
    }


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(
            'Usage: uv run python scripts/build_explosive_play_matchup.py '
            '"path/to/workbook.xlsx"'
        )
        sys.exit(1)
    build(sys.argv[1])
