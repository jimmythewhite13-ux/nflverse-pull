"""
Builds "Run Defense Matchup" -- Part B of claude_code_spec_defensive_matchup_engine.md
(Version 1 of the post-backtest roadmap). Team-level, 5-section decay-weighted pattern,
identical shape to Pass Defense Matchup / Front Seven Index / Secondary Index (no rookie
concept, no Section 2B, no individual-player Section 6).

Five scored metrics, all real and phase-SPLIT run-defense-allowed rates (efficiency.
compute_team_season_matchup_metrics -- same real, phase-split source Pass Defense Matchup
uses, confirmed with the user before building either tab that efficiency.py's OLD combined
epa_def/success_def would have put the identical non-split number in both tabs):
  - EPA/Rush Allowed
  - Run Success Rate Allowed
  - Yards/Carry Allowed
  - Explosive Run Rate Allowed
  - Stuff Rate Allowed

FOUR of the five are "lower raw value = better defense" (EPA/Rush, Run Success Rate, Yards/
Carry, Explosive Run Rate, all Allowed) -- but Stuff Rate Allowed is the OPPOSITE direction:
a HIGHER share of opponent rushes stuffed for 0-or-negative yards is a BETTER run defense.
Section 5 therefore uses a per-metric `invert` flag (unlike Pass Defense Matchup, where all
five metrics happened to share the same direction and could be sign-flipped uniformly) --
this asymmetry is real and deliberate, not an oversight; getting it backwards would score
a good stuff-heavy run defense as bad.

No REFERENCED columns from Front Seven Index on this tab -- the spec's own Part B wording
only asked for Sack Rate/Pressure Proxy/Blitz Rate context on Pass Defense Matchup, not
here (a blitz's real value shows up in the passing game it's designed to disrupt, not run
defense).

NOT VALIDATED YET: this tab is real, verified-correct plumbing -- there is no backtesting
harness yet to prove phase-specific matchups actually improve predictions over the existing
PPG-based blend. Stated explicitly on the tab itself, per the spec's own instruction. Feeds
Part C's matchup differential on Week 1 Matchups, which adds this ALONGSIDE the existing
PPG-based prediction rather than replacing it, specifically so a future backtest can compare
the two directly.

Usage:
    uv run python scripts/build_run_defense_matchup.py "C:\\path\\to\\NFL_Prediction_Model.xlsx"
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from nflverse_pull.efficiency import compute_team_season_matchup_metrics, fetch_pbp  # noqa: E402

HISTORICAL_YEARS = [2023, 2024, 2025]
SHEET_NAME = "Run Defense Matchup"
PASS_DEFENSE_SHEET = "Pass Defense Matchup"
FRONT_SEVEN_SHEET = "Front Seven Index"
TEAM_RATINGS_SHEET = "Team Ratings"

METRICS = [
    {"key": "epa_rush", "col": "EPA/Rush Allowed (Def)", "label": "EPA/Rush\nAllowed",
     "fmt": "0.000", "invert": True},
    {"key": "rsr", "col": "Run Success Rate Allowed (Def)", "label": "Run Success\nRate Allowed",
     "fmt": "0.00%", "invert": True},
    {"key": "ypc", "col": "Yards/Carry Allowed (Def)", "label": "Yards/Carry\nAllowed",
     "fmt": "0.00", "invert": True},
    {"key": "exp", "col": "Explosive Run Rate Allowed (Def)",
     "label": "Explosive Run\nRate Allowed", "fmt": "0.00%", "invert": True},
    {"key": "stuff", "col": "Stuff Rate Allowed (Def)", "label": "Stuff Rate\nAllowed",
     "fmt": "0.00%", "invert": False},
]

TEAM_ORDER = [
    "Buffalo Bills", "Miami Dolphins", "New England Patriots", "New York Jets",
    "Baltimore Ravens", "Cincinnati Bengals", "Cleveland Browns", "Pittsburgh Steelers",
    "Houston Texans", "Indianapolis Colts", "Jacksonville Jaguars", "Tennessee Titans",
    "Denver Broncos", "Kansas City Chiefs", "Las Vegas Raiders", "Los Angeles Chargers",
    "Dallas Cowboys", "New York Giants", "Philadelphia Eagles", "Washington Commanders",
    "Chicago Bears", "Detroit Lions", "Green Bay Packers", "Minnesota Vikings",
    "Atlanta Falcons", "Carolina Panthers", "New Orleans Saints", "Tampa Bay Buccaneers",
    "Arizona Cardinals", "Los Angeles Rams", "San Francisco 49ers", "Seattle Seahawks",
]

TITLE_FONT = Font(name="Arial", size=10, bold=True)
TITLE_FILL = PatternFill("solid", fgColor="FFD9E1F2")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_ALIGN = Alignment(wrap_text=True, horizontal="center", vertical="center")
INPUT_FONT = Font(name="Arial", size=10, color="FF0000FF")
FORMULA_FONT = Font(name="Arial", size=10, color="FF000000")
NOTE_FONT = Font(name="Arial", size=9, color="FF808080")
ASSUMPTION_FILL = PatternFill("solid", fgColor="FFFFFF00")


def _section_title(ws, row: int, last_col: int, text: str) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL


def _header_row(ws, row: int, headers: list[str], height: float = 28) -> None:
    ws.row_dimensions[row].height = height
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN


def add_model_assumptions_weights(wb: openpyxl.Workbook) -> None:
    ws = wb["Model Assumptions"]
    title_row = 93
    ws.merge_cells(f"A{title_row}:D{title_row}")
    title = ws.cell(row=title_row, column=1, value=(
        "Run Defense Matchup Weighting & Conversion (pts per std. dev.; reuses QB Index's "
        "points-scale constants C37/C38 -- see 'Run Defense Matchup' tab)"
    ))
    title.font = HEADER_FONT
    title.fill = HEADER_FILL

    rows = [
        (94, "EPA/Rush Allowed Weight (Run Defense Matchup, pts per SD)", 0.35,
         "The most complete single run-defense value stat -- weighted highest, same "
         "logic as every other EPA-based weighting in this model."),
        (95, "Run Success Rate Allowed Weight (Run Defense Matchup, pts per SD)", 0.25,
         "Correlates with EPA but isolates consistency (positive-EPA-allowed rate) "
         "rather than magnitude."),
        (96, "Yards/Carry Allowed Weight (Run Defense Matchup, pts per SD)", 0.20,
         "Traditional, well-understood counting stat -- a sanity check alongside the "
         "EPA-based metrics."),
        (97, "Explosive Run Rate Allowed Weight (Run Defense Matchup, pts per SD)", 0.10,
         "A real tail-risk signal (big-play rate allowed) distinct from the average-"
         "based metrics above -- weighted lowest as a supplementary signal."),
        (98, "Stuff Rate Allowed Weight (Run Defense Matchup, pts per SD)", 0.10,
         "Real havoc-rate signal (share of carries stopped at or behind the line) -- "
         "the ONE metric on this tab where HIGHER raw value means better defense; "
         "Section 5 inverts this one's Z-score direction relative to the other four "
         "(see this tab's own module docstring/closing note)."),
        (99, "Run Defense Matchup Points-to-Game-Points Conversion", 0.10,
         "A starting guess, like every other coefficient in this model -- a "
         "structurally different, dedicated constant, NOT reused from any other "
         "tab's conversion factor (this measures a phase-specific MATCHUP "
         "differential, not a team-level quality adjustment or a Replacement Value "
         "swap). NOT YET VALIDATED against real outcomes -- see the tab's own "
         "closing note and claude_code_spec_defensive_matchup_engine.md."),
    ]
    for row, label, value, note in rows:
        ws.cell(row=row, column=2, value=label)
        c = ws.cell(row=row, column=3, value=value)
        c.font = INPUT_FONT
        c.fill = ASSUMPTION_FILL
        c.number_format = "0.00"
        n = ws.cell(row=row, column=4, value=note)
        n.font = NOTE_FONT
        n.alignment = Alignment(wrap_text=True, vertical="top")


def _pull_data() -> pd.DataFrame:
    print(f"Pulling {HISTORICAL_YEARS} play-by-play data for run-defense matchup metrics...")
    pbp = fetch_pbp(HISTORICAL_YEARS)
    return compute_team_season_matchup_metrics(pbp)


def build(workbook_path: str) -> dict:
    season_stats = _pull_data()

    wb = openpyxl.load_workbook(workbook_path)
    add_model_assumptions_weights(wb)

    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    for candidate in (PASS_DEFENSE_SHEET, FRONT_SEVEN_SHEET):
        if candidate in wb.sheetnames:
            insert_after = candidate
            break
    else:
        insert_after = wb.sheetnames[0]
    ws = wb.create_sheet(SHEET_NAME, index=wb.sheetnames.index(insert_after) + 1)
    ws.column_dimensions["A"].width = 20.0
    ws.column_dimensions["C"].width = 20.0

    ws.merge_cells("A1:H1")
    t = ws.cell(row=1, column=1, value=(
        "Run Defense Matchup -- Multi-Year Decay-Weighted TEAM-Level Run Defense Rating "
        "(EPA/Rush, Run Success Rate, Yards/Carry, Explosive Run Rate, Stuff Rate, all "
        "ALLOWED; real, phase-SPLIT pbp metrics). NOT YET VALIDATED against real outcomes "
        "-- see the closing note. Feeds Part C's matchup differential on Week 1 Matchups, "
        "ALONGSIDE (not replacing) the existing PPG-based prediction."
    ))
    t.font = Font(name="Arial", size=12, bold=True)

    # ==== Section 1: Raw 3-year TEAM-level data =============================================
    sec1_first_row = 5
    sec1_last_row = sec1_first_row + len(season_stats) - 1
    _section_title(
        ws, 3, 7,
        "Section 1 — Raw 3-Year TEAM-Level Run-Defense-Allowed Rates (real nflverse "
        "play-by-play, phase-split -- see this tab's opening note).",
    )
    _header_row(ws, 4, ["Team", "Season", *[m["label"] for m in METRICS]])
    for i, r in enumerate(season_stats.to_dict("records")):
        row = sec1_first_row + i
        values = [r["Team"], int(r["Season"])]
        for m in METRICS:
            v = r.get(m["col"])
            values.append(float(v) if pd.notna(v) else None)
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.font = INPUT_FONT
        for j, m in enumerate(METRICS):
            ws.cell(row=row, column=3 + j).number_format = m["fmt"]

    team_range = f"$A${sec1_first_row}:$A${sec1_last_row}"
    season_range = f"$B${sec1_first_row}:$B${sec1_last_row}"
    sec1_col_of = {m["key"]: get_column_letter(3 + j) for j, m in enumerate(METRICS)}
    metric_ranges = {
        m["key"]: (
            f"${sec1_col_of[m['key']]}${sec1_first_row}:${sec1_col_of[m['key']]}${sec1_last_row}"
        )
        for m in METRICS
    }

    # ==== Section 2: League average per season ==============================================
    sec2_title_row = sec1_last_row + 2
    sec2_header_row = sec2_title_row + 1
    season_rows = {yr: sec2_header_row + 1 + i for i, yr in enumerate(HISTORICAL_YEARS)}

    _section_title(
        ws, sec2_title_row, 6,
        "Section 2 — League Average per Season (simple average across all 32 teams)",
    )
    _header_row(ws, sec2_header_row, ["Season", *[m["label"] for m in METRICS]], height=20)

    for yr in HISTORICAL_YEARS:
        row = season_rows[yr]
        c = ws.cell(row=row, column=1, value=yr)
        c.font = INPUT_FONT
        for j, m in enumerate(METRICS):
            formula = f"=AVERAGEIF({season_range},{yr},{metric_ranges[m['key']]})"
            cell = ws.cell(row=row, column=2 + j, value=formula)
            cell.font = FORMULA_FONT
            cell.number_format = m["fmt"]

    season_avg_col = {m["key"]: get_column_letter(2 + j) for j, m in enumerate(METRICS)}

    # ==== Section 3: Per-team decay-weighted, regressed baseline ===========================
    sec3_title_row = season_rows[HISTORICAL_YEARS[-1]] + 2
    sec3_header_row = sec3_title_row + 1
    sec3_first_row = sec3_header_row + 1
    n_teams = len(TEAM_ORDER)
    sec3_last_row = sec3_first_row + n_teams - 1
    sec3_last_col = 2 + len(METRICS) * 7

    _section_title(
        ws, sec3_title_row, sec3_last_col,
        "Section 3 — Per-Team 3-Yr Decay-Weighted, Regressed Baseline. A missing year "
        "substitutes THAT season's own Section 2 league average (no rookie concept for a "
        "team-level stat).",
    )
    headers = ["Team", "Years of\nReal History"]
    metric_block_start_col: dict[str, int] = {}
    col_cursor = 3
    for m in METRICS:
        metric_block_start_col[m["key"]] = col_cursor
        headers += [
            f"{m['label']} Y-1", f"{m['label']} Y-2", f"{m['label']} Y-3",
            f"Weighted\n{m['label']} Avg\n(3-Yr decay)", f"Team History\n{m['label']}",
            f"League Baseline\n{m['label']} (Y-1)", f"Projected 3-Yr\n{m['label']} Baseline",
        ]
        col_cursor += 7
    _header_row(ws, sec3_header_row, headers)

    for i, team in enumerate(TEAM_ORDER):
        row = sec3_first_row + i
        ws.cell(row=row, column=1, value=team).font = FORMULA_FONT

        years_hist_terms = "+".join(
            f"--(COUNTIFS({team_range},$A{row},{season_range},'Model Assumptions'!$C$18-{k})>0)"
            for k in (1, 2, 3)
        )
        yh = ws.cell(row=row, column=2, value=f"={years_hist_terms}")
        yh.font = FORMULA_FONT
        yh.number_format = "0"

        for m in METRICS:
            base = metric_block_start_col[m["key"]]
            y1, y2, y3, wavg, th, lb, pb = (get_column_letter(base + k) for k in range(7))
            mrange = metric_ranges[m["key"]]
            sec2_col = season_avg_col[m["key"]]

            def _ysub(offset: int, mrange=mrange, sec2_col=sec2_col) -> str:
                season_avg_lookup = (
                    f"INDEX(${sec2_col}${season_rows[HISTORICAL_YEARS[0]]}:"
                    f"${sec2_col}${season_rows[HISTORICAL_YEARS[-1]]},"
                    f"MATCH('Model Assumptions'!$C$18-{offset},"
                    f"$A${season_rows[HISTORICAL_YEARS[0]]}:"
                    f"$A${season_rows[HISTORICAL_YEARS[-1]]},0))"
                )
                return (
                    f"=IF(COUNTIFS({team_range},$A{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset})=0,{season_avg_lookup},"
                    f"SUMIFS({mrange},{team_range},$A{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset}))"
                )

            f_y1 = ws.cell(row=row, column=base, value=_ysub(1))
            f_y2 = ws.cell(row=row, column=base + 1, value=_ysub(2))
            f_y3 = ws.cell(row=row, column=base + 2, value=_ysub(3))
            f_wavg = ws.cell(row=row, column=base + 3, value=(
                f"=({y1}{row}*1+{y2}{row}*'Model Assumptions'!$C$20+"
                f"{y3}{row}*('Model Assumptions'!$C$20^2))/"
                f"(1+'Model Assumptions'!$C$20+'Model Assumptions'!$C$20^2)"
            ))
            f_th = ws.cell(row=row, column=base + 4, value=(
                f"={y1}{row}*'Model Assumptions'!$C$22+{wavg}{row}*(1-'Model Assumptions'!$C$22)"
            ))
            f_lb = ws.cell(row=row, column=base + 5, value=(
                f"=INDEX(${sec2_col}${season_rows[HISTORICAL_YEARS[0]]}:"
                f"${sec2_col}${season_rows[HISTORICAL_YEARS[-1]]},"
                f"MATCH('Model Assumptions'!$C$18-1,$A${season_rows[HISTORICAL_YEARS[0]]}:"
                f"$A${season_rows[HISTORICAL_YEARS[-1]]},0))"
            ))
            f_pb = ws.cell(row=row, column=base + 6, value=(
                f"={th}{row}*'Model Assumptions'!$C$21+{lb}{row}*(1-'Model Assumptions'!$C$21)"
            ))
            for cell in (f_y1, f_y2, f_y3, f_wavg, f_th, f_lb, f_pb):
                cell.font = FORMULA_FONT
                cell.number_format = m["fmt"]

    # ==== Section 4: League average/std-dev of Section 3's Projected Baselines =============
    sec4_title_row = sec3_last_row + 2
    sec4_header_row = sec4_title_row + 1
    avg_row, std_row = sec4_header_row + 1, sec4_header_row + 2

    _section_title(
        ws, sec4_title_row, 6, "Section 4 — League Average & Std. Dev. of the 3-Yr Baselines"
    )
    _header_row(ws, sec4_header_row, ["Stat", *[m["label"] for m in METRICS]], height=18)

    ws.cell(row=avg_row, column=1, value="League Average").font = FORMULA_FONT
    ws.cell(row=std_row, column=1, value="League Std. Dev.").font = FORMULA_FONT

    proj_baseline_col = {
        m["key"]: get_column_letter(metric_block_start_col[m["key"]] + 6) for m in METRICS
    }
    for j, m in enumerate(METRICS):
        pcol = proj_baseline_col[m["key"]]
        pb_range = f"{pcol}${sec3_first_row}:{pcol}${sec3_last_row}"
        a = ws.cell(row=avg_row, column=2 + j, value=f"=AVERAGE({pb_range})")
        s = ws.cell(row=std_row, column=2 + j, value=f"=STDEVP({pb_range})")
        for cell in (a, s):
            cell.font = FORMULA_FONT
            cell.number_format = m["fmt"]

    # ==== Section 5: Z-scores, weighted composite, points-scale Run Defense Score ==========
    sec5_title_row = std_row + 2
    sec5_header_row = sec5_title_row + 1
    sec5_first_row = sec5_header_row + 1
    sec5_last_row = sec5_first_row + n_teams - 1

    _section_title(
        ws, sec5_title_row, 8,
        "Section 5 — Z-Scores and Run Defense Score. FOUR of five metrics are ALLOWED "
        "rates where LOWER is better (sign-flipped: league avg minus raw); Stuff Rate "
        "Allowed is the OPPOSITE (higher = better, NOT sign-flipped) -- a real, deliberate "
        "asymmetry, not an oversight (see this tab's own closing note). Baseline/points-"
        "per-SD reuse QB Index's own Model Assumptions cells C37/C38.",
    )
    _header_row(
        ws, sec5_header_row,
        ["Team", *[f"{m['label']}\nZ" for m in METRICS], "Weighted\nZ-Score Sum",
         "Run Defense\nScore (Points)", "Years of\nReal History"],
    )

    avg_cell_ref = {
        m["key"]: f"${get_column_letter(2 + j)}${avg_row}" for j, m in enumerate(METRICS)
    }
    std_cell_ref = {
        m["key"]: f"${get_column_letter(2 + j)}${std_row}" for j, m in enumerate(METRICS)
    }
    weight_cells = {"epa_rush": "$C$94", "rsr": "$C$95", "ypc": "$C$96", "exp": "$C$97",
                     "stuff": "$C$98"}

    for i in range(n_teams):
        sec3_row = sec3_first_row + i
        row = sec5_first_row + i
        f = ws.cell(row=row, column=1, value=f"=A{sec3_row}")
        f.font = FORMULA_FONT

        z_cols = []
        for j, m in enumerate(METRICS):
            pcol = proj_baseline_col[m["key"]]
            # Four metrics are "allowed" rates where lower=better (league avg minus raw);
            # Stuff Rate Allowed is the opposite (higher=better, raw minus league avg) --
            # per-metric invert flag, deliberately NOT uniform like Pass Defense Matchup.
            if m["invert"]:
                formula = f"=({avg_cell_ref[m['key']]}-{pcol}{sec3_row})/{std_cell_ref[m['key']]}"
            else:
                formula = f"=({pcol}{sec3_row}-{avg_cell_ref[m['key']]})/{std_cell_ref[m['key']]}"
            col = 2 + j
            cell = ws.cell(row=row, column=col, value=formula)
            cell.font = FORMULA_FONT
            cell.number_format = "0.00"
            z_cols.append(get_column_letter(col))

        weighted_terms = " + ".join(
            f"{z_cols[j]}{row}*'Model Assumptions'!{weight_cells[m['key']]}"
            for j, m in enumerate(METRICS)
        )
        wz_col = 2 + len(METRICS)
        wz = ws.cell(row=row, column=wz_col, value=f"={weighted_terms}")
        wz.font = FORMULA_FONT
        wz.number_format = "0.00"

        score = ws.cell(row=row, column=wz_col + 1, value=(
            f"='Model Assumptions'!$C$37+{get_column_letter(wz_col)}{row}*'Model Assumptions'!$C$38"
        ))
        score.font = FORMULA_FONT
        score.number_format = "0.0;(0.0)"

        yh_link = ws.cell(row=row, column=wz_col + 2, value=f"=B{sec3_row}")
        yh_link.font = FORMULA_FONT
        yh_link.number_format = "0"

    # ---- Closing note --------------------------------------------------------------------
    note_row = sec5_last_row + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)
    note = ws.cell(row=note_row, column=1, value=(
        "claude_code_spec_defensive_matchup_engine.md, Version 1 of the post-backtest "
        "roadmap. NOT YET VALIDATED: there is no backtesting harness in this project yet "
        "to prove phase-specific run-defense matchups actually improve predictions over "
        "the existing PPG-based Off/Def blend -- this tab is real, verified-correct "
        "plumbing (every metric traces to real nflverse pbp, the same decay-weighted/"
        "regressed/Z-scored chain every prior tab uses), NOT a proven improvement. Do not "
        "treat its output as more reliable than the existing prediction until a real "
        "backtest exists to check that claim. All five scored metrics are real, phase-"
        "SPLIT run-defense-ALLOWED rates (efficiency.compute_team_season_matchup_metrics) "
        "-- deliberately NOT the combined pass+run epa_def/success_def efficiency.py "
        "already had, which would have put the identical non-split number in this tab and "
        "Pass Defense Matchup, defeating the point of splitting by phase (confirmed with "
        "the user before building either tab). STUFF RATE ALLOWED IS INTENTIONALLY NOT "
        "SIGN-FLIPPED, unlike the other four metrics -- a higher share of opponent rushes "
        "stuffed for 0-or-negative yards is a BETTER run defense, the opposite direction "
        "from EPA/Rush, Run Success Rate, Yards/Carry, and Explosive Run Rate (all "
        "Allowed, where lower is better). No REFERENCED Front-Seven-Index context columns "
        "on this tab (unlike Pass Defense Matchup's Sack Rate/Pressure Proxy/Blitz Rate) "
        "-- the spec's own Part B wording only called for that context on the pass side. "
        "NO man/zone coverage tendency column exists anywhere on this tab -- confirmed "
        "proprietary (paid FTN/Fantasy Points charting only), not approximated with a "
        "fabricated proxy. This tab has NO Team Ratings wiring -- Part C references its "
        "Section 5 Score directly from Week 1 Matchups instead, alongside (not replacing) "
        "the existing PPG-based prediction, per the spec's own instruction."
    ))
    note.font = NOTE_FONT
    note.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(workbook_path)
    print(
        f"Built '{SHEET_NAME}': Section 1 {len(season_stats)} rows, Section 3/5 {n_teams} "
        f"teams. NOT wired into Team Ratings -- Part C references Section 5 directly from "
        f"Week 1 Matchups."
    )
    print(f"Saved to {workbook_path}")
    return {
        "sec1_rows": len(season_stats), "n_teams": n_teams,
        "sec3_range": (sec3_first_row, sec3_last_row),
        "sec5_range": (sec5_first_row, sec5_last_row),
    }


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print('Usage: uv run python scripts/build_run_defense_matchup.py "path/to/workbook.xlsx"')
        sys.exit(1)
    build(sys.argv[1])
