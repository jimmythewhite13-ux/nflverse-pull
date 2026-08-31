"""
Builds "Secondary Index" -- completes Phase 4 of the multi-phase roadmap in
claude_code_spec_rb_index.md ("Defense -- counting-stats index only (sacks/TFL/INT/PBU),
explicitly NOT an EPA-equivalent value model"), covering the rest of the defense (CB/S)
that build_defense_index.py's Front Seven & D-Line Index deliberately left out. Built per
explicit user instruction to apply that same real-data-only approach here. No separate
written spec exists for this phase -- designed directly, mirroring build_defense_index.py's
own structure (see that script's own docstring for the full reasoning this one shares).

Simpler than the front seven: the 5 secondary depth-chart positions (LCB/RCB/NB/FS/SS) are
each populated for all 32 teams, always exactly one rank-1 starter apiece -- verified live.
No scheme-variation priority-order mapping is needed the way interior line/linebacker
required (current_roster.py's own comment covers why those varied and these don't).

Section 1: raw 3-year TEAM-level rates (INT Rate, PBU Rate) from defense_stats.py
Section 2: league average per season (simple average across 32 teams)
Section 3: per-team 3-Yr decay-weighted, regressed baseline -- a missing year substitutes
           that season's own Section 2 league average (no rookie concept for a team stat)
Section 4: league average/std-dev of Section 3's projected baselines
Section 5: Z-scores, weighted composite, points-scale Team Secondary Score (reuses QB
           Index's own Model Assumptions C37/C38 scale constants)
Section 6: Individual Real Production -- the current real starters at CB1/CB2/NB/FS/SS
           (current_roster.py's live depth-chart pull -- a direct, fixed mapping, no
           priority-order juggling needed), each showing THEIR OWN real season INT/PBU
           (defense_stats.compute_player_season_secondary_stats) -- informational, not fed
           back into team-level scoring (Section 1 already captures it via the team rollup)
Section 7: direct (not Replacement-Value-swap) Team Ratings adjustment, same unconditional
           pattern as every other position tab wired directly -- there's no "backup
           secondary" to swap in as a unit

Usage:
    uv run python scripts/build_secondary_index.py "C:\\path\\to\\NFL_Prediction_Model.xlsx"
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from current_season_blend import add_current_season_blend  # noqa: E402

from nflverse_pull.current_roster import (  # noqa: E402
    compute_current_starters,
    fetch_depth_charts,
    resolve_scored_population,
)
from nflverse_pull.defense_stats import (  # noqa: E402
    compute_player_season_secondary_stats,
    compute_team_season_secondary_stats,
)
from nflverse_pull.efficiency import fetch_pbp  # noqa: E402

HISTORICAL_YEARS = [2023, 2024, 2025]
CURRENT_ROSTER_YEAR = 2026
SHEET_NAME = "Secondary Index"
FRONT7_SHEET = "Front Seven Index"
OLINE_SHEET = "Offensive Line Index"
KICKING_SHEET = "Kicking Index"
WR_TE_SHEET = "WR-TE Value Index"
RB_INDEX_SHEET = "RB Value Index"
QB_INDEX_SHEET = "QB Index"
TEAM_RATINGS_SHEET = "Team Ratings"

METRICS = [
    {"key": "int", "col": "INT Rate", "label": "INT\nRate", "fmt": "0.00%"},
    {"key": "pbu", "col": "PBU Rate", "label": "PBU\nRate", "fmt": "0.00%"},
]

TEAM_ORDER = [
    "Buffalo Bills", "Miami Dolphins", "New England Patriots", "New York Jets",
    "Baltimore Ravens", "Cincinnati Bengals", "Cleveland Browns", "Pittsburgh Steelers",
    "Houston Texans", "Indianapolis Colts", "Jacksonville Jaguars", "Tennessee Titans",
    "Denver Broncos", "Kansas City Chiefs", "Las Vegas Raiders", "Los Angeles Chargers",
    "Dallas Cowboys", "New York Giants", "Philadelphia Eagles", "Washington Commanders",
    "Chicago Bears", "Detroit Lions", "Green Bay Packers", "Minnesota Vikings",
    "Atlanta Falcons", "Carolina Panthers", "New Orleans Saints", "Tampa Bay Buccaneers",
    "Arizona Cardinals", "Los Angeles Rams", "San Francisco 49ers", "Seattle Seahawks",
]

# Direct, fixed mapping -- every one of these 5 real depth-chart positions is populated for
# all 32 teams, always exactly one rank-1 starter apiece (verified live), unlike the front
# seven's scheme-variable interior line/linebacker groups.
SECONDARY_MAP = [("LCB", "CB1"), ("RCB", "CB2"), ("NB", "NB"), ("FS", "FS"), ("SS", "SS")]

TITLE_FONT = Font(name="Arial", size=10, bold=True)
TITLE_FILL = PatternFill("solid", fgColor="FFD9E1F2")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_ALIGN = Alignment(wrap_text=True, horizontal="center", vertical="center")
INPUT_FONT = Font(name="Arial", size=10, color="FF0000FF")
FORMULA_FONT = Font(name="Arial", size=10, color="FF000000")
LINK_FONT = Font(name="Arial", size=10, color="FF008000")
NOTE_FONT = Font(name="Arial", size=9, color="FF808080")
ASSUMPTION_FILL = PatternFill("solid", fgColor="FFFFFF00")


def _section_title(ws, row: int, last_col: int, text: str) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL


def _header_row(ws, row: int, headers: list[str], height: float = 28) -> None:
    ws.row_dimensions[row].height = height
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN


def add_model_assumptions_weights(wb: openpyxl.Workbook) -> None:
    ws = wb["Model Assumptions"]
    title_row = 73
    ws.merge_cells(f"A{title_row}:D{title_row}")
    title = ws.cell(row=title_row, column=1, value=(
        "Secondary Index Weighting & Conversion (pts per std. dev.; reuses QB Index's "
        "points-scale constants C37/C38 -- see 'Secondary Index' tab)"
    ))
    title.font = HEADER_FONT
    title.fill = HEADER_FILL

    rows = [
        (74, "INT Rate Weight (Secondary Index, pts per SD)", 0.5,
         "Real, individually-credited interceptions per real pass play faced -- weighted "
         "equally with PBU Rate since both are genuinely real signals, unlike the front "
         "seven's 3-metric split (there's no honest third team-level secondary metric "
         "available)."),
        (75, "PBU Rate Weight (Secondary Index, pts per SD)", 0.5,
         "Real pass breakups per real pass play faced -- a much larger, less variance-"
         "prone sample than interceptions, so it's a meaningful complement even at an "
         "equal weight."),
        (76, "Secondary Index Points-to-Game-Points Conversion", 0.06,
         "A starting guess, like every other coefficient in this model. This is a "
         "DIRECT quality adjustment, not a Replacement Value swap -- there's no 'backup "
         "secondary' to swap in as a unit, same reasoning as every other position-group "
         "tab wired directly into Team Ratings."),
    ]
    for row, label, value, note in rows:
        ws.cell(row=row, column=2, value=label)
        c = ws.cell(row=row, column=3, value=value)
        c.font = INPUT_FONT
        c.fill = ASSUMPTION_FILL
        c.number_format = "0.00"
        n = ws.cell(row=row, column=4, value=note)
        n.font = NOTE_FONT
        n.alignment = Alignment(wrap_text=True, vertical="top")


def _pull_data() -> dict:
    print(f"Pulling {HISTORICAL_YEARS} play-by-play data for secondary counting stats...")
    pbp = fetch_pbp(HISTORICAL_YEARS)
    team_stats = compute_team_season_secondary_stats(pbp)
    player_stats = compute_player_season_secondary_stats(pbp)

    print(f"Pulling {CURRENT_ROSTER_YEAR} depth charts for the secondary starters...")
    depth_charts = fetch_depth_charts([CURRENT_ROSTER_YEAR])
    current_starters = compute_current_starters(depth_charts)
    empty_overrides = pd.DataFrame(
        columns=["Team", "Manual Starter Override", "Manual Backup Override"]
    )

    slot_frames = []
    for pos, slot in SECONDARY_MAP:
        pop = resolve_scored_population(current_starters, empty_overrides, pos)
        pop = pop.copy()
        pop["Role"] = slot
        slot_frames.append(pop)
    starters = pd.concat(slot_frames, ignore_index=True)

    role_priority = {slot: i for i, (_, slot) in enumerate(SECONDARY_MAP, start=1)}
    starters["_priority"] = starters["Role"].map(role_priority)
    starters = starters.sort_values(["Team", "_priority"]).drop(columns="_priority")
    starters = starters.reset_index(drop=True)

    print(f"{len(starters)} current-roster secondary starters identified across "
          f"{len(SECONDARY_MAP)} slots/team.")

    return {"team_stats": team_stats, "player_stats": player_stats, "starters": starters}


def build(workbook_path: str) -> dict:
    data = _pull_data()
    team_stats = data["team_stats"]
    player_stats = data["player_stats"]
    starters = data["starters"]

    wb = openpyxl.load_workbook(workbook_path)
    add_model_assumptions_weights(wb)

    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    for candidate in (FRONT7_SHEET, OLINE_SHEET, KICKING_SHEET, WR_TE_SHEET,
                      RB_INDEX_SHEET, QB_INDEX_SHEET):
        if candidate in wb.sheetnames:
            insert_after = candidate
            break
    else:
        insert_after = wb.sheetnames[0]
    ws = wb.create_sheet(SHEET_NAME, index=wb.sheetnames.index(insert_after) + 1)
    ws.column_dimensions["A"].width = 20.0
    ws.column_dimensions["C"].width = 20.0

    ws.merge_cells("A1:H1")
    t = ws.cell(row=1, column=1, value=(
        "Secondary Index -- Multi-Year Decay-Weighted TEAM-Level Coverage Rating (INT "
        "Rate, PBU Rate; real nflverse play-by-play counting stats). Completes the "
        "defense alongside Front Seven & D-Line Index. NO invented per-position grade "
        "appears anywhere -- see the closing note."
    ))
    t.font = Font(name="Arial", size=12, bold=True)

    # ==== Section 1: Raw 3-year TEAM-level data ============================================
    sec1_first_row = 5
    sec1_last_row = sec1_first_row + len(team_stats) - 1
    _section_title(
        ws, 3, 4,
        "Section 1 \u2014 Raw 3-Year TEAM-Level Rates (real nflverse play-by-play counting "
        "stats, rolled up from every real contributing player)",
    )
    _header_row(ws, 4, ["Team", "Season", *[m["label"] for m in METRICS]])
    for i, r in enumerate(team_stats.to_dict("records")):
        row = sec1_first_row + i
        values = [r["Team"], int(r["Season"]), r["INT Rate"], r["PBU Rate"]]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.font = INPUT_FONT
        for col in (3, 4):
            ws.cell(row=row, column=col).number_format = "0.00%"

    team_range = f"$A${sec1_first_row}:$A${sec1_last_row}"
    season_range = f"$B${sec1_first_row}:$B${sec1_last_row}"
    sec1_col_of = {"int": "C", "pbu": "D"}
    metric_ranges = {
        m["key"]: (
            f"${sec1_col_of[m['key']]}${sec1_first_row}:${sec1_col_of[m['key']]}${sec1_last_row}"
        )
        for m in METRICS
    }

    # ==== Section 2: League average per season ==============================================
    sec2_title_row = sec1_last_row + 2
    sec2_header_row = sec2_title_row + 1
    season_rows = {yr: sec2_header_row + 1 + i for i, yr in enumerate(HISTORICAL_YEARS)}

    _section_title(
        ws, sec2_title_row, 3,
        "Section 2 \u2014 League Average per Season (simple average across all 32 teams)",
    )
    _header_row(ws, sec2_header_row, ["Season", *[m["label"] for m in METRICS]], height=20)

    for yr in HISTORICAL_YEARS:
        row = season_rows[yr]
        c = ws.cell(row=row, column=1, value=yr)
        c.font = INPUT_FONT
        for j, m in enumerate(METRICS):
            formula = f"=AVERAGEIF({season_range},{yr},{metric_ranges[m['key']]})"
            cell = ws.cell(row=row, column=2 + j, value=formula)
            cell.font = FORMULA_FONT
            cell.number_format = m["fmt"]

    season_avg_col = {m["key"]: get_column_letter(2 + j) for j, m in enumerate(METRICS)}

    # ==== Section 3: Per-team decay-weighted, regressed baseline ===========================
    sec3_title_row = season_rows[HISTORICAL_YEARS[-1]] + 2
    sec3_header_row = sec3_title_row + 1
    sec3_first_row = sec3_header_row + 1
    n_teams = len(TEAM_ORDER)
    sec3_last_row = sec3_first_row + n_teams - 1
    sec3_last_col = 2 + len(METRICS) * 7

    _section_title(
        ws, sec3_title_row, sec3_last_col,
        "Section 3 \u2014 Per-Team 3-Yr Decay-Weighted, Regressed Baseline. A missing year "
        "substitutes THAT season's own Section 2 league average (no rookie concept for a "
        "team-level stat).",
    )
    headers = ["Team", "Years of\nReal History"]
    metric_block_start_col: dict[str, int] = {}
    col_cursor = 3
    for m in METRICS:
        metric_block_start_col[m["key"]] = col_cursor
        headers += [
            f"{m['label']} Y-1", f"{m['label']} Y-2", f"{m['label']} Y-3",
            f"Weighted\n{m['label']} Avg\n(3-Yr decay)", f"Team History\n{m['label']}",
            f"League Baseline\n{m['label']} (Y-1)", f"Projected 3-Yr\n{m['label']} Baseline",
        ]
        col_cursor += 7
    _header_row(ws, sec3_header_row, headers)

    for i, team in enumerate(TEAM_ORDER):
        row = sec3_first_row + i
        ws.cell(row=row, column=1, value=team).font = FORMULA_FONT

        years_hist_terms = "+".join(
            f"--(COUNTIFS({team_range},$A{row},{season_range},'Model Assumptions'!$C$18-{k})>0)"
            for k in (1, 2, 3)
        )
        yh = ws.cell(row=row, column=2, value=f"={years_hist_terms}")
        yh.font = FORMULA_FONT
        yh.number_format = "0"

        for m in METRICS:
            base = metric_block_start_col[m["key"]]
            y1, y2, y3, wavg, th, lb, pb = (get_column_letter(base + k) for k in range(7))
            mrange = metric_ranges[m["key"]]
            sec2_col = season_avg_col[m["key"]]

            def _ysub(offset: int, mrange=mrange, sec2_col=sec2_col) -> str:
                season_avg_lookup = (
                    f"INDEX(${sec2_col}${season_rows[HISTORICAL_YEARS[0]]}:"
                    f"${sec2_col}${season_rows[HISTORICAL_YEARS[-1]]},"
                    f"MATCH('Model Assumptions'!$C$18-{offset},"
                    f"$A${season_rows[HISTORICAL_YEARS[0]]}:"
                    f"$A${season_rows[HISTORICAL_YEARS[-1]]},0))"
                )
                return (
                    f"=IF(COUNTIFS({team_range},$A{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset})=0,{season_avg_lookup},"
                    f"SUMIFS({mrange},{team_range},$A{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset}))"
                )

            f_y1 = ws.cell(row=row, column=base, value=_ysub(1))
            f_y2 = ws.cell(row=row, column=base + 1, value=_ysub(2))
            f_y3 = ws.cell(row=row, column=base + 2, value=_ysub(3))
            f_wavg = ws.cell(row=row, column=base + 3, value=(
                f"=({y1}{row}*1+{y2}{row}*'Model Assumptions'!$C$20+"
                f"{y3}{row}*('Model Assumptions'!$C$20^2))/"
                f"(1+'Model Assumptions'!$C$20+'Model Assumptions'!$C$20^2)"
            ))
            f_th = ws.cell(row=row, column=base + 4, value=(
                f"={y1}{row}*'Model Assumptions'!$C$22+{wavg}{row}*(1-'Model Assumptions'!$C$22)"
            ))
            f_lb = ws.cell(row=row, column=base + 5, value=(
                f"=INDEX(${sec2_col}${season_rows[HISTORICAL_YEARS[0]]}:"
                f"${sec2_col}${season_rows[HISTORICAL_YEARS[-1]]},"
                f"MATCH('Model Assumptions'!$C$18-1,$A${season_rows[HISTORICAL_YEARS[0]]}:"
                f"$A${season_rows[HISTORICAL_YEARS[-1]]},0))"
            ))
            f_pb = ws.cell(row=row, column=base + 6, value=(
                f"={th}{row}*'Model Assumptions'!$C$21+{lb}{row}*(1-'Model Assumptions'!$C$21)"
            ))
            for cell in (f_y1, f_y2, f_y3, f_wavg, f_th, f_lb, f_pb):
                cell.font = FORMULA_FONT
                cell.number_format = m["fmt"]

    # ==== Section 4: League average/std-dev of Section 3's Projected Baselines =============
    sec4_title_row = sec3_last_row + 2
    sec4_header_row = sec4_title_row + 1
    avg_row, std_row = sec4_header_row + 1, sec4_header_row + 2

    # claude_code_spec_current_season_blending.md: blend Section 3's Projected Baseline
    # with a real Current-Season value, using the EXACT SAME Blend Weight mechanism Team
    # Ratings' own PPG blend already uses (Model Assumptions C12/C13/C14). Games Played is
    # cross-referenced from Team Ratings' own col H, not re-entered here.
    proj_baseline_col = {
        m["key"]: get_column_letter(metric_block_start_col[m["key"]] + 6) for m in METRICS
    }
    blended_cols = add_current_season_blend(
        ws, team_col="A", sec3_first_row=sec3_first_row, sec3_last_row=sec3_last_row,
        start_col=sec3_last_col + 1,
        metrics=[
            {"key": m["key"], "label": m["label"], "pb_col": proj_baseline_col[m["key"]]}
            for m in METRICS
        ],
    )

    _section_title(
        ws, sec4_title_row, 3,
        "Section 4 \u2014 League Average & Std. Dev. of the 3-Yr, Current-Season-Blended "
        "Baselines (see the new columns appended to the right of Section 3)",
    )
    _header_row(ws, sec4_header_row, ["Stat", *[m["label"] for m in METRICS]], height=18)

    ws.cell(row=avg_row, column=1, value="League Average").font = FORMULA_FONT
    ws.cell(row=std_row, column=1, value="League Std. Dev.").font = FORMULA_FONT

    for j, m in enumerate(METRICS):
        pcol = blended_cols[m["key"]]
        pb_range = f"{pcol}${sec3_first_row}:{pcol}${sec3_last_row}"
        a = ws.cell(row=avg_row, column=2 + j, value=f"=AVERAGE({pb_range})")
        s = ws.cell(row=std_row, column=2 + j, value=f"=STDEVP({pb_range})")
        for cell in (a, s):
            cell.font = FORMULA_FONT
            cell.number_format = m["fmt"]

    # ==== Section 5: Z-scores, weighted composite, points-scale Team Secondary Score =======
    sec5_title_row = std_row + 2
    sec5_header_row = sec5_title_row + 1
    sec5_first_row = sec5_header_row + 1
    sec5_last_row = sec5_first_row + n_teams - 1

    _section_title(
        ws, sec5_title_row, 5,
        "Section 5 \u2014 Z-Scores and Team Secondary Score (both metrics are \"higher is "
        "better\" -- no sign-flip needed. Baseline/points-per-SD reuse QB Index's own "
        "Model Assumptions cells C37/C38.)",
    )
    _header_row(
        ws, sec5_header_row,
        ["Team", *[f"{m['label']}\nZ" for m in METRICS], "Weighted\nZ-Score Sum",
         "Team Secondary\nScore (Points)", "Years of\nReal History"],
    )

    avg_cell_ref = {
        m["key"]: f"${get_column_letter(2 + j)}${avg_row}" for j, m in enumerate(METRICS)
    }
    std_cell_ref = {
        m["key"]: f"${get_column_letter(2 + j)}${std_row}" for j, m in enumerate(METRICS)
    }
    weight_cells = {"int": "$C$74", "pbu": "$C$75"}

    for i in range(n_teams):
        sec3_row = sec3_first_row + i
        row = sec5_first_row + i
        f = ws.cell(row=row, column=1, value=f"=A{sec3_row}")
        f.font = FORMULA_FONT

        z_cols = []
        for j, m in enumerate(METRICS):
            pcol = blended_cols[m["key"]]
            formula = f"=({pcol}{sec3_row}-{avg_cell_ref[m['key']]})/{std_cell_ref[m['key']]}"
            col = 2 + j
            cell = ws.cell(row=row, column=col, value=formula)
            cell.font = FORMULA_FONT
            cell.number_format = "0.00"
            z_cols.append(get_column_letter(col))

        weighted_terms = " + ".join(
            f"{z_cols[j]}{row}*'Model Assumptions'!{weight_cells[m['key']]}"
            for j, m in enumerate(METRICS)
        )
        wz_col = 2 + len(METRICS)
        wz = ws.cell(row=row, column=wz_col, value=f"={weighted_terms}")
        wz.font = FORMULA_FONT
        wz.number_format = "0.00"

        score = ws.cell(row=row, column=wz_col + 1, value=(
            f"='Model Assumptions'!$C$37+{get_column_letter(wz_col)}{row}*'Model Assumptions'!$C$38"
        ))
        score.font = FORMULA_FONT
        score.number_format = "0.0;(0.0)"

        yh_link = ws.cell(row=row, column=wz_col + 2, value=f"=B{sec3_row}")
        yh_link.font = FORMULA_FONT
        yh_link.number_format = "0"

    sec5_score_col = get_column_letter(2 + len(METRICS) + 1)
    sec5_score_range = f"${sec5_score_col}${sec5_first_row}:${sec5_score_col}${sec5_last_row}"
    sec5_team_range = f"$A${sec5_first_row}:$A${sec5_last_row}"

    # ==== Section 6: Individual Real Production (informational) ============================
    sec6_title_row = sec5_last_row + 2
    sec6_header_row = sec6_title_row + 1
    sec6_first_row = sec6_header_row + 1
    n_starters = len(starters)
    sec6_last_row = sec6_first_row + n_starters - 1

    _section_title(
        ws, sec6_title_row, 7,
        "Section 6 \u2014 Individual Real Production (the current real starters at CB1/"
        "CB2/NB/FS/SS -- current_roster.py's live 2026 depth-chart pull, a direct fixed "
        "mapping since all 5 spots are always populated). INT/PBU are THAT PLAYER'S OWN "
        "real 2025 season totals (defense_stats.py) -- not a grade, and not fed back into "
        "Section 1-5 scoring (already captured there via the team rollup).",
    )
    _header_row(ws, sec6_header_row, [
        "Team", "Slot", "Player Name", "Player ID", "2025 INT", "2025 PBU",
        "Team Secondary\nScore (Points)",
    ], height=24)

    for i, r in enumerate(starters.to_dict("records")):
        row = sec6_first_row + i
        team = r["Team"]
        pid = r["Player ID"]
        ws.cell(row=row, column=1, value=team).font = INPUT_FONT
        ws.cell(row=row, column=2, value=r["Role"]).font = INPUT_FONT
        ws.cell(row=row, column=3, value=r["Player Name"]).font = INPUT_FONT
        ws.cell(row=row, column=4, value=pid).font = INPUT_FONT

        recent = player_stats[
            (player_stats["Player ID"] == pid) & (player_stats["Season"] == HISTORICAL_YEARS[-1])
        ]
        int_v = float(recent["INT"].iloc[0]) if len(recent) else None
        pbu_v = float(recent["PBU"].iloc[0]) if len(recent) else None

        for col, v in ((5, int_v), (6, pbu_v)):
            cell = ws.cell(row=row, column=col, value=v)
            cell.font = INPUT_FONT
            cell.number_format = "0.0"

        team_score = ws.cell(row=row, column=7, value=(
            f'=IFERROR(INDEX({sec5_score_range},MATCH(A{row},{sec5_team_range},0)),"")'
        ))
        team_score.font = FORMULA_FONT
        team_score.number_format = "0.0;(0.0)"

    # ==== Section 7: Team Ratings adjustment (direct, unconditional) =======================
    sec7_title_row = sec6_last_row + 2
    sec7_header_row = sec7_title_row + 1
    sec7_first_row = sec7_header_row + 1
    sec7_last_row = sec7_first_row + len(TEAM_ORDER) - 1

    _section_title(
        ws, sec7_title_row, 4,
        "Section 7 \u2014 Secondary Adjustment (pts) -- a DIRECT quality measure (Team "
        "Secondary Score minus the league-average baseline of 50, converted to game "
        "points), NOT a Replacement Value swap -- there's no 'backup secondary' to swap "
        "in as a unit.",
    )
    _header_row(ws, sec7_header_row, [
        "Team", "Secondary Adjustment\n(Index Points)", "Secondary Adjustment\n(Game Points)",
    ])

    for i, team in enumerate(TEAM_ORDER):
        row = sec7_first_row + i
        ws.cell(row=row, column=1, value=team).font = INPUT_FONT

        adj_index = ws.cell(row=row, column=2, value=(
            f'=IFERROR(INDEX({sec5_score_range},MATCH(A{row},{sec5_team_range},0))'
            f"-'Model Assumptions'!$C$37,\"\")"
        ))
        adj_game = ws.cell(row=row, column=3, value=(
            f'=IF(B{row}="","",B{row}*\'Model Assumptions\'!$C$76)'
        ))
        adj_index.font = FORMULA_FONT
        adj_game.font = FORMULA_FONT
        adj_index.number_format = "0.0;(0.0)"
        adj_game.number_format = "0.00;(0.00)"

    note_row = sec7_last_row + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
    note = ws.cell(row=note_row, column=1, value=(
        "NO INVENTED PER-POSITION GRADE APPEARS ANYWHERE ON THIS TAB. Section 1-5 use REAL, "
        "individually-attributed nflverse play-by-play counting stats "
        "(interception_player_id, pass_defense_1-2_player_id), rolled up to team-level "
        "rates and run through the same decay-weighted/regressed/Z-scored chain as every "
        "other tab. Section 6 shows the current real starters at CB1/CB2 (LCB/RCB), NB "
        "(nickel), FS, and SS -- verified live that all 5 real depth-chart positions are "
        "populated for all 32 teams, always, so this is a direct fixed mapping (no "
        "scheme-variation priority order needed, unlike Front Seven Index's interior "
        "line/linebacker groups). Their INT/PBU are THAT PLAYER'S OWN real 2025 totals -- "
        "informational only, not re-fed into scoring. Together with Front Seven & D-Line "
        "Index, this completes the defense's counting-stats coverage per the roadmap's own "
        "guidance. NOT INCLUDED: coverage-shell scheme charting (single-high/two-high) has "
        "no free public source at all -- same limitation noted on Front Seven Index. The "
        "2026 depth-chart snapshot Section 6's starters are pulled from was pulled BEFORE "
        "final 53-man roster cuts -- same caveat as every other current-roster-driven tab "
        "in this workbook."
    ))
    note.font = NOTE_FONT
    note.alignment = Alignment(wrap_text=True, vertical="top")

    # ==== Wire into Team Ratings (new, unconditional additive column) ======================
    tr = wb[TEAM_RATINGS_SHEET]
    tr.cell(row=2, column=22, value="Secondary\nAdjustment (pts)").font = HEADER_FONT
    tr.cell(row=2, column=22).fill = HEADER_FILL
    tr.cell(row=2, column=22).alignment = HEADER_ALIGN

    adj_game_range = f"'{SHEET_NAME}'!$C${sec7_first_row}:$C${sec7_last_row}"
    adj_team_range = f"'{SHEET_NAME}'!$A${sec7_first_row}:$A${sec7_last_row}"

    for row in range(3, 3 + len(TEAM_ORDER)):
        adj = tr.cell(row=row, column=22, value=(
            f"=IFERROR(INDEX({adj_game_range},MATCH(A{row},{adj_team_range},0)),0)"
        ))
        adj.font = LINK_FONT
        adj.number_format = "0.00;(0.00)"

        # Net Power Rating (N) now also includes the Secondary Adjustment, additive
        # alongside QB (P), RB (R), Kicking (S), OL (T), and Front 7 (U) -- none disturb
        # each other. This is now the MOST COMPLETE version of the N formula -- see
        # main.py's own ordering comment.
        net = tr.cell(row=row, column=14, value=(
            f"=J{row}-K{row}+L{row}+M{row}+P{row}+R{row}+S{row}+T{row}+U{row}+V{row}"
        ))
        net.font = FORMULA_FONT
        net.number_format = "0.0;(0.0)"

    note_row_tr = 3 + len(TEAM_ORDER) + 5
    tr.merge_cells(start_row=note_row_tr, start_column=1, end_row=note_row_tr, end_column=22)
    tr_note = tr.cell(row=note_row_tr, column=1, value=(
        "Secondary Adjustment (V) is unconditional, same reasoning as every other "
        "position-group tab's direct adjustment -- there's no 'backup secondary' to "
        "switch to. It pulls directly from 'Secondary Index' Section 7 (that team's real, "
        "Z-scored Secondary Score minus the league-average baseline, converted to game "
        "points) and applies to Net Power Rating (N) for every team, every time, alongside "
        "whatever QB/RB/Kicking/OL/Front-7 adjustments are also active."
    ))
    tr_note.font = NOTE_FONT
    tr_note.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(workbook_path)
    print(
        f"Built '{SHEET_NAME}': Section 1 {len(team_stats)} rows, Section 3/5 {n_teams} "
        f"teams, Section 6 {n_starters} starters, Section 7 {len(TEAM_ORDER)} teams. "
        f"Wired into '{TEAM_RATINGS_SHEET}' (col V)."
    )
    print(f"Saved to {workbook_path}")
    return {
        "sec1_rows": len(team_stats), "n_teams": n_teams, "n_starters": n_starters,
        "sec3_range": (sec3_first_row, sec3_last_row),
        "sec5_range": (sec5_first_row, sec5_last_row),
        "sec6_range": (sec6_first_row, sec6_last_row),
        "sec7_range": (sec7_first_row, sec7_last_row),
    }


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print('Usage: uv run python scripts/build_secondary_index.py "path/to/workbook.xlsx"')
        sys.exit(1)
    build(sys.argv[1])
