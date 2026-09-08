"""
Step 4 extension: extends the Step 1 core-formula manifest (v35_core_formula_components.csv,
Season Matchups' own Z/AA) to every weighted metric across every position/matchup index
tab's own "Section 5" (the real Z-score-and-weighted-composite section every one of these
tabs uses, confirmed identical in structure across 16 sheets before writing this).

Mechanical, not memory-based: for each real "Section 5" title found, reads the real header
row (metric names, e.g. "EPA/Play\nZ") and the real "Weighted Z-Score Sum" column's own
formula at the first real data row, parses out each `<col><row>*'Model Assumptions'!$C$<n>`
term via regex, and zips the parsed (column, Model Assumptions row) pairs against the header
labels at those same columns -- in the SAME order the formula itself lists them, not assumed.
Cross-references each Model Assumptions row against v35_model_assumptions_manifest.csv (Step
1's own output) for the real Coefficient/Label/Note, so a change to that manifest doesn't
silently drift out of sync with this one.

Read-only throughout -- never touches the frozen workbook or the live one.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
WORKBOOK = HERE / "frozen_baselines" / "NFL_Prediction_Model_v35.xlsx"
MA_MANIFEST = HERE / "manifests" / "v35_model_assumptions_manifest.csv"
OUT_PATH = HERE / "manifests" / "v35_all_weighted_components.csv"

TERM_RE = re.compile(r"([A-Z]+)(\d+)\*'Model Assumptions'!\$C\$(\d+)")


def load_ma_manifest() -> dict[int, dict]:
    out = {}
    with open(MA_MANIFEST, encoding="utf-8") as f:
        for row in csv.DictReader(f):
            if row["Label"]:
                out[int(row["Row"])] = row
    return out


def find_section5(ws) -> int | None:
    for r in range(1, min(ws.max_row, 700) + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v.strip().startswith("Section 5"):
            return r
    return None


def extract_sheet(ws, ma: dict[int, dict]) -> list[dict]:
    title_row = find_section5(ws)
    if title_row is None:
        return []
    header_row = title_row + 1
    first_row = header_row + 1

    # Find the "Weighted...Z-Score Sum" header column (may be phrased slightly differently
    # per tab, e.g. "Weighted\nZ-Score Sum" everywhere observed so far).
    weighted_col = None
    for c in range(1, ws.max_column + 1):
        h = ws.cell(row=header_row, column=c).value
        if isinstance(h, str) and "Weighted" in h and "Z-Score" in h:
            weighted_col = c
            break
    if weighted_col is None:
        return []

    formula = ws.cell(row=first_row, column=weighted_col).value
    if not isinstance(formula, str):
        return []

    terms = TERM_RE.findall(formula)
    if not terms:
        return []

    rows_out = []
    for col_letter, _row_num, ma_row_str in terms:
        col_idx = openpyxl.utils.column_index_from_string(col_letter)
        header_label = ws.cell(row=header_row, column=col_idx).value
        header_label = str(header_label).replace("\n", " ") if header_label else ""
        ma_row = int(ma_row_str)
        ma_entry = ma.get(ma_row, {})
        rows_out.append({
            "Worksheet": ws.title,
            "Metric_Z_Column": col_letter,
            "Metric_Label": header_label,
            "Coefficient_Cell": f"Model Assumptions!C{ma_row}",
            "Coefficient_Value": ma_entry.get("Value", ""),
            "Coefficient_Label": ma_entry.get("Label", ""),
            "Coefficient_Note": ma_entry.get("Note", ""),
            "Section5_First_Data_Row": first_row,
            "Weighted_Sum_Formula": formula,
        })
    return rows_out


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK, data_only=False)
    ma = load_ma_manifest()

    all_rows = []
    sheets_no_match = []
    for name in wb.sheetnames:
        ws = wb[name]
        if find_section5(ws) is None:
            continue
        rows = extract_sheet(ws, ma)
        if rows:
            all_rows.extend(rows)
        else:
            sheets_no_match.append(name)

    with open(OUT_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=[
            "Worksheet", "Metric_Z_Column", "Metric_Label", "Coefficient_Cell",
            "Coefficient_Value", "Coefficient_Label", "Coefficient_Note",
            "Section5_First_Data_Row", "Weighted_Sum_Formula",
        ])
        writer.writeheader()
        writer.writerows(all_rows)

    print(f"Extracted {len(all_rows)} real weighted-metric components across "
          f"{len({r['Worksheet'] for r in all_rows})} worksheets -> {OUT_PATH}")
    if sheets_no_match:
        print(f"Sheets with a 'Section 5' title but no parseable weighted-sum formula "
              f"(needs manual review): {sheets_no_match}")


if __name__ == "__main__":
    main()
