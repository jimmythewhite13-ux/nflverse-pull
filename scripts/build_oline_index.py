"""
Builds "Offensive Line Index" -- addresses Phase 5 of the multi-phase roadmap in
claude_code_spec_rb_index.md, which explicitly said this phase was "not attempted; no
public per-player data exists to grade individual linemen honestly -- would need PFF-style
charting, which is paid/proprietary." This build does NOT contradict that finding: it does
not grade individual linemen. It uses REAL, free, per-TEAM data (Pro Football Reference,
via oline_stats.py) for the actual grading, and REAL per-player roster data (NFL
experience, via current_roster.py) for the individual-level component -- no fabricated
per-lineman skill number appears anywhere on this tab. See oline_stats.py's and
current_roster.py's own module docstrings for the full reasoning.

This was adapted from a user-supplied draft engine (SuperOffensiveLineEngine) that used
invented per-lineman grades (LT/LG/C/RG/RT numeric scores, Pass_Protection/Run_Blocking as
free-form inputs) as its foundation -- exactly the fabricated-data pattern this project has
never used anywhere else. What's kept from that draft: the decay-weighted multi-year
baseline, Z-score-to-points-scale conversion, and positional-weighting concept (LT/RT
emphasized) -- all real techniques, just applied to real inputs instead of invented ones.
What's NOT kept (yet): the opponent-matchup Z-score / Pressure Risk / Effective_QB_Index
layer (would need real weekly opponent pass-rush data this project doesn't pull yet) and
the sklearn-fitted matchup coefficient (needs real historical game-margin backtest data
this pipeline doesn't have) -- both are real, legitimate ideas worth a future extension,
just out of scope for this build.

UPDATED: added a third real metric, Sack-Free Rate (Fault-Adjusted), per explicit user
instruction to incorporate FTN Fantasy's real per-play charting (already accessible via
nfl_data_py.import_ftn_data). PFR's Pass_Protection blames the O-line for every sack, even
ones charted as the QB's own fault -- FTN's real is_qb_fault_sack column (verified live:
100% coverage on 2025's real sack plays, 34.6% charted as QB-fault) lets this tab correct
for that. See oline_stats.py's updated docstring for the full reasoning.

Section 1: raw 3-year TEAM-level data (Pass_Protection, Run_Blocking from Pro Football
           Reference; Sack-Free Rate (Fault-Adjusted) from real nflverse pbp + FTN
           charting) from oline_stats.py -- genuinely team-level, not per-player, unlike
           every other position in this project (there is no honest "this specific left
           tackle's own stat line").
Section 2: league average per season (simple average across all 32 teams -- there's no
           per-player Role/threshold filter to apply to a team-level stat)
Section 3: per-team 3-Yr decay-weighted, regressed baseline -- population is all 32 teams
           directly; a missing year substitutes THAT season's own Section 2 league average
           (not a flat "Rookie Baseline" -- there's no rookie concept for a team stat)
Section 4: league average/std-dev of Section 3's projected baselines
Section 5: Z-scores, weighted composite, points-scale Team OL Score (reuses QB Index's own
           Model Assumptions C37/C38 scale constants, same design choice as every other tab)
Section 6: Individual Starters -- the 5 real current-roster starters per team (LT/LG/C/RG/
           RT, current_roster.py), each with REAL NFL experience (current_roster.
           attach_experience) and a Rookie-Starter risk modifier applied to their team's
           real Section 5 score (a true rookie's first NFL start, especially at Center, is
           a genuinely different real risk profile -- see the Model Assumptions notes)
Section 7: Team-Level Positional-Weighted Blend -- combines the 5 individual
           rookie-adjusted scores per team using real positional-importance weights
           (LT/RT emphasized, adapted from the user's draft engine) into one Team Blended
           OL Score, then a direct (not Replacement-Value-swap) Team Ratings adjustment,
           same unconditional pattern as Kicking Index (no "backup line" to swap in)

Usage:
    uv run python scripts/build_oline_index.py "C:\\path\\to\\NFL_Prediction_Model.xlsx"
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from add_manual_override_table import read_existing_overrides  # noqa: E402

from nflverse_pull.current_roster import (  # noqa: E402
    attach_experience,
    compute_current_starters,
    fetch_depth_charts,
    fetch_seasonal_rosters,
    resolve_scored_population,
)
from nflverse_pull.efficiency import fetch_pbp  # noqa: E402
from nflverse_pull.oline_stats import (  # noqa: E402
    compute_team_season_oline_stats,
    compute_team_season_sack_fault_stats,
    fetch_ftn,
    fetch_pfr_pass,
    fetch_pfr_rush,
)

HISTORICAL_YEARS = [2023, 2024, 2025]
CURRENT_ROSTER_YEAR = 2026
SHEET_NAME = "Offensive Line Index"
KICKING_SHEET = "Kicking Index"
WR_TE_SHEET = "WR-TE Value Index"
RB_INDEX_SHEET = "RB Value Index"
QB_INDEX_SHEET = "QB Index"
TEAM_RATINGS_SHEET = "Team Ratings"
OL_POSITIONS = ["LT", "LG", "C", "RG", "RT"]

METRICS = [
    {"key": "pp", "col": "Pass_Protection", "label": "Pass\nProtection", "fmt": "0.0"},
    {"key": "rb", "col": "Run_Blocking", "label": "Run\nBlocking", "fmt": "0.00"},
    {"key": "sfa", "col": "Sack-Free Rate (Fault-Adjusted)", "label": "Sack-Free\nRate",
     "fmt": "0.00%"},
]

TEAM_ORDER = [
    "Buffalo Bills", "Miami Dolphins", "New England Patriots", "New York Jets",
    "Baltimore Ravens", "Cincinnati Bengals", "Cleveland Browns", "Pittsburgh Steelers",
    "Houston Texans", "Indianapolis Colts", "Jacksonville Jaguars", "Tennessee Titans",
    "Denver Broncos", "Kansas City Chiefs", "Las Vegas Raiders", "Los Angeles Chargers",
    "Dallas Cowboys", "New York Giants", "Philadelphia Eagles", "Washington Commanders",
    "Chicago Bears", "Detroit Lions", "Green Bay Packers", "Minnesota Vikings",
    "Atlanta Falcons", "Carolina Panthers", "New Orleans Saints", "Tampa Bay Buccaneers",
    "Arizona Cardinals", "Los Angeles Rams", "San Francisco 49ers", "Seattle Seahawks",
]

TITLE_FONT = Font(name="Arial", size=10, bold=True)
TITLE_FILL = PatternFill("solid", fgColor="FFD9E1F2")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_ALIGN = Alignment(wrap_text=True, horizontal="center", vertical="center")
INPUT_FONT = Font(name="Arial", size=10, color="FF0000FF")
FORMULA_FONT = Font(name="Arial", size=10, color="FF000000")
LINK_FONT = Font(name="Arial", size=10, color="FF008000")
NOTE_FONT = Font(name="Arial", size=9, color="FF808080")
ASSUMPTION_FILL = PatternFill("solid", fgColor="FFFFFF00")


def _section_title(ws, row: int, last_col: int, text: str) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL


def _header_row(ws, row: int, headers: list[str], height: float = 28) -> None:
    ws.row_dimensions[row].height = height
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN


def add_model_assumptions_weights(wb: openpyxl.Workbook) -> None:
    ws = wb["Model Assumptions"]
    title_row = 57
    ws.merge_cells(f"A{title_row}:D{title_row}")
    title = ws.cell(row=title_row, column=1, value=(
        "Offensive Line Index Weighting & Conversion (pts per std. dev.; reuses QB Index's "
        "points-scale constants C37/C38 -- see 'Offensive Line Index' tab)"
    ))
    title.font = HEADER_FONT
    title.fill = HEADER_FILL

    rows = [
        (58, "Pass Protection Weight (OL Index, pts per SD)", 0.4,
         "Real, team-level pressure-rate-allowed proxy (Pro Football Reference) -- "
         "weighted down from this tab's original 0.5 now that Sack-Free Rate "
         "(Fault-Adjusted, C81) exists as a more precise, fault-attributed real "
         "signal specifically for sacks; Pass_Protection still captures pressure "
         "beyond just sacks (hits, hurries), so it stays the largest single weight."),
        (59, "Run Blocking Weight (OL Index, pts per SD)", 0.35,
         "Real, team-level Yards-Before-Contact-per-Attempt proxy (Pro Football "
         "Reference) -- weighted down from 0.5 for the same 3-metric rebalancing "
         "as C58."),
        (60, "Rookie Starter Penalty -- LT/LG/RG/RT (pts)", 1.0,
         "Applied to a real current starter's individual score when he's a true "
         "rookie (0 years NFL experience, from real roster data) -- a documented, "
         "honest proxy for continuity risk, NOT a fabricated skill grade. A guess, "
         "like every other coefficient in this model."),
        (61, "Rookie Starter Penalty -- Center (pts)", 2.0,
         "Larger than the other 4 spots -- the center makes protection calls and "
         "communicates the whole line's assignments, so a true rookie there is a "
         "meaningfully bigger real risk than at guard/tackle. Same proxy logic as "
         "C60, just weighted higher for this one spot."),
        (62, "LT Positional Weight (Team Blend)", 0.22,
         "Blind-side protector for a right-handed QB (the large majority) -- "
         "weighted highest alongside RT."),
        (63, "LG Positional Weight (Team Blend)", 0.17, ""),
        (64, "C Positional Weight (Team Blend)", 0.20,
         "Communication hub for the whole line -- weighted above LG/RG."),
        (65, "RG Positional Weight (Team Blend)", 0.17, ""),
        (66, "RT Positional Weight (Team Blend)", 0.24,
         "Weighted highest of all 5 -- faces the opponent's best pass rusher most "
         "often in most modern defensive fronts."),
        (67, "OL Index Points-to-Game-Points Conversion", 0.06,
         "A starting guess, like every other coefficient in this model. This is a "
         "DIRECT quality adjustment, not a Replacement Value swap -- there's no "
         "'backup offensive line' to swap in, same reasoning as Kicking Index's C56."),
        (81, "Sack-Free Rate (Fault-Adjusted) Weight (OL Index, pts per SD)", 0.25,
         "Added after the fact once FTN Fantasy's real per-play charting "
         "(is_qb_fault_sack) was found -- appended here rather than inserted next "
         "to C58/C59 to avoid shifting every row reference every later tab's build "
         "script already hardcodes (Front Seven/Secondary/Special Teams start at "
         "C68/C73/C77). Real, fault-attributed sack rate: excludes sacks charted "
         "as the QB's own fault from the O-line's blame -- weighted lower than "
         "Pass Protection/Run Blocking since it's a narrower, partially-redundant "
         "slice of what Pass_Protection already measures."),
    ]
    for row, label, value, note in rows:
        ws.cell(row=row, column=2, value=label)
        c = ws.cell(row=row, column=3, value=value)
        c.font = INPUT_FONT
        c.fill = ASSUMPTION_FILL
        c.number_format = "0.00"
        n = ws.cell(row=row, column=4, value=note)
        n.font = NOTE_FONT
        n.alignment = Alignment(wrap_text=True, vertical="top")


def _pull_data(overrides: pd.DataFrame) -> dict:
    print(f"Pulling {HISTORICAL_YEARS} Pro Football Reference data for team OL stats...")
    pfr_pass = fetch_pfr_pass(HISTORICAL_YEARS)
    pfr_rush = fetch_pfr_rush(HISTORICAL_YEARS)
    season_stats = compute_team_season_oline_stats(pfr_pass, pfr_rush)

    print(f"Pulling {HISTORICAL_YEARS} play-by-play + FTN charting for fault-adjusted "
          "sack rate...")
    pbp = fetch_pbp(HISTORICAL_YEARS)
    ftn = fetch_ftn(HISTORICAL_YEARS)
    sack_fault_stats = compute_team_season_sack_fault_stats(pbp, ftn)
    season_stats = season_stats.merge(sack_fault_stats, on=["Team", "Season"], how="outer")

    print(f"Pulling {CURRENT_ROSTER_YEAR} depth charts for the 5 current-roster OL starters...")
    depth_charts = fetch_depth_charts([CURRENT_ROSTER_YEAR])
    current_starters = compute_current_starters(depth_charts)
    starters = pd.concat(
        [resolve_scored_population(current_starters, overrides, pos) for pos in OL_POSITIONS],
        ignore_index=True,
    )
    print(f"{len(starters)} current-roster OL starters identified across "
          f"{len(OL_POSITIONS)} spots.")

    print(f"Pulling {CURRENT_ROSTER_YEAR} roster data for real NFL experience...")
    rosters = fetch_seasonal_rosters([CURRENT_ROSTER_YEAR])
    starters = attach_experience(starters, rosters)
    n_unknown = starters["Years of NFL Experience"].isna().sum()
    n_rookie = (starters["Is Rookie"] == True).sum()  # noqa: E712
    print(f"{n_rookie} true-rookie starters identified (0 years experience); "
          f"{n_unknown} starter(s) missing from the roster snapshot (experience unknown).")

    return {"season_stats": season_stats, "starters": starters}


def build(workbook_path: str) -> dict:
    wb = openpyxl.load_workbook(workbook_path)
    overrides = read_existing_overrides(wb, SHEET_NAME, OL_POSITIONS)
    print(f"Read back {len(overrides)} existing manual-override row(s) from Section 7 "
          "before rebuilding the sheet.")

    data = _pull_data(overrides)
    season_stats = data["season_stats"]
    starters = data["starters"]

    add_model_assumptions_weights(wb)

    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    for candidate in (KICKING_SHEET, WR_TE_SHEET, RB_INDEX_SHEET, QB_INDEX_SHEET):
        if candidate in wb.sheetnames:
            insert_after = candidate
            break
    else:
        insert_after = wb.sheetnames[0]
    ws = wb.create_sheet(SHEET_NAME, index=wb.sheetnames.index(insert_after) + 1)
    ws.column_dimensions["A"].width = 20.0
    ws.column_dimensions["C"].width = 20.0

    ws.merge_cells("A1:H1")
    t = ws.cell(row=1, column=1, value=(
        "Offensive Line Index -- Multi-Year Decay-Weighted TEAM-Level Blocking Rating "
        "(Pass Protection and Run Blocking from Pro Football Reference; Sack-Free Rate "
        "Fault-Adjusted from real nflverse pbp + FTN Fantasy charting). NO individual "
        "lineman is graded anywhere on this tab -- see the closing note for why, and how "
        "the individual-level section instead uses real NFL experience."
    ))
    t.font = Font(name="Arial", size=12, bold=True)

    # ==== Section 1: Raw 3-year TEAM-level data =============================================
    sec1_first_row = 5
    sec1_last_row = sec1_first_row + len(season_stats) - 1
    _section_title(
        ws, 3, 2 + len(METRICS),
        "Section 1 \u2014 Raw 3-Year TEAM-Level Data (Pass Protection/Run Blocking: Pro "
        "Football Reference; Sack-Free Rate (Fault-Adjusted): real nflverse pbp + FTN "
        "Fantasy charting -- genuinely team-level, not per-player; see this tab's opening "
        "note)",
    )
    _header_row(ws, 4, ["Team", "Season", *[m["label"] for m in METRICS]])
    for i, r in enumerate(season_stats.to_dict("records")):
        row = sec1_first_row + i
        values = [r["Team"], int(r["Season"])] + [
            (float(r[m["col"]]) if pd.notna(r.get(m["col"])) else None) for m in METRICS
        ]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.font = INPUT_FONT
        for j, m in enumerate(METRICS):
            ws.cell(row=row, column=3 + j).number_format = m["fmt"]

    team_range = f"$A${sec1_first_row}:$A${sec1_last_row}"
    season_range = f"$B${sec1_first_row}:$B${sec1_last_row}"
    sec1_col_of = {"pp": "C", "rb": "D", "sfa": "E"}
    metric_ranges = {
        m["key"]: (
            f"${sec1_col_of[m['key']]}${sec1_first_row}:${sec1_col_of[m['key']]}${sec1_last_row}"
        )
        for m in METRICS
    }

    # ==== Section 2: League average per season (simple average across 32 teams) ============
    sec2_title_row = sec1_last_row + 2
    sec2_header_row = sec2_title_row + 1
    season_rows = {yr: sec2_header_row + 1 + i for i, yr in enumerate(HISTORICAL_YEARS)}

    _section_title(
        ws, sec2_title_row, 1 + len(METRICS),
        "Section 2 \u2014 League Average per Season (simple average across all 32 teams -- "
        "no per-player Role filter applies to a team-level stat)",
    )
    _header_row(ws, sec2_header_row, ["Season", *[m["label"] for m in METRICS]], height=20)

    for yr in HISTORICAL_YEARS:
        row = season_rows[yr]
        c = ws.cell(row=row, column=1, value=yr)
        c.font = INPUT_FONT
        for j, m in enumerate(METRICS):
            formula = f"=AVERAGEIF({season_range},{yr},{metric_ranges[m['key']]})"
            cell = ws.cell(row=row, column=2 + j, value=formula)
            cell.font = FORMULA_FONT
            cell.number_format = m["fmt"]

    season_avg_col = {m["key"]: get_column_letter(2 + j) for j, m in enumerate(METRICS)}

    # ==== Section 3: Per-team decay-weighted, regressed baseline ===========================
    sec3_title_row = season_rows[HISTORICAL_YEARS[-1]] + 2
    sec3_header_row = sec3_title_row + 1
    sec3_first_row = sec3_header_row + 1
    n_teams = len(TEAM_ORDER)
    sec3_last_row = sec3_first_row + n_teams - 1
    sec3_last_col = 2 + len(METRICS) * 7

    _section_title(
        ws, sec3_title_row, sec3_last_col,
        "Section 3 \u2014 Per-Team 3-Yr Decay-Weighted, Regressed Baseline. A missing year "
        "substitutes THAT season's own Section 2 league average (not a flat 'rookie' "
        "concept -- there's no rookie for a team-level stat).",
    )
    headers = ["Team", "Years of\nReal History"]
    metric_block_start_col: dict[str, int] = {}
    col_cursor = 3
    for m in METRICS:
        metric_block_start_col[m["key"]] = col_cursor
        headers += [
            f"{m['label']} Y-1", f"{m['label']} Y-2", f"{m['label']} Y-3",
            f"Weighted\n{m['label']} Avg\n(3-Yr decay)", f"Team History\n{m['label']}",
            f"League Baseline\n{m['label']} (Y-1)", f"Projected 3-Yr\n{m['label']} Baseline",
        ]
        col_cursor += 7
    _header_row(ws, sec3_header_row, headers)

    for i, team in enumerate(TEAM_ORDER):
        row = sec3_first_row + i
        ws.cell(row=row, column=1, value=team).font = FORMULA_FONT

        years_hist_terms = "+".join(
            f"--(COUNTIFS({team_range},$A{row},{season_range},'Model Assumptions'!$C$18-{k})>0)"
            for k in (1, 2, 3)
        )
        yh = ws.cell(row=row, column=2, value=f"={years_hist_terms}")
        yh.font = FORMULA_FONT
        yh.number_format = "0"

        for m in METRICS:
            base = metric_block_start_col[m["key"]]
            y1, y2, y3, wavg, th, lb, pb = (get_column_letter(base + k) for k in range(7))
            mrange = metric_ranges[m["key"]]
            sec2_col = season_avg_col[m["key"]]

            def _ysub(offset: int, mrange=mrange, sec2_col=sec2_col) -> str:
                season_avg_lookup = (
                    f"INDEX(${sec2_col}${season_rows[HISTORICAL_YEARS[0]]}:"
                    f"${sec2_col}${season_rows[HISTORICAL_YEARS[-1]]},"
                    f"MATCH('Model Assumptions'!$C$18-{offset},"
                    f"$A${season_rows[HISTORICAL_YEARS[0]]}:"
                    f"$A${season_rows[HISTORICAL_YEARS[-1]]},0))"
                )
                return (
                    f"=IF(COUNTIFS({team_range},$A{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset})=0,{season_avg_lookup},"
                    f"SUMIFS({mrange},{team_range},$A{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset}))"
                )

            f_y1 = ws.cell(row=row, column=base, value=_ysub(1))
            f_y2 = ws.cell(row=row, column=base + 1, value=_ysub(2))
            f_y3 = ws.cell(row=row, column=base + 2, value=_ysub(3))
            f_wavg = ws.cell(row=row, column=base + 3, value=(
                f"=({y1}{row}*1+{y2}{row}*'Model Assumptions'!$C$20+"
                f"{y3}{row}*('Model Assumptions'!$C$20^2))/"
                f"(1+'Model Assumptions'!$C$20+'Model Assumptions'!$C$20^2)"
            ))
            f_th = ws.cell(row=row, column=base + 4, value=(
                f"={y1}{row}*'Model Assumptions'!$C$22+{wavg}{row}*(1-'Model Assumptions'!$C$22)"
            ))
            f_lb = ws.cell(row=row, column=base + 5, value=(
                f"=INDEX(${sec2_col}${season_rows[HISTORICAL_YEARS[0]]}:"
                f"${sec2_col}${season_rows[HISTORICAL_YEARS[-1]]},"
                f"MATCH('Model Assumptions'!$C$18-1,$A${season_rows[HISTORICAL_YEARS[0]]}:"
                f"$A${season_rows[HISTORICAL_YEARS[-1]]},0))"
            ))
            f_pb = ws.cell(row=row, column=base + 6, value=(
                f"={th}{row}*'Model Assumptions'!$C$21+{lb}{row}*(1-'Model Assumptions'!$C$21)"
            ))
            for cell in (f_y1, f_y2, f_y3, f_wavg, f_th, f_lb, f_pb):
                cell.font = FORMULA_FONT
                cell.number_format = m["fmt"]

    # ==== Section 4: League average/std-dev of Section 3's Projected Baselines =============
    sec4_title_row = sec3_last_row + 2
    sec4_header_row = sec4_title_row + 1
    avg_row, std_row = sec4_header_row + 1, sec4_header_row + 2

    _section_title(
        ws, sec4_title_row, 1 + len(METRICS),
        "Section 4 \u2014 League Average & Std. Dev. of the 3-Yr Baselines",
    )
    _header_row(ws, sec4_header_row, ["Stat", *[m["label"] for m in METRICS]], height=18)

    ws.cell(row=avg_row, column=1, value="League Average").font = FORMULA_FONT
    ws.cell(row=std_row, column=1, value="League Std. Dev.").font = FORMULA_FONT

    proj_baseline_col = {
        m["key"]: get_column_letter(metric_block_start_col[m["key"]] + 6) for m in METRICS
    }
    for j, m in enumerate(METRICS):
        pcol = proj_baseline_col[m["key"]]
        pb_range = f"{pcol}${sec3_first_row}:{pcol}${sec3_last_row}"
        a = ws.cell(row=avg_row, column=2 + j, value=f"=AVERAGE({pb_range})")
        s = ws.cell(row=std_row, column=2 + j, value=f"=STDEVP({pb_range})")
        for cell in (a, s):
            cell.font = FORMULA_FONT
            cell.number_format = m["fmt"]

    # ==== Section 5: Z-scores, weighted composite, points-scale Team OL Score ==============
    sec5_title_row = std_row + 2
    sec5_header_row = sec5_title_row + 1
    sec5_first_row = sec5_header_row + 1
    sec5_last_row = sec5_first_row + n_teams - 1

    _section_title(
        ws, sec5_title_row, 1 + len(METRICS) + 3,
        "Section 5 \u2014 Z-Scores and Team OL Score (all three metrics are \"higher is "
        "better\" -- no sign-flip needed. Baseline/points-per-SD reuse QB Index's own Model "
        "Assumptions cells C37/C38.)",
    )
    _header_row(
        ws, sec5_header_row,
        ["Team", *[f"{m['label']}\nZ" for m in METRICS], "Weighted\nZ-Score Sum",
         "Team OL\nScore (Points)", "Years of\nReal History"],
    )

    avg_cell_ref = {
        m["key"]: f"${get_column_letter(2 + j)}${avg_row}" for j, m in enumerate(METRICS)
    }
    std_cell_ref = {
        m["key"]: f"${get_column_letter(2 + j)}${std_row}" for j, m in enumerate(METRICS)
    }
    weight_cells = {"pp": "$C$58", "rb": "$C$59", "sfa": "$C$81"}

    for i in range(n_teams):
        sec3_row = sec3_first_row + i
        row = sec5_first_row + i
        f = ws.cell(row=row, column=1, value=f"=A{sec3_row}")
        f.font = FORMULA_FONT

        z_cols = []
        for j, m in enumerate(METRICS):
            pcol = proj_baseline_col[m["key"]]
            formula = f"=({pcol}{sec3_row}-{avg_cell_ref[m['key']]})/{std_cell_ref[m['key']]}"
            col = 2 + j
            cell = ws.cell(row=row, column=col, value=formula)
            cell.font = FORMULA_FONT
            cell.number_format = "0.00"
            z_cols.append(get_column_letter(col))

        weighted_terms = " + ".join(
            f"{z_cols[j]}{row}*'Model Assumptions'!{weight_cells[m['key']]}"
            for j, m in enumerate(METRICS)
        )
        wz_col = 2 + len(METRICS)
        wz = ws.cell(row=row, column=wz_col, value=f"={weighted_terms}")
        wz.font = FORMULA_FONT
        wz.number_format = "0.00"

        score = ws.cell(row=row, column=wz_col + 1, value=(
            f"='Model Assumptions'!$C$37+{get_column_letter(wz_col)}{row}*'Model Assumptions'!$C$38"
        ))
        score.font = FORMULA_FONT
        score.number_format = "0.0;(0.0)"

        yh_link = ws.cell(row=row, column=wz_col + 2, value=f"=B{sec3_row}")
        yh_link.font = FORMULA_FONT
        yh_link.number_format = "0"

    sec5_score_col = get_column_letter(2 + len(METRICS) + 1)
    sec5_score_range = f"${sec5_score_col}${sec5_first_row}:${sec5_score_col}${sec5_last_row}"
    sec5_team_range = f"$A${sec5_first_row}:$A${sec5_last_row}"

    # ==== Section 6: Individual Starters (real identity + real experience) =================
    sec6_title_row = sec5_last_row + 2
    sec6_header_row = sec6_title_row + 1
    sec6_first_row = sec6_header_row + 1
    n_starters = len(starters)
    sec6_last_row = sec6_first_row + n_starters - 1

    _section_title(
        ws, sec6_title_row, 8,
        "Section 6 \u2014 Individual Starters (the 5 real current-roster starters per team, "
        "current_roster.py's live 2026 depth-chart pull). Years of NFL Experience and Is "
        "Rookie are REAL data (current_roster.attach_experience) -- NOT a skill grade. "
        "Individual Effective Score applies a real-rookie-status modifier to the team's "
        "real Section 5 score; it does not represent this specific player's own "
        "performance, since no honest per-player data exists for that (see this tab's "
        "opening note).",
    )
    _header_row(ws, sec6_header_row, [
        "Team", "Position", "Player Name", "Player ID", "Years of NFL\nExperience",
        "Is Rookie\nStarter", "Team OL\nScore (Points)", "Individual\nEffective Score",
    ], height=24)

    for i, r in enumerate(starters.to_dict("records")):
        row = sec6_first_row + i
        team = r["Team"]
        ws.cell(row=row, column=1, value=team).font = INPUT_FONT
        ws.cell(row=row, column=2, value=r["Role"]).font = INPUT_FONT
        ws.cell(row=row, column=3, value=r["Player Name"]).font = INPUT_FONT
        ws.cell(row=row, column=4, value=r["Player ID"]).font = INPUT_FONT

        years_exp = r.get("Years of NFL Experience")
        exp_cell = ws.cell(
            row=row, column=5, value=int(years_exp) if pd.notna(years_exp) else None
        )
        exp_cell.font = INPUT_FONT

        is_rookie = r.get("Is Rookie")
        rookie_cell = ws.cell(
            row=row, column=6, value=bool(is_rookie) if is_rookie is not None else None
        )
        rookie_cell.font = INPUT_FONT

        team_score = ws.cell(row=row, column=7, value=(
            f'=IFERROR(INDEX({sec5_score_range},MATCH(A{row},{sec5_team_range},0)),"")'
        ))
        team_score.font = FORMULA_FONT
        team_score.number_format = "0.0;(0.0)"

        penalty_cell = (
            "'Model Assumptions'!$C$61" if r["Role"] == "C" else "'Model Assumptions'!$C$60"
        )
        eff_score = ws.cell(row=row, column=8, value=(
            f'=IF(G{row}="","",IF(F{row}=TRUE,G{row}-{penalty_cell},G{row}))'
        ))
        eff_score.font = FORMULA_FONT
        eff_score.number_format = "0.0;(0.0)"

    # ==== Section 7: Team-Level Positional-Weighted Blend + Team Ratings adjustment ========
    sec7_title_row = sec6_last_row + 2
    sec7_header_row = sec7_title_row + 1
    sec7_first_row = sec7_header_row + 1
    sec7_last_row = sec7_first_row + len(TEAM_ORDER) - 1

    _section_title(
        ws, sec7_title_row, 5,
        "Section 7 \u2014 Team-Level Positional-Weighted Blend (LT/RT emphasized, adapted "
        "from the draft engine's positional-importance concept -- applied to the real, "
        "rookie-adjusted scores above, not to any invented per-lineman grade) and the "
        "resulting Team Ratings adjustment. This is a DIRECT quality measure, NOT a "
        "Replacement Value swap -- there's no 'backup offensive line' to swap in.",
    )
    _header_row(ws, sec7_header_row, [
        "Team", "Blended OL\nScore (Points)", "OL Adjustment\n(Index Points)",
        "OL Adjustment\n(Game Points)",
    ])

    pos_weight_cell = {
        "LT": "$C$62", "LG": "$C$63", "C": "$C$64", "RG": "$C$65", "RT": "$C$66",
    }
    key_range = f"$I${sec6_first_row}:$I${sec6_last_row}"
    eff_range = f"$H${sec6_first_row}:$H${sec6_last_row}"
    for r in range(sec6_first_row, sec6_last_row + 1):
        ws.cell(row=r, column=9, value=f'=A{r}&"|"&B{r}').font = FORMULA_FONT

    for i, team in enumerate(TEAM_ORDER):
        row = sec7_first_row + i
        ws.cell(row=row, column=1, value=team).font = INPUT_FONT

        terms = []
        for pos in OL_POSITIONS:
            key = f'"{team}|{pos}"'
            terms.append(
                f'IFERROR(INDEX({eff_range},MATCH({key},{key_range},0)),0)*'
                f"'Model Assumptions'!{pos_weight_cell[pos]}"
            )
        blend = ws.cell(row=row, column=2, value="=" + "+".join(terms))
        blend.font = FORMULA_FONT
        blend.number_format = "0.0;(0.0)"

        adj_index = ws.cell(row=row, column=3, value=(
            f"=B{row}-'Model Assumptions'!$C$37"
        ))
        adj_game = ws.cell(row=row, column=4, value=(
            f"=C{row}*'Model Assumptions'!$C$67"
        ))
        adj_index.font = FORMULA_FONT
        adj_game.font = FORMULA_FONT
        adj_index.number_format = "0.0;(0.0)"
        adj_game.number_format = "0.00;(0.00)"

    note_row = sec7_last_row + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
    note = ws.cell(row=note_row, column=1, value=(
        "NO INDIVIDUAL LINEMAN IS GRADED ANYWHERE ON THIS TAB. Section 1-5 use REAL, free, "
        "TEAM-level data: Pass Protection (= 100 - the team's attempt-weighted pressure "
        "rate allowed) and Run Blocking (= the team's attempt-weighted Yards Before "
        "Contact per Attempt) from Pro Football Reference; Sack-Free Rate (Fault-Adjusted) "
        "from real nflverse play-by-play joined against FTN Fantasy's real per-play "
        "charting (is_qb_fault_sack) -- excludes sacks charted as the QB's own fault "
        "(held the ball too long, etc.) from the O-line's blame, verified live to have "
        "100% real coverage on 2025's actual sack plays. There is no honest free per-"
        "lineman performance grade anywhere (that's PFF's proprietary domain; confirmed "
        "live that PFF IDs exist in nflverse's own player crosswalk for cross-referencing "
        "only, not grade values). Section 6's Individual Effective Score applies a "
        "real-rookie-status-based modifier (Model Assumptions C60/C61) to the TEAM's "
        "real score -- it "
        "reflects real continuity risk (a true rookie's first NFL start, especially at "
        "Center, is a genuinely different risk profile), not this specific player's own "
        "invented skill number. Section 7's positional weights (LT/RT emphasized) are a "
        "real, tunable design choice applied to those real, rookie-adjusted numbers, "
        "adapted from a user-supplied draft engine's positional-importance concept -- "
        "unlike that draft, no invented input feeds them. NOT YET BUILT: an opponent-"
        "matchup Z-score / Pressure Risk layer (would need real weekly opponent pass-rush "
        "data this project doesn't pull yet) and an empirically-fitted matchup coefficient "
        "(would need real historical game-margin backtest data this pipeline doesn't have) "
        "-- both are legitimate ideas from that same draft, left for a future extension. "
        "The 2026 depth-chart snapshot Section 6's starters are pulled from was pulled "
        "BEFORE final 53-man roster cuts -- same caveat as every other current-roster-"
        "driven tab in this workbook."
    ))
    note.font = NOTE_FONT
    note.alignment = Alignment(wrap_text=True, vertical="top")

    # ==== Wire into Team Ratings (new, unconditional additive column) ======================
    tr = wb[TEAM_RATINGS_SHEET]
    tr.cell(row=2, column=20, value="OL\nAdjustment (pts)").font = HEADER_FONT
    tr.cell(row=2, column=20).fill = HEADER_FILL
    tr.cell(row=2, column=20).alignment = HEADER_ALIGN

    adj_game_range = f"'{SHEET_NAME}'!$D${sec7_first_row}:$D${sec7_last_row}"
    adj_team_range = f"'{SHEET_NAME}'!$A${sec7_first_row}:$A${sec7_last_row}"

    for row in range(3, 3 + len(TEAM_ORDER)):
        adj = tr.cell(row=row, column=20, value=(
            f"=IFERROR(INDEX({adj_game_range},MATCH(A{row},{adj_team_range},0)),0)"
        ))
        adj.font = LINK_FONT
        adj.number_format = "0.00;(0.00)"

        # Net Power Rating (N) now also includes the OL Adjustment, additive alongside QB
        # (P), RB (R), and Kicking (S) -- none of the four disturb each other. This is now
        # the MOST COMPLETE version of the N formula -- see main.py's own ordering comment.
        net = tr.cell(row=row, column=14, value=(
            f"=J{row}-K{row}+L{row}+M{row}+P{row}+R{row}+S{row}+T{row}"
        ))
        net.font = FORMULA_FONT
        net.number_format = "0.0;(0.0)"

    note_row_tr = 3 + len(TEAM_ORDER) + 3
    tr.merge_cells(start_row=note_row_tr, start_column=1, end_row=note_row_tr, end_column=20)
    tr_note = tr.cell(row=note_row_tr, column=1, value=(
        "OL Adjustment (T) is unconditional, same reasoning as Kicking Adjustment (S) -- "
        "there's no 'backup offensive line' to switch to. It pulls directly from "
        "'Offensive Line Index' Section 7 (that team's real, rookie-adjusted, positionally-"
        "weighted Blended OL Score minus the league-average baseline, converted to game "
        "points) and applies to Net Power Rating (N) for every team, every time, alongside "
        "whatever QB/RB/Kicking adjustments are also active."
    ))
    tr_note.font = NOTE_FONT
    tr_note.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(workbook_path)
    print(
        f"Built '{SHEET_NAME}': Section 1 {len(season_stats)} rows, Section 3/5 {n_teams} "
        f"teams, Section 6 {n_starters} starters, Section 7 {len(TEAM_ORDER)} teams. "
        f"Wired into '{TEAM_RATINGS_SHEET}' (col T)."
    )
    print(f"Saved to {workbook_path}")
    return {
        "sec1_rows": len(season_stats), "n_teams": n_teams, "n_starters": n_starters,
        "sec3_range": (sec3_first_row, sec3_last_row),
        "sec5_range": (sec5_first_row, sec5_last_row),
        "sec6_range": (sec6_first_row, sec6_last_row),
        "sec7_range": (sec7_first_row, sec7_last_row),
    }


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print('Usage: uv run python scripts/build_oline_index.py "path/to/workbook.xlsx"')
        sys.exit(1)
    build(sys.argv[1])
