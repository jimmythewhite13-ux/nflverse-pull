"""
Shared helper for claude_code_spec_current_season_blending.md -- applies the EXACT SAME
mechanism already proven for Team Ratings' own PPG blend (Model Assumptions C12/C13/C14,
'Current-Season Blend Weight' -- Base / Added per Additional Game / Maximum Cap) to every
other position/team tab's own Section 3 3-Yr Baseline, instead of inventing a tab-specific
variant. Not a new blend-weight formula -- the literal same shape, reusing the literal same
three Model Assumptions cells, verified live against Team Ratings' own real I3 formula
before writing this: `=IF(H3=0,0,MIN($C$14,$C$12+($C$13*(H3-1))))`.

Games Played is NOT re-entered per tab as a new manual input -- every tab this applies to
already has a real Team column, and Team Ratings' own Games Played (col H) is already the
single real source of truth for "how many games has this team played this season" that the
user already maintains weekly. Cross-referencing it here (INDEX/MATCH by Team, same pattern
used everywhere else in this project) avoids a second, potentially-drifting manual entry
per tab -- a stronger form of "reuse the exact mechanism" than merely reusing the formula
shape. A player-level tab (QB/RB/WR-TE/Kicking) uses his OWN team's real Games Played, not a
separate per-player count -- his sample size for "how much real current-season data exists"
is fundamentally a function of how many games have been played this season, not whether he
personally suited up for all of them.

Call `add_current_season_blend(...)` once per Section-3-populated tab, right after that
Section 3 is written and before Section 5 is built -- it appends new columns to Section 3
(never inserts mid-block) and returns the Blended-value column letter per metric key, which
the caller's own Section 5 Z-score formulas should reference INSTEAD OF the old Projected
3-Yr Baseline column (Section 4's own League Average/Std. Dev. must also be recomputed over
the NEW Blended column, not the old Baseline column, for the Z-score to reflect blended
values).
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

INPUT_FONT = Font(name="Arial", size=10, color="FF0000FF")
FORMULA_FONT = Font(name="Arial", size=10, color="FF000000")
LINK_FONT = Font(name="Arial", size=10, color="FF008000")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_ALIGN = Alignment(wrap_text=True, horizontal="center", vertical="center")
NOTE_FONT = Font(name="Arial", size=9, color="FF808080")


def add_current_season_blend(
    ws,
    team_col: str,
    sec3_first_row: int,
    sec3_last_row: int,
    start_col: int,
    metrics: list[dict],
) -> dict[str, str]:
    """
    `metrics`: list of {"key": ..., "label": ..., "pb_col": "<Section 3 Projected Baseline
    column letter>"} -- one entry per metric this tab's Section 3 already blends 3-Yr-wise.

    Appends, starting at `start_col`: Games Played (ref, green, INDEX/MATCH into Team
    Ratings col H) | Current-Season Blend Weight (black formula, the literal Team Ratings
    I-column shape) | per metric: Current Season <label> (blue input, starts 0) | Blended
    <label> (black formula).

    Returns {metric_key: blended_column_letter}.
    """
    header_row = ws.cell(row=sec3_first_row - 1, column=1).row  # Section 3's own header row
    gp_col = start_col
    bw_col = start_col + 1
    gp_letter = get_column_letter(gp_col)
    bw_letter = get_column_letter(bw_col)

    hdr_gp = ws.cell(row=header_row, column=gp_col, value="Games Played\n(Current Szn, ref)")
    hdr_bw = ws.cell(row=header_row, column=bw_col, value="Current-Season\nBlend Weight")
    for h in (hdr_gp, hdr_bw):
        h.font = HEADER_FONT
        h.fill = HEADER_FILL
        h.alignment = HEADER_ALIGN

    col_cursor = start_col + 2
    blended_cols: dict[str, str] = {}
    metric_headers: dict[str, tuple[int, int]] = {}
    for m in metrics:
        cur_col = col_cursor
        blend_col = col_cursor + 1
        metric_headers[m["key"]] = (cur_col, blend_col)
        h1 = ws.cell(row=header_row, column=cur_col, value=f"Current Season\n{m['label']}")
        h2 = ws.cell(row=header_row, column=blend_col, value=f"Blended\n{m['label']}")
        for h in (h1, h2):
            h.font = HEADER_FONT
            h.fill = HEADER_FILL
            h.alignment = HEADER_ALIGN
        blended_cols[m["key"]] = get_column_letter(blend_col)
        col_cursor += 2

    for row in range(sec3_first_row, sec3_last_row + 1):
        gp = ws.cell(row=row, column=gp_col, value=(
            f"=IFERROR(INDEX('Team Ratings'!$H$3:$H$34,"
            f"MATCH({team_col}{row},'Team Ratings'!$A$3:$A$34,0)),0)"
        ))
        gp.font = LINK_FONT
        gp.number_format = "0"

        bw = ws.cell(row=row, column=bw_col, value=(
            f"=IF({gp_letter}{row}=0,0,MIN('Model Assumptions'!$C$14,"
            f"'Model Assumptions'!$C$12+('Model Assumptions'!$C$13*({gp_letter}{row}-1))))"
        ))
        bw.font = FORMULA_FONT
        bw.number_format = "0.00"

        for m in metrics:
            cur_col, blend_col = metric_headers[m["key"]]
            cur_letter = get_column_letter(cur_col)
            cur = ws.cell(row=row, column=cur_col, value=0)
            cur.font = INPUT_FONT
            cur.number_format = "0.000"

            blend = ws.cell(row=row, column=blend_col, value=(
                f"={m['pb_col']}{row}*(1-{bw_letter}{row})+{cur_letter}{row}*{bw_letter}{row}"
            ))
            blend.font = FORMULA_FONT
            blend.number_format = "0.000"

    return blended_cols
