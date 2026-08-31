"""
Builds "EDGE/IDL Index" -- Part B of claude_code_spec_defensive_player_index.md. Real,
individual-player, rate-based (per real defensive snap, NOT raw counting totals) pass-rush
production for the 4 real front-line slots every team's current depth chart identifies:
EDGE1 (real LDE), EDGE2 (real RDE), IDL1/IDL2 (scheme-variable, real LDT/RDT/NT priority
order -- the SAME priority Front Seven Index's own Section 6 already uses).

HONESTY REQUIREMENT, stated plainly here and on the tab itself: these are real counting
stats (sacks, TFL, QB hits), genuinely attributed to individual players in public data --
but they are noisier and more scheme/teammate-dependent than QB Index's EPA-based metrics.
An EDGE rusher's sack total is meaningfully influenced by the other three real pass
rushers and the defensive scheme, in a way a QB's EPA/play isn't influenced by his
receivers to the same degree. This tab does NOT carry the same implied confidence as QB
Index, and says so on its own closing note.

This does NOT replace Front Seven Index (team-level, still Built & Verified, unchanged) --
it's an additional, complementary layer: individual player scoring feeding a real per-SLOT
Replacement Value, the same relationship QB Index has to Advanced Efficiency Metrics'
team-level passing numbers.

Real data pipeline (defense_stats.py, see that module's own docstring for the full real-
data reasoning): compute_player_season_front7_stats() (real per-player-per-season Sacks/
TFL/QB Hits, already given 3 years of pbp) -> compute_player_season_defensive_rates()
(joins real per-player-season snap counts, excludes below-MIN_QUALIFYING_DEFENSIVE_SNAPS
seasons entirely, converts to real per-snap rates, attaches the REAL per-season rookie flag
from the start).

NO Section 2B (Individual Rookie Assumptions): checked live before building this -- the
api.fantasycalc.com dynasty-value dataset used for QB/RB/WR-TE's rookie crosswalk carries
ZERO defensive-position entries at all (WR/RB/QB/TE/PICK only), so there is no real market
signal to build a Section 2B from for any defensive position. A zero-history rookie at any
slot uses the flat Section 2 Rookie Baseline directly, same as Kicking Index's own
precedent for exactly this situation.

Section 3/5's population is every REAL player who could appear as either a Starter or
Backup at ANY of the 4 real slots (current_roster.resolve_slot_with_backup /
resolve_priority_slots_with_backup) -- both sides of each pair need their own real 3-Yr
decay-weighted Score for Section 6's subtraction to mean anything.

Section 6: A SEPARATE Replacement Value pair PER REAL SLOT (EDGE1 vs its own real backup,
EDGE2 vs its own, IDL1 vs its own, IDL2 vs its own) -- confirmed with the user before
building this way: these positions already have 2+ real "starters" apiece unlike QB/RB, so
a single team-wide Starter/Backup pair would be arbitrary. Summed into ONE Team Ratings
adjustment column (matching every other tab's "one column per tab" convention), while each
slot's own Replacement Value stays independently visible in Section 6.

Usage:
    uv run python scripts/build_edge_idl_index.py "C:\\path\\to\\NFL_Prediction_Model.xlsx"
"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from add_manual_override_table import read_existing_overrides  # noqa: E402

from nflverse_pull.current_roster import (  # noqa: E402
    POSITION_ROLE_LABELS,
    compute_current_starters,
    fetch_depth_charts,
    fetch_seasonal_rosters,
    resolve_priority_slots_with_backup,
    resolve_slot_with_backup,
)
from nflverse_pull.defense_stats import (  # noqa: E402
    compute_player_season_defensive_rates,
    compute_player_season_defensive_snaps,
    compute_player_season_front7_stats,
    fetch_player_ids,
    fetch_snap_counts,
)
from nflverse_pull.efficiency import fetch_pbp  # noqa: E402

HISTORICAL_YEARS = [2023, 2024, 2025]
CURRENT_ROSTER_YEAR = 2026
# Excel sheet names can't contain "/" (same real constraint every other tab already works
# around, e.g. "WR-TE Value Index" not "WR/TE").
SHEET_NAME = "EDGE-IDL Index"
FRONT_SEVEN_SHEET = "Front Seven Index"
TEAM_RATINGS_SHEET = "Team Ratings"

EDGE_MAP = [("LDE", "EDGE1"), ("RDE", "EDGE2")]
IDL_PRIORITY = ["LDT", "RDT", "NT"]
IDL_MAX_SLOTS = 2

METRICS = [
    {"key": "sack", "col": "Sacks Rate", "label": "Sack Rate\n(per snap)", "fmt": "0.000"},
    {"key": "tfl", "col": "TFL Rate", "label": "TFL Rate\n(per snap)", "fmt": "0.000"},
    {"key": "hit", "col": "QB Hits Rate", "label": "QB Hit Rate\n(per snap)", "fmt": "0.000"},
]

TEAM_ORDER = [
    "Buffalo Bills", "Miami Dolphins", "New England Patriots", "New York Jets",
    "Baltimore Ravens", "Cincinnati Bengals", "Cleveland Browns", "Pittsburgh Steelers",
    "Houston Texans", "Indianapolis Colts", "Jacksonville Jaguars", "Tennessee Titans",
    "Denver Broncos", "Kansas City Chiefs", "Las Vegas Raiders", "Los Angeles Chargers",
    "Dallas Cowboys", "New York Giants", "Philadelphia Eagles", "Washington Commanders",
    "Chicago Bears", "Detroit Lions", "Green Bay Packers", "Minnesota Vikings",
    "Atlanta Falcons", "Carolina Panthers", "New Orleans Saints", "Tampa Bay Buccaneers",
    "Arizona Cardinals", "Los Angeles Rams", "San Francisco 49ers", "Seattle Seahawks",
]

TITLE_FONT = Font(name="Arial", size=10, bold=True)
TITLE_FILL = PatternFill("solid", fgColor="FFD9E1F2")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_ALIGN = Alignment(wrap_text=True, horizontal="center", vertical="center")
INPUT_FONT = Font(name="Arial", size=10, color="FF0000FF")
FORMULA_FONT = Font(name="Arial", size=10, color="FF000000")
LINK_FONT = Font(name="Arial", size=10, color="FF008000")
NOTE_FONT = Font(name="Arial", size=9, color="FF808080")
ASSUMPTION_FILL = PatternFill("solid", fgColor="FFFFFF00")


def _section_title(ws, row: int, last_col: int, text: str) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = TITLE_FONT
    cell.fill = TITLE_FILL


def _header_row(ws, row: int, headers: list[str], height: float = 28) -> None:
    ws.row_dimensions[row].height = height
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN


def add_model_assumptions_weights(wb: openpyxl.Workbook) -> None:
    ws = wb["Model Assumptions"]
    title_row = 103
    ws.merge_cells(f"A{title_row}:D{title_row}")
    title = ws.cell(row=title_row, column=1, value=(
        "EDGE/IDL Index Weighting & Conversion (pts per std. dev.; reuses QB Index's "
        "points-scale constants C37/C38 -- see 'EDGE-IDL Index' tab)"
    ))
    title.font = HEADER_FONT
    title.fill = HEADER_FILL

    rows = [
        (104, "Sack Rate Weight (EDGE/IDL Index, pts per SD)", 0.5,
         "Real, individually-credited sacks per real defensive snap -- weighted "
         "highest as the clearest single individual pass-rush production signal, "
         "same logic as Front Seven Index's own team-level weighting."),
        (105, "TFL Rate Weight (EDGE/IDL Index, pts per SD)", 0.25,
         "Real run-and-pass tackles for loss per real defensive snap."),
        (106, "QB Hit Rate Weight (EDGE/IDL Index, pts per SD)", 0.25,
         "Real QB hits per real defensive snap -- correlates with Sack Rate but "
         "captures pressure that didn't finish as a sack."),
        (107, "EDGE/IDL Index Replacement Value Points-to-Game-Points Conversion", 0.05,
         "A starting guess, like every other coefficient in this model -- a "
         "structurally different, dedicated constant, NOT reused from QB's/RB's own "
         "conversion factors. Smaller than QB's/RB's: these stats are noisier and "
         "more scheme/teammate-dependent (see this tab's own closing note), so a "
         "given point gap here should move the prediction less than the same gap "
         "in QB/RB Index."),
    ]
    for row, label, value, note in rows:
        ws.cell(row=row, column=2, value=label)
        c = ws.cell(row=row, column=3, value=value)
        c.font = INPUT_FONT
        c.fill = ASSUMPTION_FILL
        c.number_format = "0.00"
        n = ws.cell(row=row, column=4, value=note)
        n.font = NOTE_FONT
        n.alignment = Alignment(wrap_text=True, vertical="top")


def _resolve_slots(overrides: pd.DataFrame) -> pd.DataFrame:
    """Team | Slot | Starter Name | Starter Player ID | Backup Name | Backup Player ID,
    for EDGE1/EDGE2 (fixed) and IDL1/IDL2 (scheme-variable priority)."""
    print(f"Pulling {CURRENT_ROSTER_YEAR} depth charts for EDGE/IDL slots...")
    depth_charts = fetch_depth_charts([CURRENT_ROSTER_YEAR])
    current_starters = compute_current_starters(depth_charts)

    frames = [
        resolve_slot_with_backup(current_starters, overrides, real_pos, slot)
        for real_pos, slot in EDGE_MAP
    ]
    frames.append(
        resolve_priority_slots_with_backup(
            current_starters, overrides, IDL_PRIORITY, "IDL", IDL_MAX_SLOTS
        )
    )
    return pd.concat(frames, ignore_index=True)


def _slots_to_population(slots: pd.DataFrame) -> pd.DataFrame:
    """Melts each slot's Starter/Backup pair into two population rows (one per real
    player), Role = "{slot} Starter"/"{slot} Backup" -- both sides need their own real
    3-Yr decay-weighted Score for Section 6's subtraction to mean anything."""
    rows = []
    for r in slots.to_dict("records"):
        if pd.notna(r["Starter Player ID"]):
            rows.append({"Team": r["Team"], "Role": f"{r['Slot']} Starter",
                         "Player Name": r["Starter Name"], "Player ID": r["Starter Player ID"]})
        if pd.notna(r["Backup Player ID"]):
            rows.append({"Team": r["Team"], "Role": f"{r['Slot']} Backup",
                         "Player Name": r["Backup Name"], "Player ID": r["Backup Player ID"]})
    return pd.DataFrame(rows, columns=["Team", "Role", "Player Name", "Player ID"])


def build(workbook_path: str) -> dict:
    wb = openpyxl.load_workbook(workbook_path)
    # Override roles are the REAL underlying depth-chart positions (LDE/LDE2/RDE/RDE2/LDT/
    # LDT2/RDT/RDT2/NT/NT2), NOT the remapped EDGE1/IDL1 slot names -- resolve_scored_
    # population() (called inside resolve_slot_with_backup/resolve_priority_slots_with_
    # backup) only ever knows about real positions; EDGE_MAP/IDL_PRIORITY's remapping to
    # EDGE1/IDL1 happens AFTER an override already would have taken effect.
    override_roles = [
        role for real_pos, _ in EDGE_MAP for role in POSITION_ROLE_LABELS[real_pos].values()
    ] + [
        role for real_pos in IDL_PRIORITY for role in POSITION_ROLE_LABELS[real_pos].values()
    ]
    overrides = read_existing_overrides(wb, SHEET_NAME, override_roles)
    print(f"Read back {len(overrides)} existing manual-override row(s) from Section 7 "
          "before rebuilding the sheet.")

    print(f"Pulling {HISTORICAL_YEARS} play-by-play data for EDGE/IDL rate stats...")
    pbp = fetch_pbp(HISTORICAL_YEARS)
    raw_counts = compute_player_season_front7_stats(pbp)

    print(f"Pulling {HISTORICAL_YEARS} real snap counts + player-ID crosswalk...")
    snap_counts_raw = fetch_snap_counts(HISTORICAL_YEARS)
    player_ids = fetch_player_ids()
    snaps = compute_player_season_defensive_snaps(snap_counts_raw, player_ids)

    rosters = fetch_seasonal_rosters(HISTORICAL_YEARS)

    season_stats = compute_player_season_defensive_rates(
        raw_counts, ["Sacks", "TFL", "QB Hits"], snaps, rosters
    )
    print(f"{len(season_stats)} qualifying EDGE/IDL-style player-seasons "
          f"(>= real snap threshold, real rookie flag attached).")

    slots = _resolve_slots(overrides)
    population = _slots_to_population(slots)
    n_slots = len(slots)
    n_pop = len(population)
    print(f"{n_slots} real (team, slot) rows resolved; {n_pop} real players "
          f"(Starter+Backup) identified for scoring.")

    add_model_assumptions_weights(wb)

    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    insert_after = FRONT_SEVEN_SHEET if FRONT_SEVEN_SHEET in wb.sheetnames else wb.sheetnames[0]
    ws = wb.create_sheet(SHEET_NAME, index=wb.sheetnames.index(insert_after) + 1)
    ws.column_dimensions["A"].width = 20.0
    ws.column_dimensions["C"].width = 20.0

    ws.merge_cells("A1:H1")
    t = ws.cell(row=1, column=1, value=(
        "EDGE/IDL Index -- Real, Individual, Per-Snap-Rate Pass-Rush Production (Sack "
        "Rate, TFL Rate, QB Hit Rate) for the 4 real front-line slots (EDGE1/EDGE2/IDL1/"
        "IDL2). COMPLEMENTARY to Front Seven Index (team-level, unchanged) -- adds a real "
        "individual-player Replacement Value layer. NOISIER than QB Index -- see the "
        "closing note before treating these scores with the same confidence."
    ))
    t.font = Font(name="Arial", size=12, bold=True)

    # ==== Section 1: Raw 3-year per-player rate history =====================================
    sec1_first_row = 5
    sec1_last_row = sec1_first_row + len(season_stats) - 1
    _section_title(
        ws, 3, 7,
        "Section 1 — Raw 3-Year Per-Player Rate History (real nflverse pbp counting "
        "stats, real per-player snap counts -- see this tab's opening note). Below-"
        "threshold player-seasons are excluded entirely, not zero-filled.",
    )
    _header_row(ws, 4, [
        "Player ID", "Season", "Team", *[m["label"] for m in METRICS], "Is Rookie Season",
    ])
    for i, r in enumerate(season_stats.to_dict("records")):
        row = sec1_first_row + i
        values = [r["Player ID"], int(r["Season"]), r["Team"]]
        for m in METRICS:
            values.append(float(r[m["col"]]))
        values.append(bool(r["Is Rookie Season"]))
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=v)
            cell.font = INPUT_FONT
        for j, m in enumerate(METRICS):
            ws.cell(row=row, column=4 + j).number_format = m["fmt"]

    id_range = f"$A${sec1_first_row}:$A${sec1_last_row}"
    season_range = f"$B${sec1_first_row}:$B${sec1_last_row}"
    rookie_range = f"$G${sec1_first_row}:$G${sec1_last_row}"
    sec1_col_of = {m["key"]: get_column_letter(4 + j) for j, m in enumerate(METRICS)}
    metric_ranges = {
        m["key"]: (
            f"${sec1_col_of[m['key']]}${sec1_first_row}:${sec1_col_of[m['key']]}${sec1_last_row}"
        )
        for m in METRICS
    }

    # ==== Section 2: League average per season (qualifying player-seasons) + Rookie Baseline
    sec2_title_row = sec1_last_row + 2
    sec2_header_row = sec2_title_row + 1
    season_rows = {yr: sec2_header_row + 1 + i for i, yr in enumerate(HISTORICAL_YEARS)}
    rookie_baseline_row = sec2_header_row + 1 + len(HISTORICAL_YEARS)
    rookie_count_row = rookie_baseline_row + 1

    _section_title(
        ws, sec2_title_row, 4,
        "Section 2 — League Average per Season (all real qualifying player-seasons) + "
        "flat Rookie Baseline",
    )
    _header_row(ws, sec2_header_row, ["Season / Stat", *[m["label"] for m in METRICS]], height=20)

    for yr in HISTORICAL_YEARS:
        row = season_rows[yr]
        c = ws.cell(row=row, column=1, value=yr)
        c.font = INPUT_FONT
        for j, m in enumerate(METRICS):
            formula = f"=AVERAGEIF({season_range},{yr},{metric_ranges[m['key']]})"
            cell = ws.cell(row=row, column=2 + j, value=formula)
            cell.font = FORMULA_FONT
            cell.number_format = m["fmt"]

    lbl = ws.cell(row=rookie_baseline_row, column=1, value="Rookie Baseline (all years, flat)")
    lbl.font = FORMULA_FONT
    for j, m in enumerate(METRICS):
        formula = f"=AVERAGEIF({rookie_range},TRUE,{metric_ranges[m['key']]})"
        cell = ws.cell(row=rookie_baseline_row, column=2 + j, value=formula)
        cell.font = FORMULA_FONT
        cell.number_format = m["fmt"]

    lbl2 = ws.cell(
        row=rookie_count_row, column=1,
        value="Rookie Baseline Sample Size (qualifying rookie seasons) -- recomputes "
              "automatically as more seasons are added",
    )
    lbl2.font = NOTE_FONT
    lbl2.alignment = Alignment(wrap_text=True)
    for j in range(len(METRICS)):
        cell = ws.cell(row=rookie_count_row, column=2 + j, value=f"=COUNTIF({rookie_range},TRUE)")
        cell.font = FORMULA_FONT
        cell.number_format = "0"

    flat_rookie_cell = {
        m["key"]: f"${get_column_letter(2 + j)}${rookie_baseline_row}"
        for j, m in enumerate(METRICS)
    }
    season_avg_col = {m["key"]: get_column_letter(2 + j) for j, m in enumerate(METRICS)}

    # ==== Section 3: Per-player 3-Yr decay-weighted, regressed baseline ====================
    sec3_title_row = rookie_count_row + 2
    sec3_header_row = sec3_title_row + 1
    sec3_first_row = sec3_header_row + 1
    n_players = len(population)
    sec3_last_row = sec3_first_row + n_players - 1
    sec3_last_col = 5 + len(METRICS) * 7

    _section_title(
        ws, sec3_title_row, sec3_last_col,
        "Section 3 — Per-Player 3-Yr Decay-Weighted, Regressed Baseline. Population is "
        "every REAL player who could appear as either a Starter or Backup at ANY of the "
        "4 real slots (EDGE1/EDGE2/IDL1/IDL2) -- both sides of each pair are scored, so "
        "Section 6's subtraction means something. Missing years substitute the flat "
        "Section 2 Rookie Baseline directly -- no Section 2B (no real dynasty-fantasy "
        "market exists for any defensive position, verified live).",
    )
    headers = ["Player Name", "Player ID", "Team", "Role", "Years of\nReal History"]
    metric_block_start_col: dict[str, int] = {}
    col_cursor = 6
    for m in METRICS:
        metric_block_start_col[m["key"]] = col_cursor
        headers += [
            f"{m['label']} Y-1", f"{m['label']} Y-2", f"{m['label']} Y-3",
            f"Weighted\n{m['label']} Avg\n(3-Yr decay)", f"Team History\n{m['label']}",
            f"League Baseline\n{m['label']} (Y-1)", f"Projected 3-Yr\n{m['label']} Baseline",
        ]
        col_cursor += 7
    _header_row(ws, sec3_header_row, headers)

    for i, p in enumerate(population.to_dict("records")):
        row = sec3_first_row + i
        ws.cell(row=row, column=1, value=p["Player Name"]).font = FORMULA_FONT
        ws.cell(row=row, column=2, value=p["Player ID"]).font = FORMULA_FONT
        ws.cell(row=row, column=3, value=p["Team"]).font = FORMULA_FONT
        ws.cell(row=row, column=4, value=p["Role"]).font = FORMULA_FONT

        years_hist_terms = "+".join(
            f"--(COUNTIFS({id_range},$B{row},{season_range},'Model Assumptions'!$C$18-{k})>0)"
            for k in (1, 2, 3)
        )
        yh = ws.cell(row=row, column=5, value=f"={years_hist_terms}")
        yh.font = FORMULA_FONT
        yh.number_format = "0"

        for m in METRICS:
            base = metric_block_start_col[m["key"]]
            y1, y2, y3, wavg, th, lb, pb = (get_column_letter(base + k) for k in range(7))
            mrange = metric_ranges[m["key"]]
            flat_cell = flat_rookie_cell[m["key"]]

            def _ysub(offset: int, mrange=mrange, flat_cell=flat_cell) -> str:
                return (
                    f"=IF(COUNTIFS({id_range},$B{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset})=0,{flat_cell},"
                    f"SUMIFS({mrange},{id_range},$B{row},{season_range},"
                    f"'Model Assumptions'!$C$18-{offset}))"
                )

            f_y1 = ws.cell(row=row, column=base, value=_ysub(1))
            f_y2 = ws.cell(row=row, column=base + 1, value=_ysub(2))
            f_y3 = ws.cell(row=row, column=base + 2, value=_ysub(3))
            f_wavg = ws.cell(row=row, column=base + 3, value=(
                f"=({y1}{row}*1+{y2}{row}*'Model Assumptions'!$C$20+"
                f"{y3}{row}*('Model Assumptions'!$C$20^2))/"
                f"(1+'Model Assumptions'!$C$20+'Model Assumptions'!$C$20^2)"
            ))
            f_th = ws.cell(row=row, column=base + 4, value=(
                f"={y1}{row}*'Model Assumptions'!$C$22+{wavg}{row}*(1-'Model Assumptions'!$C$22)"
            ))
            sec2_col = season_avg_col[m["key"]]
            f_lb = ws.cell(row=row, column=base + 5, value=(
                f"=INDEX(${sec2_col}${season_rows[HISTORICAL_YEARS[0]]}:"
                f"${sec2_col}${season_rows[HISTORICAL_YEARS[-1]]},"
                f"MATCH('Model Assumptions'!$C$18-1,$A${season_rows[HISTORICAL_YEARS[0]]}:"
                f"$A${season_rows[HISTORICAL_YEARS[-1]]},0))"
            ))
            f_pb = ws.cell(row=row, column=base + 6, value=(
                f"={th}{row}*'Model Assumptions'!$C$21+{lb}{row}*(1-'Model Assumptions'!$C$21)"
            ))
            for cell in (f_y1, f_y2, f_y3, f_wavg, f_th, f_lb, f_pb):
                cell.font = FORMULA_FONT
                cell.number_format = m["fmt"]

    # ==== Section 4: League average/std-dev of Section 3's Projected Baselines =============
    sec4_title_row = sec3_last_row + 2
    sec4_header_row = sec4_title_row + 1
    avg_row, std_row = sec4_header_row + 1, sec4_header_row + 2

    _section_title(
        ws, sec4_title_row, 4, "Section 4 — League Average & Std. Dev. of the 3-Yr Baselines"
    )
    _header_row(ws, sec4_header_row, ["Stat", *[m["label"] for m in METRICS]], height=18)

    ws.cell(row=avg_row, column=1, value="League Average").font = FORMULA_FONT
    ws.cell(row=std_row, column=1, value="League Std. Dev.").font = FORMULA_FONT

    proj_baseline_col = {
        m["key"]: get_column_letter(metric_block_start_col[m["key"]] + 6) for m in METRICS
    }
    for j, m in enumerate(METRICS):
        pcol = proj_baseline_col[m["key"]]
        pb_range = f"{pcol}${sec3_first_row}:{pcol}${sec3_last_row}"
        a = ws.cell(row=avg_row, column=2 + j, value=f"=AVERAGE({pb_range})")
        s = ws.cell(row=std_row, column=2 + j, value=f"=STDEVP({pb_range})")
        for cell in (a, s):
            cell.font = FORMULA_FONT
            cell.number_format = m["fmt"]

    # ==== Section 5: Z-scores, weighted composite, points-scale EDGE/IDL Score =============
    sec5_title_row = std_row + 2
    sec5_header_row = sec5_title_row + 1
    sec5_first_row = sec5_header_row + 1
    sec5_last_row = sec5_first_row + n_players - 1

    _section_title(
        ws, sec5_title_row, 10,
        "Section 5 — Z-Scores and EDGE/IDL Index Score (all three metrics are \"higher "
        "is better\" -- no sign-flip needed. Baseline/points-per-SD reuse QB Index's own "
        "Model Assumptions cells C37/C38.)",
    )
    _header_row(
        ws, sec5_header_row,
        ["Player Name", "Player ID", "Team", "Role", *[f"{m['label']}\nZ" for m in METRICS],
         "Weighted\nZ-Score Sum", "EDGE/IDL Index\nScore (Points)", "Years of\nReal History",
         "Team|Role\n(helper)"],
    )

    avg_cell_ref = {
        m["key"]: f"${get_column_letter(2 + j)}${avg_row}" for j, m in enumerate(METRICS)
    }
    std_cell_ref = {
        m["key"]: f"${get_column_letter(2 + j)}${std_row}" for j, m in enumerate(METRICS)
    }
    weight_cells = {"sack": "$C$104", "tfl": "$C$105", "hit": "$C$106"}

    for i in range(n_players):
        sec3_row = sec3_first_row + i
        row = sec5_first_row + i
        for col, src_col in ((1, "A"), (2, "B"), (3, "C"), (4, "D")):
            f = ws.cell(row=row, column=col, value=f"={src_col}{sec3_row}")
            f.font = FORMULA_FONT

        z_cols = []
        for j, m in enumerate(METRICS):
            pcol = proj_baseline_col[m["key"]]
            formula = f"=({pcol}{sec3_row}-{avg_cell_ref[m['key']]})/{std_cell_ref[m['key']]}"
            col = 5 + j
            cell = ws.cell(row=row, column=col, value=formula)
            cell.font = FORMULA_FONT
            cell.number_format = "0.00"
            z_cols.append(get_column_letter(col))

        weighted_terms = " + ".join(
            f"{z_cols[j]}{row}*'Model Assumptions'!{weight_cells[m['key']]}"
            for j, m in enumerate(METRICS)
        )
        wz_col = 5 + len(METRICS)
        wz = ws.cell(row=row, column=wz_col, value=f"={weighted_terms}")
        wz.font = FORMULA_FONT
        wz.number_format = "0.00"

        score = ws.cell(row=row, column=wz_col + 1, value=(
            f"='Model Assumptions'!$C$37+{get_column_letter(wz_col)}{row}*'Model Assumptions'!$C$38"
        ))
        score.font = FORMULA_FONT
        score.number_format = "0.0;(0.0)"

        yh_link = ws.cell(row=row, column=wz_col + 2, value=f"=E{sec3_row}")
        yh_link.font = FORMULA_FONT
        yh_link.number_format = "0"

        key = ws.cell(row=row, column=wz_col + 3, value=f'=C{row}&"|"&D{row}')
        key.font = FORMULA_FONT

    sec5_score_col = get_column_letter(wz_col + 1)
    sec5_key_col = get_column_letter(wz_col + 3)
    score_range = f"${sec5_score_col}${sec5_first_row}:${sec5_score_col}${sec5_last_row}"
    key_range = f"${sec5_key_col}${sec5_first_row}:${sec5_key_col}${sec5_last_row}"

    # ==== Section 6: Replacement Value, per real slot =======================================
    sec6_title_row = sec5_last_row + 2
    sec6_header_row = sec6_title_row + 1
    sec6_first_row = sec6_header_row + 1
    sec6_last_row = sec6_first_row + n_slots - 1

    _section_title(
        ws, sec6_title_row, 8,
        "Section 6 — Replacement Value, PER REAL SLOT (Starter Score minus Backup Score, "
        "can be negative -- a negative value means the backup rates HIGHER, and the game-"
        "point adjustment flips sign accordingly, same requirement as QB/RB Index). A "
        "team missing a real backup at a slot shows blank Replacement Value, not a "
        "misleading 0.",
    )
    _header_row(ws, sec6_header_row, [
        "Team", "Slot", "Starter Name", "Starter Score", "Backup Name", "Backup Score",
        "Replacement Value\n(Index Points)", "Replacement Value\n(Game Points)",
    ])

    for i, s in enumerate(slots.to_dict("records")):
        row = sec6_first_row + i
        team, slot = s["Team"], s["Slot"]
        t = ws.cell(row=row, column=1, value=team)
        sl = ws.cell(row=row, column=2, value=slot)
        t.font = INPUT_FONT
        sl.font = INPUT_FONT

        starter_key = f'"{team}|{slot} Starter"'
        backup_key = f'"{team}|{slot} Backup"'
        starter_name = ws.cell(row=row, column=3, value=(
            f'=IFERROR(INDEX($A${sec5_first_row}:$A${sec5_last_row},'
            f'MATCH({starter_key},{key_range},0)),"")'
        ))
        starter_score = ws.cell(row=row, column=4, value=(
            f"=IFERROR(INDEX({score_range},MATCH({starter_key},{key_range},0)),\"\")"
        ))
        backup_name = ws.cell(row=row, column=5, value=(
            f'=IFERROR(INDEX($A${sec5_first_row}:$A${sec5_last_row},'
            f'MATCH({backup_key},{key_range},0)),"")'
        ))
        backup_score = ws.cell(row=row, column=6, value=(
            f"=IFERROR(INDEX({score_range},MATCH({backup_key},{key_range},0)),\"\")"
        ))
        for cell in (starter_name, backup_name):
            cell.font = FORMULA_FONT
        for cell in (starter_score, backup_score):
            cell.font = FORMULA_FONT
            cell.number_format = "0.0;(0.0)"

        rv_index = ws.cell(row=row, column=7, value=(
            f'=IF(OR(D{row}="",F{row}=""),"",D{row}-F{row})'
        ))
        rv_game = ws.cell(row=row, column=8, value=(
            f'=IF(G{row}="","",G{row}*\'Model Assumptions\'!$C$107)'
        ))
        rv_index.font = FORMULA_FONT
        rv_game.font = FORMULA_FONT
        rv_index.number_format = "0.0;(0.0)"
        rv_game.number_format = "0.00;(0.00)"

    rv_game_col = "H"
    rv_team_col = "A"
    rv_game_range = f"'{SHEET_NAME}'!${rv_game_col}${sec6_first_row}:${rv_game_col}${sec6_last_row}"
    rv_team_range = f"'{SHEET_NAME}'!${rv_team_col}${sec6_first_row}:${rv_team_col}${sec6_last_row}"

    # ---- Closing note --------------------------------------------------------------------
    note_row = sec6_last_row + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
    note = ws.cell(row=note_row, column=1, value=(
        "claude_code_spec_defensive_player_index.md. HONESTY REQUIREMENT: these are real "
        "counting stats (sacks, TFL, QB hits), genuinely attributed to individual players "
        "in public data, converted to real per-snap RATES (not raw totals -- a rotational "
        "player is never penalized purely for playing fewer real snaps, MIN_QUALIFYING_"
        "DEFENSIVE_SNAPS=200 excludes anyone below that entirely, not zero-filled). But "
        "they are NOISIER and more scheme/teammate-dependent than QB Index's EPA-based "
        "metrics -- an EDGE rusher's sack total is meaningfully influenced by the other "
        "three real pass rushers and the defensive scheme, in a way a QB's EPA/play isn't "
        "influenced by his receivers to the same degree. DO NOT read these scores with "
        "the same implied confidence as QB Index. Real per-season rookie identification "
        "(current_roster.attach_real_rookie_season) was used from the START on this tab -- "
        "not the flawed 'first pulled season' proxy four other tabs needed fixing after "
        "the fact. Replacement Value (Section 6) is computed PER REAL SLOT, not one team-"
        "wide Starter/Backup pair -- EDGE1/EDGE2/IDL1/IDL2 each have their own real "
        "Starter and (where a team has one) real Backup. IDL slots are scheme-variable "
        "(real LDT/RDT/NT priority order, up to 2 real slots/team, same priority Front "
        "Seven Index's own Section 6 already uses) -- a team that only runs 1 real "
        "interior-line position type has no IDL2 row at all. NO man/zone coverage "
        "tendency, no individual pass-rush-win-rate, no position-specific matchup data "
        "(LT vs. EDGE) anywhere on this tab -- all confirmed proprietary, none "
        "approximated with a fabricated proxy. The 2026 depth-chart snapshot this tab's "
        "population is built from was pulled BEFORE final 53-man roster cuts -- same "
        "caveat as every other current-roster-driven tab in this workbook."
    ))
    note.font = NOTE_FONT
    note.alignment = Alignment(wrap_text=True, vertical="top")

    # ==== Wire into Team Ratings (new, unconditional additive column -- SUMS all real slots)
    tr = wb[TEAM_RATINGS_SHEET]
    tr_col = 25  # Y -- next free after WR/TE Corps Quality's X (24)
    tr.cell(row=2, column=tr_col, value="EDGE/IDL Repl.\nValue Adj (pts)").font = HEADER_FONT
    tr.cell(row=2, column=tr_col).fill = HEADER_FILL
    tr.cell(row=2, column=tr_col).alignment = HEADER_ALIGN

    for row in range(3, 3 + len(TEAM_ORDER)):
        adj = tr.cell(row=row, column=tr_col, value=(
            f'=SUMIF({rv_team_range},A{row},{rv_game_range})'
        ))
        adj.font = LINK_FONT
        adj.number_format = "0.00;(0.00)"

    note_row_tr = 3 + len(TEAM_ORDER) + 10
    tr.merge_cells(start_row=note_row_tr, start_column=1, end_row=note_row_tr, end_column=tr_col)
    tr_note = tr.cell(row=note_row_tr, column=1, value=(
        "EDGE/IDL Repl. Value Adj (Y) sums ALL of that team's real per-slot Replacement "
        "Value (Game Points) from 'EDGE-IDL Index' Section 6 (up to 4 real slots: EDGE1/"
        "EDGE2/IDL1/IDL2) into one Team Ratings column, unconditional (no manual toggle -- "
        "each slot's own Backup-In/Starter-In reality is already reflected by whichever "
        "real player Section 3/5 currently scores at that slot). Applies to Net Power "
        "Rating (N) for every team, alongside every other adjustment."
    ))
    tr_note.font = NOTE_FONT
    tr_note.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(workbook_path)
    print(
        f"Built '{SHEET_NAME}': Section 1 {len(season_stats)} rows, Section 3/5 {n_players} "
        f"players, Section 6 {n_slots} real slots. Wired into '{TEAM_RATINGS_SHEET}' (col Y)."
    )
    print(f"Saved to {workbook_path}")
    return {
        "sec1_rows": len(season_stats), "n_players": n_players, "n_slots": n_slots,
        "sec3_range": (sec3_first_row, sec3_last_row),
        "sec5_range": (sec5_first_row, sec5_last_row),
        "sec6_range": (sec6_first_row, sec6_last_row),
    }


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print('Usage: uv run python scripts/build_edge_idl_index.py "path/to/workbook.xlsx"')
        sys.exit(1)
    build(sys.argv[1])
